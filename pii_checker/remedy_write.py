"""対処（remediate）で、ファイルの中身を書き換える部分（csv・tsv・txt・xlsx・xlsm とダミーの作り方）。

決まり:
- 元のファイルには書かない。書き換えた結果は別のファイル（一時の名前）へ書き、呼び出し側が置き換える。
- 位置（行・列・シート・値の開始と終了）は検査のときに集めたもの（engine.Collector）。生の値は持たないので、
  ダミー化は書き換えるときにその場でセルの文字を読み直し、位置の部分だけを置き換える。
- 触らない行・セルは元のまま（csv・txt は元の文字列をそのまま写す＝文字コード・BOM・改行・引用符の付け方が変わらない。
  xlsx のダミー化は zip の中の XML のうち、対象のセルだけを書き換える）。
"""

from __future__ import annotations

import codecs
import csv
import hashlib
import hmac
import io
import os
import re
import unicodedata
import zipfile
from bisect import bisect_left, bisect_right
from dataclasses import dataclass, field

from . import detectors

# 書き換える csv・txt を丸ごとメモリへ読む上限（これより大きいものは書き換えない）
TEXT_REWRITE_LIMIT = 256 * 1024**2
# 検査（misc.TEXT_READ_ALL_LIMIT）がこれより大きいテキストを別の読み方（物理行で数える）にするため、
# 行番号が食い違いうる特殊な改行文字があるときは、この大きさを超えるファイルを書き換えない
HUGE_TEXT_BYTES = 64 * 1024**2
RE_EXOTIC_BREAK = re.compile("[\x0b\x0c\x1c\x1d\x1e\x85  ]")
FIXED_DUMMY = "＊＊＊"
DUMMY_MODES = {
    "same": "同じ値は同じダミー（突き合わせが壊れない）",
    "shape": "形を保つ（数字は 0、文字は 〇 に。区切りの記号は残す）",
    "fixed": "固定の「＊＊＊」",
}


class RemedyError(Exception):
    """対処できない・途中で失敗した（文面はそのまま利用者に見せる。値は含めない）。"""


# ---------------------------------------------------------------- 位置の選び方


@dataclass
class Selection:
    """利用者が選んだ種類・確度・シートに当たる位置の集まり。"""

    groups: dict = field(default_factory=dict)  # (シート名, 種類, 確度の順位) → array('q')

    @property
    def detections(self) -> int:
        return sum(len(a) // 4 for a in self.groups.values())

    def sheets(self) -> list[str]:
        return sorted({k[0] for k in self.groups if len(self.groups[k])})

    def rows(self) -> dict[str, set[int]]:
        """シートごとの行（1 始まり）の集合。"""
        out: dict[str, set[int]] = {}
        for (sheet, _t, _r), arr in self.groups.items():
            out.setdefault(sheet, set()).update(arr[0::4])
        return out

    def cols(self) -> dict[str, set[int]]:
        """シートごとの列（0 始まり）の集合。"""
        out: dict[str, set[int]] = {}
        for (sheet, _t, _r), arr in self.groups.items():
            out.setdefault(sheet, set()).update(arr[1::4])
        return out

    def cell_count(self) -> int:
        cells: set[int] = set()
        for (sheet, _t, _r), arr in self.groups.items():
            h = hash(sheet)
            cells.update((h, arr[i], arr[i + 1]) for i in range(0, len(arr), 4))
        return len(cells)

    def spans(self) -> dict[str, dict[tuple[int, int], list[tuple[int, int, str]]]]:
        """シート → (行, 列) → [(開始, 終了, 種類)]。"""
        out: dict[str, dict[tuple[int, int], list]] = {}
        for (sheet, pii_type, _r), arr in self.groups.items():
            d = out.setdefault(sheet, {})
            for i in range(0, len(arr), 4):
                d.setdefault((arr[i], arr[i + 1]), []).append((arr[i + 2], arr[i + 3], pii_type))
        return out


def select_positions(positions: dict, types=None, ranks=None, sheets=None) -> Selection:
    """集めた位置から、種類・確度の順位・シートで絞る（None は絞らない）。"""
    sel = Selection()
    for key, arr in positions.items():
        sheet, pii_type, rank = key
        if types is not None and pii_type not in types:
            continue
        if ranks is not None and rank not in ranks:
            continue
        if sheets is not None and sheet not in sheets:
            continue
        if len(arr):
            sel.groups[key] = arr
    return sel


# ---------------------------------------------------------------- ダミー


class DummyMaker:
    """ダミーの値を作る。"same" は鍵つきの HMAC で、同じ値には同じダミーを返す（鍵は保存しない）。"""

    def __init__(self, mode: str, key: bytes | None = None):
        if mode not in DUMMY_MODES:
            raise RemedyError("ダミーの作り方が正しくありません")
        self.mode = mode
        self.key = key or os.urandom(32)
        # この回で作ったダミーの文字列（検査し直しで「ダミーそのもの」を数えないために呼び出し側が使う。値ではない）
        self.produced: set[str] = set()

    def make(self, pii_type: str, value: str) -> str:
        out = self._make(pii_type, value)
        self.produced.add(out)
        return out

    def _make(self, pii_type: str, value: str) -> str:
        if self.mode == "fixed":
            return FIXED_DUMMY
        if self.mode == "shape":
            return shape_of(value)
        from .engine import distinct_key

        digest = hmac.new(self.key, f"{pii_type}\0{distinct_key(pii_type, value)}".encode("utf-8"), hashlib.sha256)
        return f"{detectors.PII_TYPES.get(pii_type, pii_type)}_{digest.hexdigest()[:8]}"


def shape_of(value: str) -> str:
    """形を保つダミー: 数字は 0・文字（英字・日本語）は 〇。区切りの記号・空白はそのまま残す。

    英字を x にすると「xxxx@xxxx.xxx」がメールアドレスとして再び検出されるので、文字は 〇 にそろえる。
    """
    out = []
    for ch in value:
        if ch.isdigit():
            out.append("0")
        elif ch.isalpha() or ch in "々〆ヶー":
            out.append("〇")
        else:
            out.append(ch)
    return "".join(out)


def _orig_bounds(text: str, norm: str):
    """正規化（NFKC）した文字列の位置を、元の文字列の位置へ直す関数を返す。直せないときは None。"""
    if text == norm:
        return lambda s, e: (s, e)
    pieces = [unicodedata.normalize("NFKC", ch) for ch in text]
    if "".join(pieces) != norm:
        return None
    offs = []
    pos = 0
    for p in pieces:
        offs.append(pos)
        pos += len(p)

    def bounds(s: int, e: int) -> tuple[int, int]:
        return bisect_right(offs, s) - 1, bisect_right(offs, max(s, e - 1))

    return bounds


def replace_spans(text: str, spans: list[tuple[int, int, str]], maker: DummyMaker) -> str:
    """文字列の中の、検出した値の部分（正規化後の位置 spans）だけをダミーに置き換える。

    位置が元の文字へ直せないとき（NFKC が文字をまたいで変わる）は、安全側に倒してセル・行の全体をダミーにする。
    重なる位置は1つにまとめる（長いほうの種類でダミーを作る）。
    """
    norm = detectors.normalize(text)
    spans = sorted((s, min(e, len(norm)), t) for s, e, t in spans if 0 <= s < min(e, len(norm)))
    if not spans:
        return text
    merged: list[list] = []
    for s, e, t in spans:
        if merged and s < merged[-1][1]:
            if e - s > merged[-1][1] - merged[-1][0]:
                merged[-1][2] = t
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e, t])
    bounds = _orig_bounds(text, norm)
    if bounds is None:
        return maker.make(merged[0][2], text)
    out = []
    last = 0
    for s, e, t in merged:
        os_, oe = bounds(s, e)
        out.append(text[last:os_])
        out.append(maker.make(t, norm[s:e]))
        last = oe
    out.append(text[last:])
    return "".join(out)


# ---------------------------------------------------------------- 文字コード


def decode_for_rewrite(data: bytes) -> tuple[str, str, bytes]:
    """書き換えるテキストを（文字列, 文字コード, BOM のバイト列）にする。元のバイト列へ正確に戻せないときは断る。"""
    from .textdecode import decode_text

    bom = b""
    name = ""
    if data.startswith(codecs.BOM_UTF8):
        name, bom = "utf-8", codecs.BOM_UTF8
    elif data.startswith(codecs.BOM_UTF16_LE):
        name, bom = "utf-16-le", codecs.BOM_UTF16_LE
    elif data.startswith(codecs.BOM_UTF16_BE):
        name, bom = "utf-16-be", codecs.BOM_UTF16_BE
    if name:
        try:
            text = data[len(bom) :].decode(name)
        except UnicodeDecodeError:
            raise RemedyError("文字コードを読み取れないため、書き換えできません") from None
    else:
        text, label = decode_text(data)
        if label not in ("utf-8", "cp932", "euc_jp", "utf-16-le", "utf-16-be"):
            raise RemedyError(
                f"文字コード（{label}）のファイルは、書き換えに対応していません（ファイルの削除はできます）"
            )
        name = label
    try:
        again = bom + text.encode(name)
    except UnicodeEncodeError:
        again = b""
    if again != data:
        raise RemedyError(
            "文字コードを変換して戻しても元と同じ中身にならないため、書き換えません（元のファイルは変えていません）"
        )
    return text, name, bom


def encode_for_rewrite(text: str, name: str, bom: bytes) -> bytes:
    try:
        return bom + text.encode(name)
    except UnicodeEncodeError:
        raise RemedyError("置き換えた文字を元の文字コードに直せませんでした（元のファイルは変えていません）") from None


# ---------------------------------------------------------------- csv・tsv


def _quote(text: str) -> str:
    return '"' + text.replace('"', '""') + '"'


def split_raw_fields(raw: str, delim: str) -> list[tuple[str, bool]]:
    """csv の1レコードの文字列（行末は除く）を、元の書き方のままの欄（文字列, 引用符つきか）に分ける。

    csv モジュールと同じ読み方（引用符で始まる欄だけ引用符として読み、閉じた後の文字はそのまま続く）。
    """
    fields: list[tuple[str, bool]] = []
    i, n = 0, len(raw)
    if n == 0:
        return fields
    while True:
        start = i
        if i < n and raw[i] == '"':
            i += 1
            while i < n:
                if raw[i] == '"':
                    if i + 1 < n and raw[i + 1] == '"':
                        i += 2
                        continue
                    i += 1
                    break
                i += 1
            while i < n and raw[i] != delim:
                i += 1
            fields.append((raw[start:i], True))
        else:
            while i < n and raw[i] != delim:
                i += 1
            fields.append((raw[start:i], False))
        if i >= n:
            break
        i += 1  # 区切り
        if i >= n:
            fields.append(("", False))  # 行末の区切りの後の空の欄
            break
    return fields


def _split_terminator(raw: str) -> tuple[str, str]:
    for t in ("\r\n", "\n", "\r"):
        if raw.endswith(t):
            return raw[: -len(t)], t
    return raw, ""


def rewrite_delimited(
    text: str,
    delim: str,
    method: str,
    sel: Selection,
    maker: DummyMaker | None,
) -> tuple[str, dict]:
    """csv・tsv の文字列を書き換える。触らないレコードは元の文字列のまま写す。"""
    rows = sel.rows().get("", set())
    cols = sel.cols().get("", set())
    spans = sel.spans().get("", {})
    span_rows: dict[int, dict[int, list]] = {}
    for (r, c), lst in spans.items():
        span_rows.setdefault(r, {})[c] = lst
    csv.field_size_limit(2**31 - 1)
    consumed: list[str] = []

    def tee():
        for ln in io.StringIO(text):
            consumed.append(ln)
            yield ln

    reader = csv.reader(tee(), delimiter=delim)
    out: list[str] = []
    stats = {"rows": 0, "cols": 0, "cells": 0, "detections": 0}
    r = 0
    while True:
        consumed.clear()
        try:
            fields = next(reader)
        except StopIteration:
            break
        except csv.Error as e:
            raise RemedyError(f"csv として読み取れない行があります（{e}）。書き換えていません") from None
        r += 1
        raw = "".join(consumed)
        if method == "delete_rows":
            if r in rows:
                stats["rows"] += 1
                continue
            out.append(raw)
            continue
        if method == "delete_columns":
            if not cols or not fields:
                out.append(raw)
                continue
            body, term = _split_terminator(raw)
            raws = split_raw_fields(body, delim)
            if len(raws) != len(fields):
                raise RemedyError(f"{r} 行目の区切りを元の書き方のまま分けられないため、書き換えません")
            kept = [rf for i, (rf, _q) in enumerate(raws) if i not in cols]
            if len(kept) != len(raws):
                out.append(delim.join(kept) + term)
            else:
                out.append(raw)
            continue
        if method == "dummy":
            targets = span_rows.get(r)
            if not targets or not fields:
                out.append(raw)
                continue
            body, term = _split_terminator(raw)
            raws = split_raw_fields(body, delim)
            if len(raws) != len(fields):
                raise RemedyError(f"{r} 行目の区切りを元の書き方のまま分けられないため、書き換えません")
            changed = False
            for c, lst in targets.items():
                if c >= len(fields):
                    continue
                new = replace_spans(fields[c], lst, maker)
                if new == fields[c]:
                    continue
                was_quoted = raws[c][1]
                need = was_quoted or any(ch in new for ch in (delim, '"', "\r", "\n"))
                raws[c] = (_quote(new) if need else new, was_quoted)
                stats["cells"] += 1
                stats["detections"] += len(lst)
                changed = True
            out.append(delim.join(rf for rf, _q in raws) + term if changed else raw)
            continue
        raise RemedyError("この方法は csv に使えません")
    if method == "delete_columns":
        stats["cols"] = len(cols)
    return "".join(out), stats


# ---------------------------------------------------------------- txt


def rewrite_lines(text: str, method: str, sel: Selection, maker: DummyMaker | None) -> tuple[str, dict]:
    """テキストを行単位で書き換える（行は str.splitlines と同じ区切り。検査の行番号と同じ数え方）。"""
    rows = sel.rows().get("", set())
    spans = sel.spans().get("", {})
    by_row: dict[int, list] = {}
    for (r, _c), lst in spans.items():
        by_row.setdefault(r, []).extend(lst)
    out: list[str] = []
    stats = {"rows": 0, "cols": 0, "cells": 0, "detections": 0}
    for i, line in enumerate(text.splitlines(keepends=True), start=1):
        if method == "delete_rows":
            if i in rows:
                stats["rows"] += 1
                continue
        elif method == "dummy" and i in by_row:
            body = line.splitlines()[0] if line.splitlines() else ""
            term = line[len(body) :]
            new = replace_spans(body, by_row[i], maker)
            if new != body:
                stats["cells"] += 1  # 行を1つと数える
                stats["detections"] += len(by_row[i])
                line = new + term
        elif method not in ("delete_rows", "dummy"):
            raise RemedyError("この方法はテキストに使えません")
        out.append(line)
    return "".join(out), stats


def rewrite_text_file(src: str, dst: str, layout: str, method: str, sel: Selection, maker: DummyMaker | None) -> dict:
    """csv・tsv・txt のファイル src を書き換えた結果を dst へ書く。"""
    from .engine import long_path

    size = os.path.getsize(long_path(src))
    if size > TEXT_REWRITE_LIMIT:
        raise RemedyError(
            f"大きすぎるため書き換えできません（{TEXT_REWRITE_LIMIT // 1024**2}MB まで。ファイルの削除はできます）"
        )
    with open(long_path(src), "rb") as f:
        data = f.read()
    text, name, bom = decode_for_rewrite(data)
    if layout == "lines" and size > HUGE_TEXT_BYTES and RE_EXOTIC_BREAK.search(text):
        raise RemedyError("特殊な改行文字を含む大きなファイルは、行の位置を確かめられないため書き換えません")
    if layout.startswith("csv"):
        new_text, stats = rewrite_delimited(text, layout[3:], method, sel, maker)
    elif layout == "lines":
        new_text, stats = rewrite_lines(text, method, sel, maker)
    else:
        raise RemedyError("この形式の中身の書き換えには対応していません")
    out = encode_for_rewrite(new_text, name, bom)
    with open(long_path(dst), "wb") as f:
        f.write(out)
        f.flush()
        os.fsync(f.fileno())
    return stats


# ---------------------------------------------------------------- xlsx・xlsm


XLSX_WARNINGS = [
    "行・列を消すと、数式の参照（別の行・列を指す式）はずれたままになります。消した行・列を指していた式は壊れます",
    "条件付き書式・入力規則・オートフィルタ・印刷範囲・名前の定義は、行・列の位置に合わせて直しません",
    "保存し直すため、図形・一部のグラフ・マクロの画面部品などは失われることがあります。数式の計算結果の控えも消えます（Excel で開くと計算し直されます）",
]


def _col_letters(idx0: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(idx0 + 1)


def _shifted_span(lo: int, hi: int, removed: list[int]) -> tuple[int, int] | None:
    """範囲 [lo, hi]（行または列の番号）から、消した番号（昇順）を詰めたあとの範囲。全部消えたら None。"""
    gone = bisect_right(removed, hi) - bisect_left(removed, lo)
    kept = hi - lo + 1 - gone
    if kept <= 0:
        return None
    new_lo = lo - bisect_left(removed, lo)
    return new_lo, new_lo + kept - 1


def _fix_ref_rows(ref: str | None, rows: list[int]) -> str | None:
    """範囲の文字列（"A1:E20"）を、行を消したあとの範囲にする。全部消えたら None。"""
    from openpyxl.utils import get_column_letter, range_boundaries

    if not ref:
        return ref
    min_c, min_r, max_c, max_r = range_boundaries(ref)
    span = _shifted_span(min_r, max_r, rows)
    if span is None:
        return None
    return f"{get_column_letter(min_c)}{span[0]}:{get_column_letter(max_c)}{span[1]}"


def _fix_ref_cols(ref: str | None, cols: list[int]) -> str | None:
    """範囲の文字列を、列（1 始まり）を消したあとの範囲にする。全部消えたら None。"""
    from openpyxl.utils import get_column_letter, range_boundaries

    if not ref:
        return ref
    min_c, min_r, max_c, max_r = range_boundaries(ref)
    span = _shifted_span(min_c, max_c, cols)
    if span is None:
        return None
    return f"{get_column_letter(span[0])}{min_r}:{get_column_letter(span[1])}{max_r}"


def _fix_tables_rows(ws, rows: list[int]) -> None:
    """行を消したあと、Excel の表（ListObject）の範囲・オートフィルタ・シートのオートフィルタを合わせる。
    表の見出し行・集計行が消える、または表のデータ行が1行も残らないときは、壊れた表を作らないよう断る。"""
    from openpyxl.utils import get_column_letter, range_boundaries

    for table in list(ws.tables.values()):  # items() は（名前, 範囲の文字列）なので、表そのものを values() で取る
        name = table.name
        min_c, min_r, max_c, max_r = range_boundaries(table.ref)
        header = 1 if table.headerRowCount is None else int(table.headerRowCount)
        totals = int(table.totalsRowCount or 0)
        in_table = [r for r in rows if min_r <= r <= max_r]
        if not in_table:
            span = _shifted_span(min_r, max_r, rows)
            new_min, new_max = span
        else:
            head_rows = range(min_r, min_r + header)
            tail_rows = range(max_r - totals + 1, max_r + 1) if totals else range(0)
            if any(r in head_rows or r in tail_rows for r in in_table):
                raise RemedyError(
                    f"シート「{ws.title}」の Excel の表「{name}」の見出し行（または集計行）が、削除する行に含まれています。"
                    "表が壊れるため、行を削除できません。表の見出し・集計行を含まない選びにしてください。元のファイルは変えていません"
                )
            data_left = (max_r - min_r + 1 - header - totals) - len(in_table)
            if data_left <= 0:
                raise RemedyError(
                    f"シート「{ws.title}」の Excel の表「{name}」のデータ行がすべて削除対象です。"
                    "データ行が1行も無い表は作れないため、行を削除できません（表ごと消すなら、シートを削除してください）。"
                    "元のファイルは変えていません"
                )
            span = _shifted_span(min_r, max_r, rows)
            new_min, new_max = span
        table.ref = f"{get_column_letter(min_c)}{new_min}:{get_column_letter(max_c)}{new_max}"
        if table.autoFilter is not None:
            table.autoFilter.ref = f"{get_column_letter(min_c)}{new_min}:{get_column_letter(max_c)}{new_max - totals}"
        if getattr(table, "sortState", None) is not None:
            table.sortState = (
                None  # 並べ替えの範囲は古い行番号のまま。残すと壊れるので外す（並べ替えの結果の順は保たれる）
            )
    if ws.auto_filter is not None and ws.auto_filter.ref:
        new = _fix_ref_rows(ws.auto_filter.ref, rows)
        ws.auto_filter.ref = new
        if new is None and getattr(ws.auto_filter, "sortState", None) is not None:
            ws.auto_filter.sortState = None


def _fix_tables_cols(ws, cols: list[int]) -> None:
    """列を消したあと、Excel の表の範囲・列の定義（tableColumns）・オートフィルタを合わせる。列が全部消える表は取り除く。"""
    from openpyxl.utils import get_column_letter, range_boundaries

    colset = set(cols)
    for table in list(ws.tables.values()):  # items() は（名前, 範囲の文字列）なので、表そのものを values() で取る
        name = table.name
        min_c, min_r, max_c, max_r = range_boundaries(table.ref)
        span = _shifted_span(min_c, max_c, cols)
        if span is None:
            del ws.tables[name]  # 表の列がすべて消える（データも消える）。表の定義だけ残さない
            continue
        new_min, new_max = span
        kept_idx = [i for i in range(max_c - min_c + 1) if (min_c + i) not in colset]
        if len(table.tableColumns) == max_c - min_c + 1:
            table.tableColumns = [table.tableColumns[i] for i in kept_idx]
        table.ref = f"{get_column_letter(new_min)}{min_r}:{get_column_letter(new_max)}{max_r}"
        af = table.autoFilter
        if af is not None:
            totals = int(table.totalsRowCount or 0)
            af.ref = f"{get_column_letter(new_min)}{min_r}:{get_column_letter(new_max)}{max_r - totals}"
            new_filters = []
            for fc in list(af.filterColumn):
                if fc.colId in kept_idx:
                    fc.colId = kept_idx.index(fc.colId)
                    new_filters.append(fc)
            af.filterColumn = new_filters
        if getattr(table, "sortState", None) is not None:
            table.sortState = None
    if ws.auto_filter is not None and ws.auto_filter.ref:
        ws.auto_filter.ref = _fix_ref_cols(ws.auto_filter.ref, cols)


def _delete_rows(ws, rows: list[int]) -> None:
    """シートから行を消す（セル・結合・行の高さを詰める。数式の参照は直さない）。"""
    rows = sorted(set(rows))
    rowset = set(rows)
    from openpyxl.worksheet.cell_range import CellRange

    _fix_tables_rows(ws, rows)

    merged = [(rg.min_row, rg.min_col, rg.max_row, rg.max_col) for rg in ws.merged_cells.ranges]
    for rg in list(ws.merged_cells.ranges):
        ws.merged_cells.remove(rg)
    new_cells = {}
    for (r, c), cell in ws._cells.items():
        if r in rowset:
            continue
        nr = r - bisect_left(rows, r)
        cell.row = nr
        new_cells[(nr, c)] = cell
        if getattr(cell, "_hyperlink", None) is not None:
            cell._hyperlink.ref = cell.coordinate
    ws._cells = new_cells
    old_dims = dict(ws.row_dimensions)
    ws.row_dimensions.clear()
    for r, dim in old_dims.items():
        if r in rowset:
            continue
        nr = r - bisect_left(rows, r)
        dim.index = nr
        ws.row_dimensions[nr] = dim
    for min_r, min_c, max_r, max_c in merged:
        kept = [r for r in range(min_r, max_r + 1) if r not in rowset] if max_r - min_r < 100000 else None
        if kept is None:
            continue
        if not kept:
            continue
        nmin = kept[0] - bisect_left(rows, kept[0])
        nmax = nmin + len(kept) - 1
        if nmin == nmax and min_c == max_c:
            continue
        ws.merged_cells.add(CellRange(min_col=min_c, min_row=nmin, max_col=max_c, max_row=nmax))


def _delete_cols(ws, cols1: list[int]) -> None:
    """シートから列（1 始まり）を消す（セル・結合・列の幅を詰める。数式の参照は直さない）。"""
    cols = sorted(set(cols1))
    colset = set(cols)
    from openpyxl.utils import column_index_from_string, get_column_letter
    from openpyxl.worksheet.cell_range import CellRange

    _fix_tables_cols(ws, cols)

    merged = [(rg.min_row, rg.min_col, rg.max_row, rg.max_col) for rg in ws.merged_cells.ranges]
    for rg in list(ws.merged_cells.ranges):
        ws.merged_cells.remove(rg)
    new_cells = {}
    for (r, c), cell in ws._cells.items():
        if c in colset:
            continue
        nc = c - bisect_left(cols, c)
        cell.column = nc
        new_cells[(r, nc)] = cell
        if getattr(cell, "_hyperlink", None) is not None:
            cell._hyperlink.ref = cell.coordinate
    ws._cells = new_cells
    new_dims = {}
    for key, dim in list(ws.column_dimensions.items()):
        lo = dim.min or column_index_from_string(key)
        hi = dim.max or lo
        kept = [c for c in range(lo, min(hi, 16384) + 1) if c not in colset]
        if not kept:
            continue
        nlo = kept[0] - bisect_left(cols, kept[0])
        nhi = nlo + len(kept) - 1
        dim.min, dim.max = nlo, nhi
        dim.index = get_column_letter(nlo)
        new_dims[dim.index] = dim
    ws.column_dimensions.clear()
    for k, dim in new_dims.items():
        ws.column_dimensions[k] = dim
    for min_r, min_c, max_r, max_c in merged:
        kept = [c for c in range(min_c, max_c + 1) if c not in colset]
        if not kept:
            continue
        nmin = kept[0] - bisect_left(cols, kept[0])
        nmax = nmin + len(kept) - 1
        if nmin == nmax and min_r == max_r:
            continue
        ws.merged_cells.add(CellRange(min_col=nmin, min_row=min_r, max_col=nmax, max_row=max_r))


def rewrite_xlsx_structure(src: str, dst: str, method: str, sel: Selection, keep_vba: bool) -> dict:
    """行・列・シートの削除（openpyxl で読み込み直して保存する）。"""
    import openpyxl

    from .engine import long_path

    try:
        wb = openpyxl.load_workbook(long_path(src), keep_vba=keep_vba)
    except Exception as e:
        raise RemedyError(f"ブックを読み込めませんでした（{type(e).__name__}）。元のファイルは変えていません") from None
    stats = {"rows": 0, "cols": 0, "sheets": 0, "cells": 0, "detections": 0}
    try:
        if method == "delete_sheet":
            targets = [s for s in sel.sheets() if s in wb.sheetnames]
            if len(targets) >= len(wb.sheetnames):
                raise RemedyError(
                    "ブックのすべてのシートは削除できません（シートが1つも残らなくなります）。ファイル全体を消すなら「ファイルを削除」を選んでください"
                )
            for name in targets:
                del wb[name]
            stats["sheets"] = len(targets)
        elif method == "delete_rows":
            for sheet, rows in sel.rows().items():
                if sheet in wb.sheetnames:
                    _delete_rows(wb[sheet], list(rows))
                    stats["rows"] += len(rows)
        elif method == "delete_columns":
            for sheet, cols in sel.cols().items():
                if sheet in wb.sheetnames:
                    _delete_cols(wb[sheet], [c + 1 for c in cols])
                    stats["cols"] += len(cols)
        else:
            raise RemedyError("この方法は選べません")
        try:
            wb.save(long_path(dst))
        except Exception as e:
            raise RemedyError(
                f"ブックを保存できませんでした（{type(e).__name__}）。元のファイルは変えていません"
            ) from None
    finally:
        wb.close()
    return stats


RE_CELL = re.compile(r"<c\b((?:\s+[\w:.-]+=\"[^\"]*\")*)\s*(/>|>(.*?)</c>)", re.S)
RE_ATTR = re.compile(r"([\w:.-]+)=\"([^\"]*)\"")
RE_REF = re.compile(r"^([A-Z]+)(\d+)$")
RE_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _col_index0(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _xml_escape(text: str) -> str:
    text = RE_XML_BAD.sub("", text)
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _sheet_parts(z: zipfile.ZipFile) -> dict[str, str]:
    """シート名 → zip の中のシートの部品名。"""
    import xml.etree.ElementTree as ET

    ns_main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    ns_rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    out = {}
    for sh in wb.iter(f"{ns_main}sheet"):
        target = targets.get(sh.get(f"{ns_rel}id"))
        if not target:
            continue
        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        out[sh.get("name")] = part
    return out


def _shared_strings_text(xml: str) -> list[str]:
    """sharedStrings.xml の各項目の文字（リッチテキストは run をつなぐ）。"""
    import html

    out = []
    for si in re.finditer(r"<si\b[^>]*?(?:/>|>.*?</si>)", xml, re.S):
        out.append(html.unescape("".join(m.group(1) for m in re.finditer(r"<t\b[^>]*>(.*?)</t>", si.group(0), re.S))))
    return out


def rewrite_xlsx_dummy(src: str, dst: str, sel: Selection, maker: DummyMaker) -> dict:
    """該当セルの検出した値だけをダミーに置き換える（zip の中の XML のうち、対象のセルだけを書き換える）。

    セルは文字列（inlineStr）に置き換える。図形・グラフ・ピボット・書式・マクロなど、ほかの部品は一切触らない。
    数式のセルは書き換えない（式を壊さないため）。置き換えて参照されなくなった共有文字列は空にする（元の値が残らないように）。
    """
    import openpyxl

    from .engine import _cell_to_str, long_path

    spans_by_sheet = sel.spans()
    stats = {"rows": 0, "cols": 0, "sheets": 0, "cells": 0, "detections": 0, "skipped_formula": 0}
    # 1) 検査と同じ読み方で、対象セルの文字を得る
    texts: dict[tuple[str, int, int], str] = {}
    try:
        wb = openpyxl.load_workbook(long_path(src), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet, cells in spans_by_sheet.items():
                if sheet not in wb.sheetnames:
                    continue
                ws = wb[sheet]
                if hasattr(ws, "reset_dimensions"):
                    ws.reset_dimensions()
                want_rows = {r for r, _c in cells}
                for r_idx, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True), start=1):
                    if r_idx not in want_rows:
                        continue
                    for c_idx, v in enumerate(row):
                        if (r_idx, c_idx) in cells and v is not None:
                            texts[(sheet, r_idx, c_idx)] = _cell_to_str(v)
        finally:
            wb.close()
    except RemedyError:
        raise
    except Exception as e:
        raise RemedyError(f"ブックを読み取れませんでした（{type(e).__name__}）。元のファイルは変えていません") from None

    with zipfile.ZipFile(long_path(src)) as zin:
        parts = _sheet_parts(zin)
        sst_name = "xl/sharedStrings.xml"
        sst_texts: list[str] | None = None
        replaced: dict[str, bytes] = {}
        released_sst: set[int] = set()
        for sheet, cells in spans_by_sheet.items():
            part = parts.get(sheet)
            if part is None or part not in zin.namelist():
                continue
            xml = zin.read(part).decode("utf-8")

            def sub(m: re.Match) -> str:
                nonlocal sst_texts
                attrs = dict(RE_ATTR.findall(m.group(1)))
                ref = RE_REF.match(attrs.get("r", ""))
                if not ref:
                    return m.group(0)
                key = (int(ref.group(2)), _col_index0(ref.group(1)))
                if key not in cells:
                    return m.group(0)
                inner = m.group(3) or ""
                if "<f" in inner:
                    stats["skipped_formula"] += 1
                    return m.group(0)
                text = texts.get((sheet, key[0], key[1]))
                if text is None:
                    return m.group(0)
                new = replace_spans(text, cells[key], maker)
                if new == text:
                    return m.group(0)
                if attrs.get("t") == "s":
                    v = re.search(r"<v>\s*(\d+)\s*</v>", inner)
                    if v:
                        released_sst.add(int(v.group(1)))
                attrs["t"] = "inlineStr"
                names = [n for n, _v in RE_ATTR.findall(m.group(1))]
                if "t" not in names:
                    names.append("t")
                attr_text = " ".join(f'{n}="{attrs[n]}"' for n in names)
                stats["cells"] += 1
                stats["detections"] += len(cells[key])
                return f'<c {attr_text}><is><t xml:space="preserve">{_xml_escape(new)}</t></is></c>'

            new_xml = RE_CELL.sub(sub, xml)
            if new_xml != xml:
                replaced[part] = new_xml.encode("utf-8")
        # 2) 参照されなくなった共有文字列を空にする（元の値が部品の中に残らないように）
        if released_sst and sst_name in zin.namelist():
            still: set[int] = set()
            for p in parts.values():
                if p not in zin.namelist():
                    continue
                data = replaced.get(p) or zin.read(p)
                for m in re.finditer(
                    r"<c\b[^>]*\bt=\"s\"[^>]*>\s*(?:<f\b.*?</f>\s*)?<v>\s*(\d+)\s*</v>", data.decode("utf-8"), re.S
                ):
                    still.add(int(m.group(1)))
            gone = released_sst - still
            if gone:
                sst_xml = zin.read(sst_name).decode("utf-8")
                idx = [-1]

                def blank(m: re.Match) -> str:
                    idx[0] += 1
                    return "<si><t></t></si>" if idx[0] in gone else m.group(0)

                new_sst = re.sub(r"<si\b[^>]*?(?:/>|>.*?</si>)", blank, sst_xml, flags=re.S)
                replaced[sst_name] = new_sst.encode("utf-8")
        if stats["cells"] == 0:
            raise RemedyError(
                "書き換えるセルが見つかりませんでした（数式のセルは書き換えません）。元のファイルは変えていません"
            )
        with zipfile.ZipFile(long_path(dst), "w") as zout:
            for info in zin.infolist():
                data = replaced.get(info.filename)
                if data is None:
                    data = zin.read(info)
                zi = zipfile.ZipInfo(info.filename, info.date_time)
                zi.compress_type = info.compress_type
                zi.external_attr = info.external_attr
                zi.comment = info.comment
                zout.writestr(zi, data)
    return stats


def rewrite_xlsx_file(
    src: str, dst: str, method: str, sel: Selection, maker: DummyMaker | None, keep_vba: bool
) -> dict:
    if method == "dummy":
        return rewrite_xlsx_dummy(src, dst, sel, maker)
    return rewrite_xlsx_structure(src, dst, method, sel, keep_vba)


def sha256_file(path: str) -> str:
    """ファイルの SHA-256（16進）。"""
    from .engine import long_path

    h = hashlib.sha256()
    with open(long_path(path), "rb") as f:
        while True:
            chunk = f.read(1024 * 1024)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()
