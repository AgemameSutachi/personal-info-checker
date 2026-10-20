"""異なる値の数と、漏えい時の報告の目安（印）の試験。"""

# 番号・氏名・メールはすべて架空（マイナンバーは検査数字だけ合わせた連番）

from __future__ import annotations

import json
import sqlite3
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pii_checker import engine, reportflags  # noqa: E402
from pii_checker.engine import Collector, scan_file  # noqa: E402
from pii_checker.models import FileResult, ItemResult, ScanOptions, Status  # noqa: E402
from pii_checker.store import Store  # noqa: E402


def my_number(n: int) -> str:
    """11桁の連番に検査数字を付けた12桁（架空のマイナンバー）。"""
    body = f"{n:011d}"
    total = 0
    for i in range(1, 12):
        p = int(body[-i])
        q = i + 1 if i <= 6 else i - 5
        total += p * q
    r = total % 11
    return body + str(0 if r <= 1 else 11 - r)


def test_distinct_normalizes_values():
    item = ItemResult(path="t", kind="", status=Status.PENDING)
    col = Collector(item, 20)
    col.add_text("携帯 090-1234-5678", "a")
    col.add_text("携帯 09012345678", "b")
    col.add_text("携帯 ０９０－１２３４－５６７８", "c")
    col.add_text("携帯 080-1111-2222", "d")
    col.finish()
    assert item.finding_counts["phone"] == 4
    assert item.distinct_counts == {"phone:medium": 2}, item.distinct_counts


def test_distinct_name_honorific():
    item = ItemResult(path="t", kind="", status=Status.PENDING)
    col = Collector(item, 20)
    col.add_table([["氏名"], ["鬼瓦 権左衛門"], ["鬼瓦権左衛門様"], ["勅使河原二郎"]], lambda r, c: f"R{r}")
    col.finish()
    assert reportflags.distinct_medium(item.distinct_counts).get("name") == 2, item.distinct_counts


def test_distinct_uses_highest_confidence():
    item = ItemResult(path="t", kind="", status=Status.PENDING)
    col = Collector(item, 20)
    col.add_text(f"伝票 {my_number(5)}", "a")  # 語が無い: 低
    col.add_text(f"個人番号 {my_number(5)}", "b")  # 同じ値を語つきで: 高
    col.finish()
    assert item.distinct_counts == {"my_number:high": 1}, item.distinct_counts


def test_distinct_keeps_value_after_deferred_upgrade():
    # 手がかり待ちの住所を後から 中 へ引き上げても（明細の位置と確度を書き換えても）、異なる値の数え方は残る
    item = ItemResult(path="t", kind="", status=Status.PENDING)
    col = Collector(item, 20)
    col.add_text("住所 東京都千代田区丸の内9-9-9", "a")
    col.add_text("住所 東京都千代田区丸の内９－９－９", "b")
    col.add_text("氏名 鬼瓦 権左衛門", "c")
    col.finish()
    assert reportflags.distinct_medium(item.distinct_counts).get("address") == 1, item.distinct_counts
    assert "address:low" not in item.distinct_counts, item.distinct_counts


def test_distinct_sink_merges_items():
    # 集計先を共有する品目（zip の中の複数のファイルなど）: 先の品目の辞書をそのまま渡した後も、後の品目が正しく合流する
    sink: dict = {}
    first = ItemResult(path="a", kind="", status=Status.PENDING)
    col = Collector(first, 20, distinct_sink=sink)
    col.add_text("住所 東京都千代田区丸の内9-9-9", "a")  # 人の手がかりが無い: 低
    col.finish()
    second = ItemResult(path="b", kind="", status=Status.PENDING)
    col = Collector(second, 20, distinct_sink=sink)
    col.add_text("住所 東京都千代田区丸の内9-9-9", "a")
    col.add_text("氏名 鬼瓦 権左衛門", "b")  # 手がかりがあるので同じ住所が 中
    col.finish()
    third = ItemResult(path="c", kind="", status=Status.PENDING)
    col = Collector(third, 20, distinct_sink=sink)
    col.add_text("携帯 090-1234-5678", "a")
    col.add_text("携帯 080-1111-2222", "b")
    col.finish()
    assert first.distinct_counts == {"address:low": 1}, first.distinct_counts  # 後の合流で品目の数は変わらない
    got = engine._count_by_rank(sink)
    assert got == {"address:medium": 1, "name:medium": 1, "phone:medium": 2}, got


def test_distinct_skips_values_over_limit(monkeypatch):
    # 覚える上限を超えた値は件数だけ数え、異なる値には数えない（後から 中 へ引き上げた値も同じ）
    monkeypatch.setattr(engine, "SEEN_LIMIT", 1)
    item = ItemResult(path="t", kind="", status=Status.PENDING)
    sink: dict = {}
    col = Collector(item, 20, distinct_sink=sink)
    col.add_text("住所 東京都千代田区丸の内9-9-9", "a")  # これで上限
    col.add_text("住所 大阪府大阪市北区梅田9-9-9", "b")  # 上限の外（手がかり待ち）
    col.add_text("氏名 鬼瓦 権左衛門", "c")  # 手がかり: 2つの住所を 中 へ引き上げる
    col.finish()
    assert item.finding_counts["address"] == 2, item.finding_counts
    assert item.conf_counts.get("address:medium") == 2, item.conf_counts
    assert item.distinct_counts == {"address:medium": 1}, item.distinct_counts
    assert engine._count_by_rank(sink) == {"address:medium": 1}, sink


def test_compute_thresholds():
    assert reportflags.compute({"my_number:high": 1}, {"my_number:high": 100})["flags"] == ["my_number"]
    got = reportflags.compute({"my_number:high": 101}, {"my_number:high": 60, "my_number:medium": 41})
    assert got["flags"] == ["my_number", "many_my_number"]
    # 確度 低 だけなら印にしない
    assert reportflags.compute({"my_number:low": 500, "sensitive:low": 3}, {"my_number:low": 500})["flags"] == []
    got = reportflags.compute(
        {"email:medium": 2000, "bank_account:high": 1}, {"email:medium": 1001, "name:medium": 900}
    )
    assert got["flags"] == ["financial", "many_people"]
    assert (got["people_n"], got["people_type"]) == (1001, "email")
    got = reportflags.compute({"sensitive:medium": 1, "credit_card:medium": 1}, {"name:medium": 1000})
    assert got["flags"] == ["sensitive", "financial"]  # 1,000 ちょうどは「超える」に当たらない
    # 異なる値を数えていない（古い版）: 数に頼る印は付けず、人の数も分からない
    got = reportflags.compute({"my_number:high": 5000}, None)
    assert got == {"flags": ["my_number"], "people_n": None, "people_type": ""}
    assert reportflags.decode(reportflags.encode(["sensitive", "x"])) == ["sensitive"]


def test_scan_file_counts_distinct_across_items(tmp_path):
    # zip の中の2つの CSV に同じマイナンバーが入っていても、ファイル全体では1つと数える
    rows_a = "\n".join(f"{my_number(i)}" for i in range(1, 81))
    rows_b = "\n".join(f"{my_number(i)}" for i in range(41, 121))
    p = tmp_path / "番号.zip"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("a.csv", "個人番号\n" + rows_a + "\n")
        z.writestr("b.csv", "個人番号\n" + rows_b + "\n")
    r = scan_file(str(p), ScanOptions(ocr=False))
    per_item = [reportflags.distinct_medium(i.distinct_counts).get("my_number", 0) for i in r.items]
    assert sorted(per_item)[-2:] == [80, 80], per_item
    assert reportflags.distinct_medium(r.distinct_counts)["my_number"] == 120, r.distinct_counts
    mark = reportflags.compute({"my_number:high": 160}, r.distinct_counts)
    assert "many_my_number" in mark["flags"]


def test_store_saves_flags_and_summary(tmp_path):
    store = Store(str(tmp_path / "db.sqlite3"))
    job = store.create_job([str(tmp_path)], ScanOptions())
    top = str(tmp_path / "a.csv")
    store.add_files(job, [(top, 1, 1.0), (top + "2", 1, 1.0)])
    item = ItemResult(path=top, kind="CSV", status=Status.SCANNED)
    item.finding_counts = {"my_number": 150, "email": 1500}
    item.conf_counts = {"my_number:high": 150, "email:medium": 1500}
    item.distinct_counts = {"my_number:high": 150, "email:medium": 1200}
    store.save_result(job, FileResult(top_path=top, items=[item], distinct_counts=dict(item.distinct_counts)))
    row = store.query("SELECT flags, people_n, people_type, distinct_counts FROM files WHERE top_path=?", (top,))[0]
    assert reportflags.decode(row["flags"]) == ["my_number", "many_my_number", "many_people"]
    assert (row["people_n"], row["people_type"]) == (1200, "email")
    assert json.loads(row["distinct_counts"]) == item.distinct_counts
    s = store.summary(job)
    assert s["flags"]["my_number"] == 1 and s["flags"]["many_people"] == 1 and s["flags"]["sensitive"] == 0
    stored = store.query("SELECT distinct_counts FROM items WHERE top_path=?", (top,))[0][0]
    assert json.loads(stored) == item.distinct_counts


def test_old_db_gets_flags_without_counts(tmp_path):
    """古い版の DB（列が無い）を開くと、件数から決まる印だけが付き、異なる値は「数えていない」（NULL）のまま。"""
    path = tmp_path / "old.sqlite3"
    con = sqlite3.connect(path)
    con.executescript("""
        CREATE TABLE jobs (id TEXT PRIMARY KEY, created REAL, name TEXT, roots TEXT, options TEXT, state TEXT,
            started REAL, finished REAL, enum_done INTEGER DEFAULT 0, message TEXT DEFAULT '');
        CREATE TABLE files (job_id TEXT, top_path TEXT, state TEXT, worst_status TEXT, risk TEXT, size INTEGER,
            mtime REAL, elapsed REAL, n_findings INTEGER DEFAULT 0, PRIMARY KEY (job_id, top_path));
        CREATE TABLE items (job_id TEXT, top_path TEXT, seq INTEGER, path TEXT, kind TEXT, status TEXT, reason TEXT,
            method TEXT, size INTEGER, mtime REAL, depth INTEGER, password_kind TEXT, counts TEXT, notes TEXT,
            conf_counts TEXT DEFAULT '{}', PRIMARY KEY (job_id, top_path, seq));
        INSERT INTO jobs VALUES ('j1', 1, 'x', '["C:\\\\x"]', '{}', 'done', 1, 2, 1, '');
        INSERT INTO files VALUES ('j1', 'C:\\x\\a.csv', 'done', 'scanned', '高', 1, 1, 0, 3);
        INSERT INTO files VALUES ('j1', 'C:\\x\\b.csv', 'done', 'scanned', 'なし', 1, 1, 0, 0);
        INSERT INTO items VALUES ('j1', 'C:\\x\\a.csv', 0, 'C:\\x\\a.csv', 'CSV', 'scanned', '', '', 1, 1, 0, '',
            '{"sensitive": 2, "bank_account": 1}', '[]', '{"sensitive:medium": 2, "bank_account:high": 1}');
        INSERT INTO items VALUES ('j1', 'C:\\x\\b.csv', 0, 'C:\\x\\b.csv', 'CSV', 'scanned', '', '', 1, 1, 0, '',
            '{}', '[]', '{}');
        """)
    con.commit()
    con.close()
    store = Store(str(path))
    rows = {r["top_path"]: r for r in store.query("SELECT top_path, flags, people_n, distinct_counts FROM files")}
    assert reportflags.decode(rows["C:\\x\\a.csv"]["flags"]) == ["sensitive", "financial"]
    assert rows["C:\\x\\a.csv"]["people_n"] is None and rows["C:\\x\\a.csv"]["distinct_counts"] is None
    assert rows["C:\\x\\b.csv"]["flags"] == ""
    assert store.summary("j1")["flags"]["financial"] == 1


def _save(store, job, top, mine: int, people: int):
    store.add_files(job, [(top, 1, 1.0)])
    item = ItemResult(path=top, kind="CSV", status=Status.SCANNED)
    item.finding_counts = {"my_number": mine, "email": people}
    item.conf_counts = {"my_number:high": mine, "email:medium": people}
    item.distinct_counts = {"my_number:high": mine, "email:medium": people}
    store.save_result(job, FileResult(top_path=top, items=[item], distinct_counts=dict(item.distinct_counts)))


def test_job_totals_are_shown_in_verdict_excel_and_marked_over_limits(tmp_path):
    """ファイルごとでは 100・1,000 を超えなくても、合計が超えたら判定に目安として出す（ファイルごとの印は付かない）。"""
    from openpyxl import load_workbook

    from pii_checker.report import write_xlsx
    from pii_checker.verdict import decide

    store = Store(str(tmp_path / "db.sqlite3"))
    job = store.create_job([str(tmp_path)], ScanOptions())
    for i in range(3):
        _save(store, job, str(tmp_path / f"{i}.csv"), mine=40, people=400)
    store.set_job(job, state="done")
    s = store.summary(job)
    assert s["distinct_totals"] == {"my_number": 120, "people": 1200, "uncounted_files": 0}
    assert s["flags"]["many_my_number"] == 0 and s["flags"]["many_people"] == 0  # ファイルごとの印は付かない
    detail = decide(s, "done").detail
    joined = "\n".join(detail)
    assert (
        "異なるマイナンバーの数の合計 120（ファイルごとの数の合計。同じ人が複数のファイルにいれば重ねて数えるので上限の目安）"
        in joined
    )
    assert "100 を超える" in joined
    assert (
        "人数の目安の合計 1,200（ファイルごとの人数の目安を足したもの。重ねて数えることも、少なく数えることもある目安）"
        in joined
    )
    assert "1,000 を超える" in joined and "上限" not in [x for x in detail if x.startswith("人数の目安の合計")][0]
    assert reportflags.PER_FILE_NOTE in detail
    # Excel の判定シートにも同じ行と、概要の合計の行が出る
    out = tmp_path / "r.xlsx"
    write_xlsx(store, job, str(out))
    cells = [c for row in load_workbook(out, read_only=True)["判定"].iter_rows(values_only=True) for c in row if c]
    assert any("異なるマイナンバーの数の合計 120" in str(c) for c in cells)
    assert 1200 in cells

    # 合計が上限以下なら「超える」とは出さない
    store2 = Store(str(tmp_path / "db2.sqlite3"))
    job2 = store2.create_job([str(tmp_path)], ScanOptions())
    _save(store2, job2, str(tmp_path / "a.csv"), mine=10, people=20)
    lines = reportflags.total_lines(store2.summary(job2))
    assert lines[0].startswith("異なるマイナンバーの数の合計 10（") and not any("超える" in x for x in lines)


def test_people_estimate_ties_follow_fixed_priority():
    """人数の目安で数が同じときは、氏名＞マイナンバー＞メール＞電話＞… の順で選ぶ（書いた順・辞書の順に左右されない）。"""
    one = {"phone:medium": 1, "address:medium": 1, "email:medium": 1, "name:medium": 1, "my_number:high": 1}
    assert reportflags.people_estimate(one) == (1, "name")
    assert reportflags.people_estimate({"address:medium": 1, "phone:medium": 1}) == (1, "phone")
    assert reportflags.people_estimate({"phone:medium": 1, "email:medium": 1}) == (1, "email")
    assert reportflags.people_estimate({"email:medium": 2, "my_number:high": 2, "phone:medium": 2}) == (2, "my_number")
    # 数が多ければ後の種類でも選ぶ
    assert reportflags.people_estimate({"name:medium": 1, "phone:medium": 2}) == (2, "phone")
    assert reportflags.PEOPLE_TYPES[:4] == ("name", "my_number", "email", "phone")


def test_per_file_people_flag_wording():
    assert reportflags.FLAG_LABELS["many_people"] == "人数の目安が 1,000 を超える"


def test_financial_short_label_does_not_read_as_happened():
    assert reportflags.FLAG_SHORT["financial"] == "被害のおそれ（カード・口座）"
