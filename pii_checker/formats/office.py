"""Office 系の処理。

- OOXML（xlsx/xlsm/docx/pptx）と OpenDocument は zip の中の部品を**全部**読む。
  セルや本文だけでなく、コメント・図形・グラフ・ピボットキャッシュ（元データの写し）・
  文書プロパティ・変更履歴の削除文字列・埋め込みファイル・画像（OCR）まで見る。
- Office97-2003（xls/doc/ppt）・Outlook の msg・一太郎などの OLE 複合文書は、
  読める形式は構造を読み、加えて全ストリームを生データ走査する（削除済みの残留も拾う）。
- 暗号化は既定パスワード（VelvetSweatshop）と受け取ったパスワード候補で解除を試す。
"""

from __future__ import annotations

import io
import re
import zipfile
import zlib
import xml.etree.ElementTree as ET

from openpyxl.utils import get_column_letter

from ..engine import (
    READ_ALL_LIMIT,
    Blob,
    Collector,
    ItemResult,
    NeedPassword,
    PartialScan,
    ScanContext,
    Uncheckable,
)
from ..models import Status
from ..ooxml_text import (
    DOCX_EXTRA_RE,
    DOCX_MAIN,
    docx_structure,
    pptx_paragraphs,
    pptx_para_text,
    pptx_sheet,
    pptx_slide_parts,
)
from ..textdecode import iter_strings, iter_strings_chunked
from .misc import attribute_texts, markup_to_text, rtf_to_text

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SKIP_PARTS = re.compile(r"(^|/)(printerSettings|fonts?)/|\.(odttf|fntdata|ttf|otf)$", re.I)
PPT_OBJECT_LIMIT = 512 * 1024**2  # PPT の埋め込みオブジェクトを展開してよい大きさ
DEFAULT_OFFICE_PASSWORDS = ["VelvetSweatshop"]  # Excel の「読み取りパスワードなし」暗号化の既定値


def _col_letter(c: int) -> str:
    return get_column_letter(c + 1)


# ------------------------------------------------------------------ zip 部品の総なめ


def sweep_zip_parts(
    ctx: ScanContext,
    z: zipfile.ZipFile,
    blob: Blob,
    col: Collector,
    depth: int,
    skip: set[str] | None = None,
    raw_parts=lambda name: False,
) -> int:
    """zip の部品をすべて検査する。XML は本文と属性値、それ以外は内側のファイルとして再帰。"""
    skip = skip or set()
    n = 0
    for info in z.infolist():
        name = info.filename
        if info.is_dir() or name in skip:
            continue
        if SKIP_PARTS.search(name):
            continue
        n += 1
        lower = name.lower()
        if info.file_size > READ_ALL_LIMIT:
            # 巨大な部品（数 GB のシート XML・部品に仕込んだ圧縮爆弾）はメモリへ丸ごと読まず、文字列を少しずつ拾う
            ctx.account(info.file_size)
            with z.open(info) as f:
                for s in iter_strings_chunked(f.read):
                    col.add_text(s, name, method="大きな部品の文字列走査", raw=raw_parts(name))
            col.item.status = Status.PARTIAL
            col.item.reason = f"部品 {name}（{info.file_size:,} バイト）が大きすぎるため、その部品は文字として読める部分だけを検査した"
            continue
        data = z.read(info)
        if lower.endswith((".xml", ".rels", ".vml")) or data[:5] == b"<?xml":
            text_src = data.decode("utf-8", errors="replace")
            raw = raw_parts(name)
            col.add_lines(
                markup_to_text(text_src), location_prefix=f"{name} ", method="XML部品" if not raw else "XML部品"
            )
            for v in attribute_texts(text_src):
                col.add_text(v, f"{name}（属性値）", method="属性値", raw=raw)
            continue
        if (
            lower.endswith(".bin")
            and data[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
            and lower.startswith(("xl/", "word/", "ppt/"))
            and "embeddings" not in lower
        ):
            # xlsb のシート・共有文字列など、構造の分からない部品は生データ走査
            for s in iter_strings(data):
                col.add_text(s, name, method="生データ走査", raw=raw_parts(name))
            continue
        ctx.scan_nested(f"{blob.display}!{name}", data=data, depth=depth + 1, name_hint=name)
    return n


def _odf_encrypted(z: zipfile.ZipFile) -> bool:
    try:
        manifest = z.read("META-INF/manifest.xml")
    except KeyError:
        return False
    return b"encryption-data" in manifest


def _open_zip(blob: Blob):
    f = blob.open()
    try:
        return f, zipfile.ZipFile(f)
    except Exception:
        f.close()
        raise


# ------------------------------------------------------------------ xlsx


def handle_xlsx(ctx: ScanContext, blob: Blob, item: ItemResult, col: Collector, depth: int) -> None:
    import openpyxl

    sheets_done = False
    sheet_notes = []
    try:
        with blob.open() as f:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True, keep_links=False)
            try:
                for ws in wb.worksheets:
                    state = getattr(ws, "sheet_state", "visible")
                    if state != "visible":
                        sheet_notes.append(f"{ws.title}（{'非表示' if state == 'hidden' else 'ごく非表示'}）")
                    if hasattr(ws, "reset_dimensions"):
                        ws.reset_dimensions()  # 記録された範囲を信用せず、全セルを読む
                    title = ws.title
                    col.pos_layout, col.pos_sheet = "xlsx", title  # 対処のために、シート・行・列で位置を集める
                    col.add_table(
                        ws.iter_rows(min_row=1, min_col=1, values_only=True),
                        lambda r, c, title=title: f"{title}!{_col_letter(c)}{r}",
                        method="セル",
                    )
                sheets_done = True
            finally:
                wb.close()
                col.pos_sheet = ""
    except Exception as e:
        col.pos_layout = ""  # 途中までしか読めていない位置は、対処に使わない
        col.positions.clear()
        item.notes.append(f"openpyxl で読めなかったため XML 部品として読んだ（{type(e).__name__}: {e}）")
    if sheet_notes:
        item.notes.append("非表示のシートも検査した: " + "、".join(sheet_notes))
    f, z = _open_zip(blob)
    with f, z:
        n = sweep_zip_parts(
            ctx,
            z,
            blob,
            col,
            depth,
            raw_parts=lambda name: sheets_done
            and (name.startswith("xl/worksheets/sheet") or name == "xl/sharedStrings.xml"),
        )
    item.method = ("openpyxl で全シートのセル（非表示含む）＋" if sheets_done else "") + f"zip 内の全部品 {n} 個"
    if not sheets_done:
        raise PartialScan("セルを表として読めず、XML の文字だけを検査した（列見出しの意味を使えていない）")


# ------------------------------------------------------------------ xlsb


def handle_xlsb(ctx: ScanContext, blob: Blob, item: ItemResult, col: Collector, depth: int) -> None:
    sheets_done = False
    try:
        from pyxlsb import open_workbook

        path = ctx.blob_as_path(blob, ".xlsb")
        with open_workbook(path) as wb:
            for name in wb.sheets:
                with wb.get_sheet(name) as sh:

                    def rows(sh=sh):
                        for row in sh.rows():
                            yield [c.v for c in row]

                    col.add_table(rows(), lambda r, c, name=name: f"{name}!{_col_letter(c)}{r}", method="セル")
        sheets_done = True
    except Exception as e:
        item.notes.append(f"pyxlsb で読めなかったため生データ走査（{type(e).__name__}: {e}）")
    f, z = _open_zip(blob)
    with f, z:
        n = sweep_zip_parts(ctx, z, blob, col, depth, raw_parts=lambda name: sheets_done and name.endswith(".bin"))
    item.method = ("pyxlsb で全シート＋" if sheets_done else "") + f"zip 内の全部品 {n} 個"
    if not sheets_done:
        raise Uncheckable("xlsb のセルを構造として読めず、生データ走査のみ（数値のセルは検査できていない）")


# ------------------------------------------------------------------ docx


def _para_text(p: ET.Element) -> str:
    """段落の文字（入れ子の段落＝テキストボックスの中身は除く。そちらは別に数える）。"""
    out: list[str] = []

    def walk(el: ET.Element, top: bool) -> None:
        if not top and el.tag == f"{W}p":
            return
        if el.tag in (f"{W}t", f"{W}delText", f"{W}instrText") and el.text:
            out.append(el.text)
        elif el.tag in (f"{W}tab",):
            out.append("\t")
        elif el.tag in (f"{W}br", f"{W}cr"):
            out.append("\n")
        for ch in el:
            walk(ch, False)

    walk(p, True)
    return "".join(out)


def handle_docx(ctx: ScanContext, blob: Blob, item: ItemResult, col: Collector, depth: int) -> None:
    f, z = _open_zip(blob)
    with f, z:
        structured = False
        collecting = col.collect_positions
        if collecting:
            col.pos_layout = (
                "docx"  # 対処のために、部品・段落・表の行と列で位置を集める（remedy_office.py が書き換えに使う）
            )
        try:
            root = ET.fromstring(z.read(DOCX_MAIN))
            info = docx_structure(root) if collecting else {}
            n_par = 0
            for idx, p in enumerate(root.iter(f"{W}p"), start=1):
                text = _para_text(p)
                if text.strip():
                    n_par += 1
                    col.pos_sheet, col.cell_pos = DOCX_MAIN, (idx, info[id(p)][1] if collecting else 0)
                    try:
                        col.add_text(text, f"本文 段落{n_par}", method="本文")
                    finally:
                        col.pos_sheet, col.cell_pos = "", None
            for t_idx, tbl in enumerate(root.iter(f"{W}tbl"), start=1):
                rows = []
                for tr in tbl.findall(f"{W}tr"):
                    rows.append(["".join(_para_text(p) for p in tc.iter(f"{W}p")) for tc in tr.findall(f"{W}tc")])
                # 表は列見出しの意味を使ってもう一度読む（本文で数えた値は二重に数えない）
                col.pos_sheet = DOCX_MAIN
                col.pos_map = _cell_position_mapper(tbl, info) if collecting else None
                try:
                    col.add_table(
                        rows,
                        lambda r, c, t_idx=t_idx: f"表{t_idx} {r}行{c + 1}列",
                        method="表（列見出し）",
                        skip_seen=True,
                    )
                finally:
                    col.pos_sheet, col.cell_pos, col.pos_map = "", None, None
            structured = True
        except Exception as e:
            item.notes.append(f"本文の構造を読めなかったため XML として読んだ（{type(e).__name__}: {e}）")
        n = sweep_zip_parts(ctx, z, blob, col, depth, raw_parts=lambda name: structured and name == DOCX_MAIN)
        if collecting and structured:
            _collect_docx_extra_positions(z, col, item)
    item.method = (
        f"本文の段落・表（見出し付き）＋zip 内の全部品 {n} 個（ヘッダー・フッター・コメント・脚注・変更履歴を含む）"
    )
    if not structured:
        raise PartialScan(
            "本文の段落・表を構造として読めず、XML の文字だけを検査した（表の列見出しの意味を使えていない）"
        )


def _cell_position_mapper(tbl, info: dict):
    """表のセル（行・列。文字は、セルの中の段落の文字をつないだもの）の文字の範囲を、段落ごとの範囲にする関数を返す。

    対処は段落を単位に書き換えるので、表の列見出しを使って読んだ検出も、本文の段落の位置と同じ形（部品・段落の通し番号・段落の中の範囲）
    で残す（セルの中の検出が、本文の段落として読んだ検出と同じ鍵になり、書き換えで二重にならない）。
    """
    import unicodedata

    from ..remedy_write import _orig_bounds

    rows = tbl.findall(f"{W}tr")

    def mapper(row: int, c: int, start: int, end: int) -> list[tuple]:
        cells = rows[row - 1].findall(f"{W}tc")
        if c >= len(cells):
            return []
        paras = list(cells[c].iter(f"{W}p"))
        texts = [_para_text(p) for p in paras]
        cell = "".join(texts)
        norm = unicodedata.normalize("NFKC", cell)
        bounds = _orig_bounds(cell, norm)
        o_start, o_end = (0, len(cell)) if bounds is None else bounds(start, min(end, len(norm)))
        out = []
        offset = 0
        for p, text in zip(paras, texts):
            a, b = max(o_start, offset), min(o_end, offset + len(text))
            if a < b:
                n0 = len(unicodedata.normalize("NFKC", text[: a - offset]))
                n1 = len(unicodedata.normalize("NFKC", text[: b - offset]))
                idx, code = info[id(p)]
                out.append((DOCX_MAIN, idx, code, n0, n1))
            offset += len(text)
        return out

    return mapper


def _collect_docx_extra_positions(z: zipfile.ZipFile, col: Collector, item: ItemResult) -> None:
    """対処のための位置集め（検査の結果は変えない）: ヘッダー・フッター・脚注・コメントの段落の位置を集める。

    これらの部品は検査では XML の文字として読むだけなので、段落の通し番号で位置を取り直す。検査で数えた値は
    二重に数えない（skip_seen）。
    """
    for name in z.namelist():
        if not DOCX_EXTRA_RE.match(name):
            continue
        try:
            root = ET.fromstring(z.read(name))
        except Exception as e:
            item.notes.append(f"{name} の段落を読めなかったため、対処では変わらない（{type(e).__name__}）")
            continue
        info = docx_structure(root)
        for idx, p in enumerate(root.iter(f"{W}p"), start=1):
            text = _para_text(p)
            if not text.strip():
                continue
            col.pos_sheet, col.cell_pos = name, (idx, info[id(p)][1])
            try:
                col.add_text(text, f"{name} 段落{idx}", method="XML部品", skip_seen=True)
            finally:
                col.pos_sheet, col.cell_pos = "", None


# ------------------------------------------------------------------ pptx / ODF / その他 OOXML


ODF_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
ODF_OFFICE = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
ODF_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def _scan_ods_tables(z: zipfile.ZipFile, col: Collector) -> int:
    """ods の content.xml を表として読む（列見出しの意味を使う）。戻り値は表の数。"""
    root = ET.fromstring(z.read("content.xml"))
    n = 0
    for table in root.iter(f"{ODF_TABLE}table"):
        n += 1
        name = table.get(f"{ODF_TABLE}name") or f"表{n}"
        rows = []
        for tr in table.iter(f"{ODF_TABLE}table-row"):
            repeat_r = min(int(tr.get(f"{ODF_TABLE}number-rows-repeated", "1")), 1000)
            cells = []
            for tc in tr:
                if tc.tag not in (f"{ODF_TABLE}table-cell", f"{ODF_TABLE}covered-table-cell"):
                    continue
                repeat_c = min(int(tc.get(f"{ODF_TABLE}number-columns-repeated", "1")), 1000)
                text = "\n".join("".join(p.itertext()) for p in tc.iter(f"{ODF_TEXT}p"))
                value = tc.get(f"{ODF_OFFICE}date-value") or text or tc.get(f"{ODF_OFFICE}value") or ""
                cells.extend([value] * repeat_c)
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.extend([cells] * (repeat_r if cells else 1))
            elif rows:
                rows.append([])
        col.add_table(rows, lambda r, c, name=name: f"{name}!{_col_letter(c)}{r}", method="セル")
    return n


def handle_ooxml_generic(ctx: ScanContext, blob: Blob, item: ItemResult, col: Collector, depth: int) -> None:
    f, z = _open_zip(blob)
    with f, z:
        if _odf_encrypted(z):
            raise Uncheckable(
                "パスワードで暗号化された OpenDocument は未対応のため検査できない（要目視・パスワードを外して再検査）"
            )
        tables = 0
        try:
            mimetype = z.read("mimetype").decode("ascii", "replace") if "mimetype" in z.namelist() else ""
            if "spreadsheet" in mimetype:
                tables = _scan_ods_tables(z, col)
        except Exception as e:
            item.notes.append(f"表として読めなかったため XML として読んだ（{type(e).__name__}: {e}）")
        n = sweep_zip_parts(
            ctx, z, blob, col, depth, skip={"mimetype"}, raw_parts=lambda name: tables > 0 and name == "content.xml"
        )
    item.method = (
        f"表 {tables} 枚のセル＋" if tables else ""
    ) + f"zip 内の全部品 {n} 個（本文・ノート・コメント・埋め込み・画像）"


def handle_pptx(ctx: ScanContext, blob: Blob, item: ItemResult, col: Collector, depth: int) -> None:
    """pptx。検査は handle_ooxml_generic と同じ（zip の全部品）。対処のための位置集めのときだけ、
    スライドとノートの段落の位置を足す（検査の結果は変えない）。"""
    handle_ooxml_generic(ctx, blob, item, col, depth)
    if not col.collect_positions:
        return
    try:
        f, z = _open_zip(blob)
        with f, z:
            col.pos_layout = "pptx"
            for no, part, notes in pptx_slide_parts(z.read):
                for sheet, name in ((pptx_sheet(no, part, False), part), (pptx_sheet(no, notes, True), notes)):
                    if not name:
                        continue
                    root = ET.fromstring(z.read(name))
                    for para_no, shape, p in pptx_paragraphs(root):
                        text = pptx_para_text(p)
                        if not text.strip():
                            continue
                        col.pos_sheet, col.cell_pos = sheet, (para_no, shape)
                        try:
                            col.add_text(text, f"{name} 段落{para_no}", method="スライドの文字", skip_seen=True)
                        finally:
                            col.pos_sheet, col.cell_pos = "", None
    except Exception as e:
        col.pos_layout = ""  # 途中までしか読めていない位置は、対処に使わない
        col.positions.clear()
        item.notes.append(f"スライドの段落の位置を集められなかったため、対処できない（{type(e).__name__}: {e}）")


# ------------------------------------------------------------------ 暗号化


def try_decrypt_office(ctx: ScanContext, data: bytes) -> tuple[bytes | None, bool]:
    """Office の暗号化を解除する。戻り値は（復号した中身 or None, 暗号化されていたか）。

    暗号化されていて開けなければ NeedPassword を投げる。
    """
    import msoffcrypto

    try:
        probe = msoffcrypto.OfficeFile(io.BytesIO(data))
        encrypted = probe.is_encrypted()
    except Exception:
        return None, False
    if not encrypted:
        return None, False
    for pw in DEFAULT_OFFICE_PASSWORDS + ctx.password_candidates():
        try:
            of = msoffcrypto.OfficeFile(io.BytesIO(data))
            try:
                of.load_key(password=pw, verify_password=True)
            except TypeError:
                of.load_key(password=pw)
            out = io.BytesIO()
            of.decrypt(out)
            result = out.getvalue()
            if result:
                return result, True
        except Exception:
            continue
    raise NeedPassword("office", _pw_reason(ctx, "Office ファイルがパスワードで暗号化されている"))


def _pw_reason(ctx: ScanContext, what: str) -> str:
    n = len(ctx.password_candidates())
    return (
        f"{what}（受け取ったパスワード {n} 個と既定値を試したが開けなかった）"
        if n
        else f"{what}（パスワードを入力すると再検査できる）"
    )


# ------------------------------------------------------------------ OLE 複合文書


def handle_ole(ctx: ScanContext, blob: Blob, item: ItemResult, col: Collector, depth: int) -> None:
    import olefile

    data = blob.read_all()
    try:
        ole = olefile.OleFileIO(data)
    except Exception as e:
        raise Uncheckable(f"OLE 複合文書として開けなかった（壊れている可能性）: {e}")
    with ole:
        streams = ["/".join(e) for e in ole.listdir(streams=True, storages=False)]
        top = {s.split("/")[0] for s in streams}
        ext = blob.ext

        # 1) 暗号化（OOXML の暗号化、xls/doc/ppt の RC4 暗号化）
        if "EncryptedPackage" in top or any(
            s in top for s in ("Workbook", "Book", "WordDocument", "PowerPoint Document")
        ):
            if "EncryptedPackage" in top:
                item.kind = "暗号化された Office 文書（docx/xlsx/pptx）"
            decrypted, was_encrypted = try_decrypt_office(ctx, data)
            if not was_encrypted and "EncryptedPackage" in top:
                raise Uncheckable("暗号化された Office 文書だが、暗号化の方式を判別できず復号できなかった")
            if was_encrypted:
                item.method = "暗号化を解除して中身を検査"
                item.notes.append("暗号化されていたため、パスワード（または既定値）で復号した")
                ctx.scan_nested(
                    f"{blob.display}!（復号した中身）",
                    data=decrypted,
                    depth=depth + 1,
                    name_hint="decrypted" + (ext or ".bin"),
                )
                return

        # IRM（権限管理）で保護された文書は、本文が DataSpaces/DRMContent に暗号化されている。
        # WordDocument には「アクセスは制限されています」という表紙しか無いので、読めたことにしない
        # 最上位のストリームだけで見る（IRM 文書を埋め込んだメール等で、全体を検査不能にしない）
        if any(t.startswith("\x06DataSpaces") or "DRMContent" in t for t in top):
            raise Uncheckable(
                "IRM（Information Rights Management）で保護された文書のため中身を読めない"
                "（作成者に保護の解除を依頼するか、権限のあるアカウントで開いて保護なしで保存してから再検査）"
            )

        # 2) 種類ごとの構造の読み取り
        structured = ""
        if "Workbook" in top or "Book" in top:
            item.kind = "Excel 97-2003（xls）" if "Workbook" in top else "Excel 5.0/95（xls）"
            stream = "Workbook" if "Workbook" in top else "Book"
            try:
                _scan_xls_bytes(ole.openstream(stream).read(), col, "", from_stream=True, whole=data)
                structured = "xlrd で全シートのセル"
            except Exception as e:
                item.notes.append(f"xlrd でセルを読めなかった（{type(e).__name__}: {e}）")
        elif "WordDocument" in top:
            item.kind = "Word 97-2003（doc）"
            try:
                text = _doc_text(ole)
                col.add_lines(text, location_prefix="本文 ", method="本文")
                structured = "本文（piece table。ヘッダー・脚注・テキストボックスを含む）"
            except Exception as e:
                item.kind = "Word（doc）"
                item.notes.append(f"本文を構造として読めなかった（{type(e).__name__}: {e}）")
        elif "PowerPoint Document" in top:
            item.kind = "PowerPoint 97-2003（ppt）"
            try:
                texts, objects = _ppt_parse(ole.openstream("PowerPoint Document").read())
                for i, t in enumerate(texts, start=1):
                    col.add_text(t, f"スライドの文字 {i}", method="スライドの文字")
                for i, obj in enumerate(objects, start=1):
                    ctx.scan_nested(
                        f"{blob.display}!埋め込みオブジェクト{i}", data=obj, depth=depth + 1, name_hint="embedded.bin"
                    )
                structured = f"スライド・ノートの文字 {len(texts)} 件・埋め込み {len(objects)} 件"
            except Exception as e:
                item.notes.append(f"スライドを構造として読めなかった（{type(e).__name__}: {e}）")
        elif any(s.startswith("__substg1.0_") for s in top) or "__properties_version1.0" in top:
            item.kind = "Outlook メール（msg）"
            try:
                _scan_msg(ctx, blob, col, depth)
                structured = "件名・差出人・宛先・本文・添付"
            except Exception as e:
                item.notes.append(f"msg を構造として読めなかった（{type(e).__name__}: {e}）")
        elif ext in (".jtd", ".jtt", ".jtdc", ".jttc", ".jaw", ".jbw"):
            item.kind = "一太郎"
            structured = "ストリームの文字列抽出"
        elif ext in (".msi", ".msp", ".msm"):
            # インストーラーは中に CAB でファイルを持つ（任意の文書を入れられる）ので、対象外にせず 7-Zip で中身を展開する
            from .archive import handle_with_7z_exe

            item.kind = "Windows インストーラー（MSI）"
            return handle_with_7z_exe(ctx, blob, item, col, depth)
        elif ext == ".db" and "Catalog" in top:
            item.kind = "Thumbs.db（縮小画像の控え）"
            structured = "縮小画像の取り出し"

        # 3) 全ストリームの生データ走査と、埋め込みファイルの取り出し
        raw = bool(structured) and structured.startswith(("xlrd", "件名", "本文", "スライド"))
        n_embedded = 0
        for s in streams:
            try:
                sdata = ole.openstream(s).read()
            except Exception as e:
                item.notes.append(f"ストリーム {s} を読めなかった: {e}")
                continue
            if not sdata:
                continue
            carved = list(_carve_embedded(s, sdata))
            if s.rsplit("/", 1)[-1] in ("Pictures", "Data") and not carved:
                carved = list(_carve_all_images(s, sdata))
            for emb_name, emb in carved:
                n_embedded += 1
                ctx.scan_nested(f"{blob.display}!{emb_name}", data=emb, depth=depth + 1, name_hint=emb_name)
            # 埋め込まれた Excel（他文書の中の表）は、ストリームのまま xlrd で読める
            if "/" in s and s.rsplit("/", 1)[-1] in ("Workbook", "Book"):
                try:
                    _scan_xls_bytes(sdata, col, f"埋め込み表 {s} ", from_stream=True)
                except Exception:
                    pass
            for text in iter_strings(sdata):
                col.add_text(text, f"ストリーム {s}", method="生データ走査", raw=raw)
        item.method = (
            (structured + "＋" if structured else "")
            + f"全ストリーム {len(streams)} 本の生データ走査"
            + (f"・埋め込み {n_embedded} 件" if n_embedded else "")
        )
        if not structured or item.kind == "一太郎":
            from ..engine import PartialScan

            raise PartialScan(
                f"{item.kind if structured else 'この OLE 形式'}は構造を解釈していないため、文字として読める部分だけを検査した"
                "（圧縮された本文や数値は読めない）"
            )


def _u16(b: bytes, o: int) -> int:
    return int.from_bytes(b[o : o + 2], "little")


def _u32(b: bytes, o: int) -> int:
    return int.from_bytes(b[o : o + 4], "little")


def _doc_text(ole) -> str:
    """Word 97-2003 の本文を piece table（MS-DOC の Clx）から取り出す。"""
    wd = ole.openstream("WordDocument").read()
    if _u16(wd, 0) != 0xA5EC:
        raise ValueError("Word の見出し（FIB）ではない")
    nfib = _u16(wd, 2)
    if nfib < 101:
        raise ValueError(f"Word 6.0/95 以前の形式（nFib={nfib}）")
    flags = _u16(wd, 0x0A)
    table = ole.openstream("1Table" if flags & 0x0200 else "0Table").read()
    fc_clx, lcb_clx = _u32(wd, 0x01A2), _u32(wd, 0x01A6)
    clx = table[fc_clx : fc_clx + lcb_clx]
    pos = 0
    while pos < len(clx) and clx[pos] == 1:  # Prc（書式）は読み飛ばす
        pos += 3 + _u16(clx, pos + 1)
    if pos >= len(clx) or clx[pos] != 2:
        raise ValueError("piece table が見つからない")
    lcb = _u32(clx, pos + 1)
    plc = clx[pos + 5 : pos + 5 + lcb]
    n = (lcb - 4) // 12
    cps = [_u32(plc, 4 * i) for i in range(n + 1)]
    out = []
    for i in range(n):
        pcd = plc[4 * (n + 1) + 8 * i : 4 * (n + 1) + 8 * i + 8]
        fcv = _u32(pcd, 2)
        length = max(0, cps[i + 1] - cps[i])
        if fcv & 0x40000000:
            fc = (fcv & 0x3FFFFFFF) // 2
            out.append(wd[fc : fc + length].decode("cp1252", errors="replace"))
        else:
            fc = fcv & 0x3FFFFFFF
            out.append(wd[fc : fc + 2 * length].decode("utf-16-le", errors="replace"))
    text = "".join(out)
    if n > 0 and not text.strip() and cps[-1] > 0:
        raise ValueError("本文の文字を取り出せなかった（ストリームが欠けている可能性）")
    # 段落・セル・改ページ・フィールドの制御文字を、検査しやすい区切りに置き換える
    return text.translate(
        {0x0D: "\n", 0x07: "\t", 0x0B: "\n", 0x0C: "\n", 0x13: " ", 0x14: " ", 0x15: " ", 0x01: " ", 0x08: " "}
    )


def _ppt_parse(data: bytes) -> tuple[list[str], list[bytes]]:
    """PowerPoint 97-2003 のレコードを辿り、文字（TextChars/TextBytes/CString）と埋め込みオブジェクトを取り出す。"""
    texts: list[str] = []
    objects: list[bytes] = []
    stack = [(0, len(data))]
    while stack:
        start, end = stack.pop()
        pos = start
        while pos + 8 <= end:
            h = _u16(data, pos)
            rtype = _u16(data, pos + 2)
            length = _u32(data, pos + 4)
            body = pos + 8
            stop = min(body + length, end)
            ver, inst = h & 0x0F, h >> 4
            if ver == 0x0F:
                stack.append((body, stop))
            elif rtype in (0x0FA0, 0x0FBA):
                texts.append(data[body:stop].decode("utf-16-le", errors="replace"))
            elif rtype == 0x0FA8:
                texts.append(data[body:stop].decode("latin-1"))
            elif rtype == 0x1011:
                payload = data[body:stop]
                if inst == 1 and len(payload) > 4:
                    try:
                        # 展開後の大きさに上限を設ける（圧縮爆弾でメモリを食い潰さない）
                        d = zlib.decompressobj()
                        out = d.decompress(payload[4:], PPT_OBJECT_LIMIT)
                        if d.unconsumed_tail:
                            raise ValueError(
                                f"埋め込みオブジェクトが展開上限（{PPT_OBJECT_LIMIT // 1024**2}MB）を超える"
                            )
                        payload = out
                    except zlib.error:
                        pass
                if payload:
                    objects.append(payload)
            pos = body + length
    return texts, objects


def _carve_embedded(stream_name: str, data: bytes):
    """ストリームの中に丸ごと入っている別ファイル（埋め込み・添付）を取り出す。"""
    leaf = stream_name.rsplit("/", 1)[-1]
    if leaf == "\x01Ole10Native" and len(data) > 6:
        # [全体長4][種別2][表示名\0][元パス\0][予備4][一時パス長4][一時パス][データ長4][データ]
        try:
            pos = 6
            label_end = data.index(b"\x00", pos)
            label = data[pos:label_end].decode("cp932", errors="replace")
            pos = data.index(b"\x00", label_end + 1) + 1
            pos += 4
            tmp_len = int.from_bytes(data[pos : pos + 4], "little")
            pos += 4 + tmp_len
            size = int.from_bytes(data[pos : pos + 4], "little")
            pos += 4
            payload = data[pos : pos + size]
            if payload:
                yield f"{stream_name}（埋め込み {label}）", payload
                return
        except (ValueError, IndexError):
            pass
    magics = (b"PK\x03\x04", b"%PDF", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", b"\xff\xd8\xff", b"\x89PNG", b"{\\rtf")
    for m in magics:
        idx = data.find(m, 0, 64)
        if idx >= 0 and len(data) - idx > 64:
            if m == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" and idx == 0 and leaf not in ("Package", "CONTENTS"):
                continue
            yield f"{stream_name}（埋め込み）", data[idx:]
            return
    # Thumbs.db 等: 先頭近くに JPEG を持つストリーム
    idx = data.find(b"\xff\xd8\xff", 0, 128)
    if idx > 0:
        yield f"{stream_name}（縮小画像）", data[idx:]


def _carve_all_images(stream_name: str, data: bytes):
    """Word の Data・PowerPoint の Pictures ストリームに並ぶ画像を全部取り出す。"""
    import re as _re

    starts = sorted({m.start() for m in _re.finditer(rb"\xff\xd8\xff|\x89PNG\r\n\x1a\n", data)})
    for i, pos in enumerate(starts):
        end = starts[i + 1] if i + 1 < len(starts) else len(data)
        if end - pos > 256:
            yield f"{stream_name}（画像{i + 1}）", data[pos:end]


def _scan_xls_bytes(
    data: bytes, col: Collector, prefix: str, from_stream: bool = False, whole: bytes | None = None
) -> None:
    """xls（BIFF2〜8）のセルを xlrd で読む。日付は日付として扱う。"""
    import xlrd

    try:
        book = xlrd.open_workbook(
            file_contents=whole if whole is not None else data, on_demand=True, formatting_info=False
        )
    except Exception:
        if whole is not None:
            book = xlrd.open_workbook(file_contents=data, on_demand=True, formatting_info=False)
        else:
            raise
    try:
        for idx in range(book.nsheets):
            sh = book.sheet_by_index(idx)
            name = sh.name

            def rows(sh=sh):
                for r in range(sh.nrows):
                    out = []
                    for c in range(sh.ncols):
                        ctype = sh.cell_type(r, c)
                        v = sh.cell_value(r, c)
                        if ctype == xlrd.XL_CELL_DATE:
                            try:
                                v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                            except Exception:
                                pass
                        out.append(v)
                    yield out

            col.add_table(rows(), lambda r, c, name=name: f"{prefix}{name}!{_col_letter(c)}{r}", method="セル")
            book.unload_sheet(idx)
    finally:
        book.release_resources()


def handle_xls_legacy(ctx: ScanContext, blob: Blob, item: ItemResult, col: Collector, depth: int) -> None:
    data = blob.read_all()
    ok = False
    try:
        _scan_xls_bytes(data, col, "")
        ok = True
    except Exception as e:
        item.notes.append(f"xlrd でセルを読めなかった（{type(e).__name__}: {e}）")
    for text in iter_strings(data):
        col.add_text(text, "生データ", method="生データ走査", raw=ok)
    item.method = ("xlrd でセル＋" if ok else "") + "生データ走査"
    if not ok:
        raise Uncheckable("古い Excel 形式のセルを構造として読めず、生データ走査のみ（数値のセルは検査できていない）")


def _scan_msg(ctx: ScanContext, blob: Blob, col: Collector, depth: int) -> None:
    import extract_msg

    msg = extract_msg.openMsg(blob.read_all() if blob.data is not None else blob.path)
    try:
        _scan_msg_obj(ctx, msg, blob.display, col, depth)
    finally:
        try:
            msg.close()
        except Exception:
            pass


def _scan_msg_obj(ctx: ScanContext, msg, display: str, col: Collector, depth: int) -> None:
    for label, attr in (("件名", "subject"), ("差出人", "sender"), ("宛先", "to"), ("CC", "cc"), ("BCC", "bcc")):
        try:
            v = getattr(msg, attr, None)
        except Exception:
            v = None
        if v:
            col.add_text(str(v), label, method="ヘッダー")
    try:
        body = msg.body
    except Exception:
        body = None
    if body:
        col.add_lines(body, location_prefix="本文 ", method="本文")
    for attr in ("htmlBody", "rtfBody"):
        try:
            raw = getattr(msg, attr, None)
        except Exception:
            raw = None
        if raw:
            text = rtf_to_text(raw) if attr == "rtfBody" else markup_to_text(raw.decode("utf-8", errors="replace"))
            col.add_lines(
                text,
                location_prefix=f"{'HTML' if attr == 'htmlBody' else 'RTF'}本文 ",
                method="本文",
            )
    for i, att in enumerate(getattr(msg, "attachments", []) or [], start=1):
        name = getattr(att, "longFilename", None) or getattr(att, "shortFilename", None) or f"添付{i}"
        data = getattr(att, "data", None)
        if isinstance(data, (bytes, bytearray)):
            ctx.scan_nested(f"{display}!{name}", data=bytes(data), depth=depth + 1, name_hint=name)
        elif data is not None and hasattr(data, "subject"):
            _scan_msg_obj(ctx, data, f"{display}!{name}", col, depth + 1)
