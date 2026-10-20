"""対処の安全の評価（1周目）で出た指摘の再現の試験。

どれも「直す前は落ち、直したあと通る」ことを確かめるために、指摘の状況を関数の差し替えで起こす。
- 高1: 位置を集めたあと・確定の前にファイルが変わっても、古い位置のまま書き換えない（実ファイル・書庫）
- 高2: 「ファイルを削除」も、下見のあとにファイルが変わっていたら消さない
- 高3: ごみ箱へ送ったのにごみ箱が増えなかったら「完全に削除した」と記録・表示する／下見の時点でごみ箱が使えない場所を見分ける
- 中4: sqlite3 の .dump のように 1 行で書く CREATE TABLE でも、列の削除のあと読み込めるダンプになる
- 中5: Excel の表（ListObject）つきの xlsx で行・列を削除しても、表の範囲・列の定義が正しい
- 中7: 同じファイルは、別のジョブからも同時に対処できない
- 中8: 記録の更新が失敗しても、元の名前のファイルが残る
- 低: 書きかけの新ファイルを置かない・未完了の記録があるジョブは消さない・長いパスは始める前に断る・属性などの引き継ぎ・
  埋め込みの検査の一時フォルダ・削除後の記録の失敗の文面・立て直しで一時の新ファイルを消す
検体はすべて pytest の tmp_path に作る。本物のごみ箱へは送らない（関数を差し替える）。
"""

from __future__ import annotations

import ctypes
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from pii_checker import engine, remediate, remedy_common, remedy_meta, remedy_write  # noqa: E402
from pii_checker.models import ScanOptions  # noqa: E402
from pii_checker.remediate import RemedyService, inline_runner  # noqa: E402
from pii_checker.remedy_write import RemedyError  # noqa: E402
from pii_checker.store import Store  # noqa: E402
from test_remediate import (  # noqa: E402
    ALL_RANKS,
    REAL_SEND_TO_RECYCLE_BIN,
    FakeTrash,
    analyze,
    encode_csv,
    make_service,
    new_tmp_files,
    person_rows,
    run,
)
from test_remediate_archive import analyze_inner, zip_case  # noqa: E402,F401  （zip_case は固定具として使う）
from test_remediate_db import build_sqlite  # noqa: E402

REAL_HAS_RECYCLE_BIN = remediate.has_recycle_bin  # 差し替える前の本物（固定具が差し替える前に控える）
windows_only = pytest.mark.skipif(os.name != "nt", reason="Windows のごみ箱・属性の試験")


@pytest.fixture(autouse=True)
def trash(tmp_path, monkeypatch):
    """どの試験でも、ごみ箱へ送る関数を差し替える（本物のごみ箱へ送らない）。ごみ箱のある場所として扱う。"""
    t = FakeTrash(tmp_path / "_trash")
    monkeypatch.setattr(remediate, "send_to_recycle_bin", t)
    monkeypatch.setattr(remediate, "has_recycle_bin", lambda path: True)
    return t


def make_csv(tmp_path: Path, name: str = "a.csv") -> Path:
    p = tmp_path / "src" / name
    p.parent.mkdir(exist_ok=True)
    p.write_bytes(encode_csv(person_rows(), "utf-8", "\n"))
    return p


def start_and_wait(svc: RemedyService, job: str, p: Path, params: dict, sha: str | None) -> dict:
    started = svc.start_run(job, str(p), params, True, sha)
    svc._runs[started["run_id"]].thread.join(180)
    return svc.run_status(job, started["run_id"])


# ================================================================ 高1: 集めたあと・確定の前にファイルが変わる


def test_plain_file_changed_after_positions_were_collected_is_refused(tmp_path, monkeypatch, trash):
    """位置を集めた直後に先頭へ3行足されても、古い行番号のまま書き換えない（見出しが消え、個人情報の行が残った不具合）。"""
    p = make_csv(tmp_path)
    svc, store, job = make_service(tmp_path, [p])
    analyze(svc, job, p)
    params = {"method": "delete_rows", "confidences": ALL_RANKS}
    pv = svc.preview(job, str(p), params)
    real_collect = remediate.collect_task
    box: dict = {}

    def collect_then_edit(args):
        got = real_collect(args)
        p.write_bytes("追記1,a\n追記2,b\n追記3,c\n".encode("utf-8") + p.read_bytes())  # 集め終えた直後に先頭へ3行
        box["edited"] = p.read_bytes()
        return got

    monkeypatch.setattr(remediate, "collect_task", collect_then_edit)
    r = start_and_wait(svc, job, p, params, pv["expected_sha256"])
    assert r["state"] == "error" and "変わりました" in r["message"], r
    assert p.read_bytes() == box["edited"]  # 触っていない（編集された中身のまま）
    assert not new_tmp_files(p.parent) and trash.sent == [] and store.remediations(job) == []


def test_archive_changed_while_rewriting_is_refused(zip_case, monkeypatch, trash):
    """書庫の中のファイルの書き換えの途中で書庫が変わったら、作り直した書庫で上書きしない（下見の基準のハッシュで照合する）。"""
    c = zip_case
    archive = c["archive"]
    analyze_inner(c["svc"], c["job"], archive, c["item"])
    params = {"method": "delete_rows", "confidences": ALL_RANKS, "item_path": c["item"]}
    pv = c["svc"].preview(c["job"], str(archive), params)
    real_rewrite = remediate.rewrite_file
    box: dict = {}

    def rewrite_then_edit(*a, **k):
        out = real_rewrite(*a, **k)
        with open(archive, "ab") as f:
            f.write(b"\0" * 16)  # 書き換えの途中に、書庫の末尾へ書き足された
        box["edited"] = archive.read_bytes()
        return out

    monkeypatch.setattr(remediate, "rewrite_file", rewrite_then_edit)
    r = start_and_wait(c["svc"], c["job"], archive, params, pv["expected_sha256"])
    assert r["state"] == "error" and "変わりました" in r["message"], r
    assert archive.read_bytes() == box["edited"] and trash.sent == []
    assert not new_tmp_files(archive.parent)


# ================================================================ 高2: 削除でも下見のあとの変化を確かめる


def test_delete_file_refuses_when_file_was_replaced_after_preview(tmp_path, monkeypatch, trash):
    p = make_csv(tmp_path)
    monkeypatch.setattr(
        remediate, "has_recycle_bin", lambda path: False
    )  # ごみ箱の無い場所（差し替えると完全に消える）
    svc, store, job = make_service(tmp_path, [p])
    analyze(svc, job, p)
    pv = svc.preview(job, str(p), {"method": "delete_file"})
    p.write_bytes("差し替えた、大事なファイル\n".encode("utf-8"))
    r = start_and_wait(svc, job, p, {"method": "delete_file"}, pv["expected_sha256"])
    assert r["state"] == "error" and "変わりました" in r["message"], r
    assert p.read_text(encoding="utf-8") == "差し替えた、大事なファイル\n" and trash.sent == []
    assert store.remediations(job) == []


def test_delete_file_without_expected_hash_is_refused(tmp_path):
    p = make_csv(tmp_path)
    with pytest.raises(RemedyError, match="ハッシュ"):
        remediate.apply_task({"path": str(p), "method": "delete_file", "options": ScanOptions()})
    assert p.exists()


# ================================================================ 高3: 送ったのにごみ箱が増えない


def _info_v2(path: str) -> bytes:
    """Windows 10 以降の $I ファイル（版2）の中身。"""
    name = (path + "\0").encode("utf-16-le")
    return (
        (2).to_bytes(8, "little")
        + (1).to_bytes(8, "little")
        + (0).to_bytes(8, "little")
        + (len(name) // 2).to_bytes(4, "little")
        + name
    )


def _info_v1(path: str) -> bytes:
    """Vista〜8.1 の $I ファイル（版1。パスは 520 バイト固定）の中身。"""
    raw = path.encode("utf-16-le")
    return (1).to_bytes(8, "little") + (1).to_bytes(8, "little") + (0).to_bytes(8, "little") + raw.ljust(520, b"\0")


def test_read_recycle_info_reads_both_versions(tmp_path):
    target = "C:\\データ\\名簿.csv"
    v2, v1 = tmp_path / "$IAAA.csv", tmp_path / "$IBBB.csv"
    v2.write_bytes(_info_v2(target))
    v1.write_bytes(_info_v1(target))
    assert remediate._read_recycle_info(str(v2)) == target
    assert remediate._read_recycle_info(str(v1)) == target
    (tmp_path / "$Ibad").write_bytes(b"xx")
    assert remediate._read_recycle_info(str(tmp_path / "$Ibad")) is None


def test_find_in_recycle_bin_matches_only_new_entry_for_the_path(tmp_path):
    """そのドライブのごみ箱の、送ったあとにできた $I の元のパスが一致したときだけ「入った」とみなす。"""
    import time as _time

    root = str(tmp_path) + os.sep
    mine = tmp_path / "$Recycle.Bin" / "S-1-5-21-1"
    mine.mkdir(parents=True)
    target = str(tmp_path / "data" / "名簿.csv")
    since = _time.time() - 1
    assert remediate._find_in_recycle_bin(root, target, since) is False  # まだ入っていない
    (mine / "$IOTHER.csv").write_bytes(_info_v2(str(tmp_path / "別のファイル.csv")))
    assert remediate._find_in_recycle_bin(root, target, since) is False  # 別のファイルだけ
    old = mine / "$IOLD.csv"
    old.write_bytes(_info_v2(target))
    os.utime(old, (since - 100, since - 100))
    assert remediate._find_in_recycle_bin(root, target, since) is False  # 前に送った同じ名前のもの（古い）
    (mine / "$INEW.csv").write_bytes(_info_v2(target.upper()))  # 大文字小文字の違いはそろえて比べる
    assert remediate._find_in_recycle_bin(root, target, since) is True


def test_find_in_recycle_bin_without_bin_folder_is_unknown(tmp_path):
    assert remediate._find_in_recycle_bin(str(tmp_path) + os.sep, str(tmp_path / "a.txt"), 0) is None


@windows_only
def test_send_that_does_not_enter_the_bin_is_permanent(tmp_path, monkeypatch):
    p = tmp_path / "a.txt"
    p.write_text("x", encoding="utf-8")

    def shell_delete(abs_path):
        os.remove(abs_path)  # SHFileOperationW が 0 を返しつつ、実は完全に削除した状態
        return 0, False

    monkeypatch.setattr(remediate, "_shell_delete", shell_delete)
    monkeypatch.setattr(remediate, "_find_in_recycle_bin", lambda root, path, since: False)
    assert REAL_SEND_TO_RECYCLE_BIN(str(p)) == "permanent"
    assert not p.exists()


@windows_only
def test_send_that_enters_the_bin_is_recycle_bin(tmp_path, monkeypatch):
    p = tmp_path / "a.txt"
    p.write_text("x", encoding="utf-8")
    seen = {}
    monkeypatch.setattr(remediate, "_shell_delete", lambda abs_path: (os.remove(abs_path), (0, False))[1])

    def find(root, path, since):
        seen["path"] = path
        return True

    monkeypatch.setattr(remediate, "_find_in_recycle_bin", find)
    assert REAL_SEND_TO_RECYCLE_BIN(str(p)) == "recycle_bin"
    assert seen["path"] == os.path.abspath(str(p))


@windows_only
def test_send_is_permanent_when_the_bin_cannot_be_looked_at(tmp_path, monkeypatch):
    p = tmp_path / "a.txt"
    p.write_text("x", encoding="utf-8")
    monkeypatch.setattr(remediate, "_shell_delete", lambda abs_path: (os.remove(abs_path), (0, False))[1])
    monkeypatch.setattr(
        remediate, "_find_in_recycle_bin", lambda root, path, since: None
    )  # 確かめられないものは、完全に削除したと扱う
    assert REAL_SEND_TO_RECYCLE_BIN(str(p)) == "permanent"


@windows_only
@pytest.mark.parametrize(
    "has_bin_dir,policy,expected",
    [
        pytest.param(True, {"nuke": False, "max_mb": 1000}, True, id="使える"),
        pytest.param(False, {"nuke": False, "max_mb": 1000}, False, id="ごみ箱のフォルダが無い"),
        pytest.param(True, {"nuke": True, "max_mb": 1000}, False, id="NukeOnDelete"),
        pytest.param(True, {"nuke": False, "max_mb": 0}, False, id="最大の大きさが0"),
    ],
)
def test_has_recycle_bin_detects_unusable_bin(tmp_path, monkeypatch, has_bin_dir, policy, expected):
    import ctypes as ct

    if ct.windll.kernel32.GetDriveTypeW(remediate.drive_root(str(tmp_path))) != 3:
        pytest.skip("試験用の場所が固定ディスクではない")
    real_isdir = os.path.isdir
    monkeypatch.setattr(
        remediate.os.path,
        "isdir",
        lambda q: has_bin_dir if str(q).endswith(remediate.RECYCLE_DIR) else real_isdir(q),
    )
    monkeypatch.setattr(remediate, "_recycle_bin_policy", lambda root: policy)
    assert REAL_HAS_RECYCLE_BIN(str(tmp_path / "a.txt")) is expected


@windows_only
def test_recycle_bin_policy_reads_without_error():
    policy = remediate._recycle_bin_policy(remediate.drive_root(os.getcwd()))
    assert isinstance(policy["nuke"], bool) and (policy["max_mb"] is None or policy["max_mb"] >= 0)


def test_rewrite_whose_original_did_not_enter_the_bin_is_recorded_as_gone(tmp_path, monkeypatch, trash):
    """ごみ箱へ送ったのに入らなかったとき: 新しい中身は元の名前のまま巻き戻さず、記録・結果を「元は残らなかった」にする。"""
    p = make_csv(tmp_path)
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])

    def vanishes(path):
        trash(path)  # 元は消えた（ごみ箱の代わりの置き場へ移しているが、アプリには「入らなかった」と答える）
        return "permanent"

    monkeypatch.setattr(remediate, "send_to_recycle_bin", vanishes)
    r = run(svc, job, p, method="delete_rows", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    assert r["disposal"] == "replaced" and r["bin_failed"] is True and r["pre_name"] is None
    assert remediate.BIN_FAILED_TEXT in r["notes"]
    assert p.read_bytes() != original and not new_tmp_files(p.parent)  # 新しい中身が元の名前にある（巻き戻していない）
    rec = store.remediations(job)[0]
    assert (
        rec["state"] == "applied" and rec["changed"]["bin_failed"] is True and rec["changed"]["disposal"] == "replaced"
    )


def test_delete_file_that_did_not_enter_the_bin_is_recorded_as_permanent(tmp_path, monkeypatch, trash):
    p = make_csv(tmp_path)
    svc, store, job = make_service(tmp_path, [p])

    def vanishes(path):
        trash(path)
        return "permanent"

    monkeypatch.setattr(remediate, "send_to_recycle_bin", vanishes)
    r = run(svc, job, p, method="delete_file")
    assert r["state"] == "done" and r["disposal"] == "permanent", r
    assert store.remediations(job)[0]["changed"]["disposal"] == "permanent"


# ================================================================ 中4: sqlite3 の .dump と同じ 1 行の CREATE


def test_one_line_create_table_columns_are_dropped_and_dump_loads(tmp_path, trash):
    src = tmp_path / "base.sqlite"
    build_sqlite(src, n=40)
    con = sqlite3.connect(src)
    dump_text = "\n".join(con.iterdump()) + "\n"
    con.close()
    p = tmp_path / "src" / "people.sql"
    p.parent.mkdir()
    p.write_text(dump_text, encoding="utf-8", newline="\n")
    create_line = next(x for x in dump_text.splitlines() if x.startswith("CREATE TABLE people"))
    assert create_line.count(",") >= 4  # 検体の確認: CREATE は 1 行に全部の列がある（sqlite3 の .dump の形）
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_columns", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    result = p.read_text(encoding="utf-8")
    mem = sqlite3.connect(":memory:")
    mem.executescript(result)  # 読み込めること（CREATE と INSERT の列の数が合っている）
    names = [row[1] for row in mem.execute("PRAGMA table_info(people)")]
    assert names == ["id", "備考"], names
    assert mem.execute("SELECT COUNT(*) FROM people").fetchone()[0] == 40
    assert mem.execute("SELECT COUNT(*) FROM items").fetchone()[0] == 30  # ほかの表は元のまま
    mem.close()
    assert r["remaining"]["total"] == 0


def test_dropping_a_column_missing_from_create_is_refused(tmp_path, trash):
    """CREATE の列の並びを読み取れず、INSERT だけ列が減る形になるなら、断って元を残す。"""
    text = "CREATE TABLE people (id INTEGER, memo TEXT, email TEXT);\n" + "".join(
        f"INSERT INTO people VALUES({i},'m{i}','user{i}@example.com');\n" for i in range(1, 40)
    )
    # 列名の並びと食い違う CREATE（INSERT は3列だが、CREATE は別の表の定義にしか見えない形）にして、断ることを見る
    broken = text.replace(
        "CREATE TABLE people (id INTEGER, memo TEXT, email TEXT);", "CREATE TABLE people (LIKE other);"
    )
    p = tmp_path / "src" / "people.sql"
    p.parent.mkdir()
    p.write_text(broken, encoding="utf-8", newline="\n")
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_columns", confidences=ALL_RANKS)
    assert r["state"] == "error", r
    assert p.read_bytes() == original and not new_tmp_files(p.parent) and trash.sent == []


# ================================================================ 中5: Excel のテーブル（ListObject）


def build_table_xlsx(path: Path, n: int = 30) -> list[list[str]]:
    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "顧客"
    rows = person_rows(n)
    for r in rows:
        ws.append(r)
    tab = Table(displayName="Customers", ref=f"A1:D{len(rows)}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    wb.save(path)
    return rows


def open_table(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb["顧客"]
    return wb, ws, ws.tables["Customers"]


def test_xlsx_table_follows_row_deletion(tmp_path, trash):
    p = tmp_path / "src" / "t.xlsx"
    p.parent.mkdir()
    rows = build_table_xlsx(p)
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_rows", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    kept = [x for x in rows[1:] if not x[3]]
    wb, ws, tab = open_table(p)
    try:
        assert tab.ref == f"A1:D{1 + len(kept)}", tab.ref
        assert tab.autoFilter is not None and tab.autoFilter.ref == tab.ref
        assert [tc.name for tc in tab.tableColumns] == rows[0]
        assert [c.value for c in ws[1]][:4] == rows[0] and ws.max_row == 1 + len(kept)
    finally:
        wb.close()


def test_xlsx_table_follows_column_deletion(tmp_path, trash):
    p = tmp_path / "src" / "t.xlsx"
    p.parent.mkdir()
    rows = build_table_xlsx(p)
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_columns", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    wb, ws, tab = open_table(p)
    try:
        assert tab.ref == f"A1:A{len(rows)}", tab.ref
        assert [tc.name for tc in tab.tableColumns] == ["備考"]
        assert tab.autoFilter is not None and tab.autoFilter.ref == tab.ref
        assert ws["A1"].value == "備考" and ws["B1"].value is None
    finally:
        wb.close()


def test_xlsx_table_header_row_cannot_be_deleted(tmp_path):
    import openpyxl
    from openpyxl.worksheet.table import Table

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in person_rows(5):
        ws.append(r)
    ws.add_table(Table(displayName="T1", ref="A1:D6"))
    with pytest.raises(RemedyError, match="見出し行"):
        remedy_write._delete_rows(ws, [1, 3])
    with pytest.raises(RemedyError, match="データ行がすべて"):
        remedy_write._delete_rows(ws, [2, 3, 4, 5, 6])
    remedy_write._delete_rows(ws, [3, 4])  # 見出しとデータが残る削除は通り、範囲が詰まる
    assert ws.tables["T1"].ref == "A1:D4"
    wb.close()


# ================================================================ 中7: 別ジョブの同じファイル


def test_same_file_in_another_job_cannot_run_at_once(tmp_path, trash):
    p = make_csv(tmp_path)
    svc, store, job1 = make_service(tmp_path, [p])
    job2 = store.create_job([str(p.parent)], ScanOptions(temp_dir=str(tmp_path / "scantmp")))
    store.add_files(job2, [(str(p), p.stat().st_size, p.stat().st_mtime)])
    with store.tx() as c:
        c.execute("UPDATE files SET state='done' WHERE job_id=?", (job2,))
    analyze(svc, job1, p)
    analyze(svc, job2, p)
    params = {"method": "delete_rows", "confidences": ALL_RANKS}
    sha = svc.preview(job1, str(p), params)["expected_sha256"]
    svc._begin_run(job1, str(p), params, True, sha)  # job1 の実行が始まっている（記録は「始めた」）
    with pytest.raises(RemedyError, match="別の対処"):
        svc._begin_run(job2, str(p), params, True, sha)
    # 実行中の記録を、止まったものとみなして立て直さない
    assert svc.recover(str(p)) == []
    assert [x["state"] for x in store.unfinished_remediations()] == ["started"]


# ================================================================ 中8: 記録の更新の失敗


@pytest.mark.parametrize("failing", ["renamed", "swapped"])
def test_record_update_failure_keeps_the_original_name(tmp_path, monkeypatch, trash, failing):
    p = make_csv(tmp_path)
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])
    real = store.update_remediation

    def flaky(rid, **fields):
        if fields.get("state") == failing:
            raise sqlite3.OperationalError("database is locked（試験）")
        return real(rid, **fields)

    monkeypatch.setattr(store, "update_remediation", flaky)
    r = run(svc, job, p, method="delete_rows", confidences=ALL_RANKS)
    assert r["state"] == "error", r
    assert p.read_bytes() == original  # 元の名前に、元の中身が残る
    assert not new_tmp_files(p.parent) and not list(p.parent.glob("*対処前*")) and trash.sent == []
    assert store.remediations(job) == []


def test_record_update_failure_at_the_end_does_not_fail_the_run(tmp_path, monkeypatch, trash):
    p = make_csv(tmp_path)
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])
    real = store.update_remediation

    def flaky(rid, **fields):
        if fields.get("state") == "applied":
            raise sqlite3.OperationalError("database is locked（試験）")
        return real(rid, **fields)

    monkeypatch.setattr(store, "update_remediation", flaky)
    r = run(svc, job, p, method="delete_rows", confidences=ALL_RANKS)
    assert r["state"] == "done", r  # ファイルは確定済み。記録の最後の更新だけ失敗しても、成功として返す
    assert p.read_bytes() != original and len(trash.files()) == 1


def test_unreadable_original_at_commit_leaves_no_residue(tmp_path, monkeypatch, trash):
    p = make_csv(tmp_path)
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])
    analyze(svc, job, p)
    params = {"method": "delete_rows", "confidences": ALL_RANKS}
    pv = svc.preview(job, str(p), params)
    rid, run_state, ctx = svc._begin_run(job, str(p), params, True, pv["expected_sha256"])
    got = inline_runner("apply", ctx[2], 10)

    def locked(path):
        raise PermissionError("使用中（試験）")

    monkeypatch.setattr(remediate, "sha256_file", locked)
    with pytest.raises(RemedyError, match="読めませんでした"):
        remediate.commit_rewrite(store, ctx[6], str(p), ctx[5], got["orig_sha256"])
    monkeypatch.undo()
    assert p.read_bytes() == original and not new_tmp_files(p.parent) and store.remediations(job) == []


# ================================================================ 低


def _rec(path: Path, plan: dict, new_sha: str | None) -> dict:
    return {"top_path": str(path), "plan": plan, "new_sha256": new_sha, "id": 1}


@pytest.mark.parametrize("mode", ["recycle", "replace"])
def test_recovery_does_not_place_an_unverified_new_file(tmp_path, mode):
    """元の名前のファイルが無く、一時の新ファイルが書き終えたと確かめられないときは、元の名前へ置かない（書きかけを置かない）。"""
    path = tmp_path / "a.csv"
    tmp_new = tmp_path / "~piichk-new-abcdefgh-a.csv"
    tmp_new.write_bytes(b"half-written")
    plan = {"mode": mode, "tmp_new": str(tmp_new), "dest": str(path), "pre_path": str(tmp_path / "a（対処前 x）.csv")}
    state, msg = remediate.recover_record(_rec(path, plan, "0" * 64))
    assert state == "recovered" and not path.exists() and tmp_new.exists() and "書きかけ" in msg
    # 書き終えたと確かめられる（ハッシュが一致する）なら置く
    state, _ = remediate.recover_record(_rec(path, plan, remedy_write.sha256_file(str(tmp_new))))
    assert state == "recovered" and path.read_bytes() == b"half-written" and not tmp_new.exists()


def test_recovery_removes_leftover_temp_copy_and_sidecars(tmp_path):
    """SQLite の対処は元の完全な写しを一時に作る。落ちて残った写しと付属のファイルは、立て直しで消える。"""
    path = tmp_path / "people.sqlite"
    path.write_bytes(b"original")
    tmp_new = tmp_path / "~piichk-new-abcdefgh-people.sqlite"
    tmp_new.write_bytes(b"copy-with-personal-data")
    (tmp_path / (tmp_new.name + "-journal")).write_bytes(b"j")
    plan = {"mode": "replace", "tmp_new": str(tmp_new), "dest": str(path), "pre_path": None}
    state, _ = remediate.recover_record(_rec(path, plan, None))
    assert state == "recovered" and path.read_bytes() == b"original"
    assert not tmp_new.exists() and not (tmp_path / (tmp_new.name + "-journal")).exists()


def test_job_with_unfinished_remediation_cannot_be_deleted(tmp_path):
    store = Store(str(tmp_path / "jobs.sqlite3"))
    job = store.create_job([str(tmp_path)], ScanOptions())
    rid = store.add_remediation(job, str(tmp_path / "a.csv"), "delete_rows", {}, {}, state="renamed")
    with pytest.raises(ValueError, match="途中で止まった"):
        store.delete_job(job)
    assert store.job(job) is not None and store.remediation(job, rid) is not None
    store.update_remediation(rid, state="recovered")
    store.delete_job(job)  # 立て直したあとは消せる
    assert store.job(job) is None


def test_too_long_path_is_refused_before_starting(tmp_path, monkeypatch):
    long_path = str(tmp_path / ("d" * 120) / ("e" * 120) / "a.csv")
    with pytest.raises(RemedyError, match="長すぎ"):
        remediate.make_commit_plan(long_path, False)
    with pytest.raises(RemedyError, match="長すぎ"):
        remediate.check_recycle_path_length(long_path, extra=1)
    remediate.check_recycle_path_length(str(tmp_path / "a.csv"))  # 短ければ通る


def test_embedded_file_scan_uses_the_job_temp_dir(tmp_path, monkeypatch):
    seen: list[str] = []
    real_scan = engine.scan_file

    def spy(path, options, book):
        seen.append(path)
        return real_scan(path, options, book)

    monkeypatch.setattr(engine, "scan_file", spy)
    work = tmp_path / "jobtemp"
    work.mkdir()
    remedy_common.set_temp_dir(str(work))
    try:
        res = remedy_common.Residue(values=(), types={"email"}, ranks={3, 2, 1})
        assert res.embedded_has_pii("連絡先 taro@example.com".encode("utf-8"), "memo.txt") in (True, False)
    finally:
        remedy_common.set_temp_dir(None)
    assert seen and all(Path(s).is_relative_to(work) for s in seen), seen


def test_delete_succeeds_with_a_note_when_the_record_cannot_be_added(tmp_path, monkeypatch, trash):
    """ファイルを消したあとに記録の追加が失敗しても、「変わっていません」とは出さず、削除したことと記録を残せなかったことを知らせる。"""
    p = make_csv(tmp_path)
    svc, store, job = make_service(tmp_path, [p])
    analyze(svc, job, p)
    pv = svc.preview(job, str(p), {"method": "delete_file"})

    def broken(*a, **k):
        raise sqlite3.OperationalError("disk I/O error（試験）")

    monkeypatch.setattr(store, "add_remediation", broken)
    r = start_and_wait(svc, job, p, {"method": "delete_file"}, pv["expected_sha256"])
    assert r["state"] == "done" and not p.exists() and len(trash.files()) == 1, r
    assert any("記録を残せませんでした" in n for n in r["notes"]) and r["remediation_id"] is None


@windows_only
def test_new_file_inherits_attributes_creation_time_acl_and_streams(tmp_path):
    src = tmp_path / "src.csv"
    dst = tmp_path / "dst.csv"
    src.write_bytes(b"a,b\n")
    dst.write_bytes(b"a\n")
    with open(str(src) + ":Zone.Identifier", "wb") as f:
        f.write(b"[ZoneTransfer]\r\nZoneId=3\r\n")
    with open(str(src) + ":secret", "wb") as f:
        f.write(b"hidden-stream")
    ctypes.windll.kernel32.SetFileAttributesW(str(src), 0x2)  # 隠しファイル
    old = 1_500_000_000  # 2017 年
    k32 = remedy_meta._kernel32()
    ft_value = int((old + 11644473600) * 10_000_000)
    ft = remedy_meta._FILETIME(ft_value & 0xFFFFFFFF, ft_value >> 32)
    h = k32.CreateFileW(str(src), 0x100, 7, None, 3, 0x02000000, None)
    assert k32.SetFileTime(h, ctypes.byref(ft), None, None)
    k32.CloseHandle(h)
    subprocess.run(["icacls", str(src), "/grant", "Guests:(R)"], check=True, capture_output=True)  # 元だけの明示の許可
    notes = remedy_meta.inherit_file_metadata(str(src), str(dst))
    assert notes == [], notes
    assert ctypes.windll.kernel32.GetFileAttributesW(str(dst)) & 0x2  # 隠し属性
    assert open(str(dst) + ":Zone.Identifier", "rb").read().startswith(b"[ZoneTransfer]")
    assert open(str(dst) + ":secret", "rb").read() == b"hidden-stream"
    created = remedy_meta._FILETIME()
    h = k32.CreateFileW(str(dst), 0x80, 7, None, 3, 0x02000000, None)
    assert k32.GetFileTime(h, ctypes.byref(created), None, None)
    k32.CloseHandle(h)
    assert (created.hi << 32 | created.lo) == ft_value  # 作成日時
    acl = subprocess.run(["icacls", str(dst)], check=True, capture_output=True, text=True).stdout
    assert "Guests" in acl, acl  # 元のファイルにだけ付けた許可が引き継がれた


@windows_only
def test_rewrite_keeps_hidden_attribute_through_service(tmp_path, trash):
    p = make_csv(tmp_path)
    ctypes.windll.kernel32.SetFileAttributesW(str(p), 0x2)
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_rows", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    assert ctypes.windll.kernel32.GetFileAttributesW(str(p)) & 0x2
