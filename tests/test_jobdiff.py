"""前回の検査との違いの試験（保存した件数どうしを比べる）。"""

from __future__ import annotations

import time

from openpyxl import load_workbook

from pii_checker import jobdiff
from pii_checker.models import FileResult, ItemResult, ScanOptions, Status
from pii_checker.report import write_xlsx
from pii_checker.store import Store


def _job(store: Store, roots: list[str], files: dict[str, tuple[int, Status]], done: bool = True, **opts) -> str:
    """files: パス → (検出件数, 状態)。件数 0 は検出なし。件数 None はまだ調べていない。"""
    job = store.create_job(roots, ScanOptions(**opts))
    store.add_files(job, [(p, 1, 1.0) for p in files])
    for p, (n, status) in files.items():
        if n is None:
            continue
        item = ItemResult(path=p, kind="テキスト", status=status)
        if n:
            item.finding_counts = {"phone": n}
            item.conf_counts = {"phone:medium": n}
        store.save_result(job, FileResult(top_path=p, items=[item]))
    store.set_job(job, state="done" if done else "interrupted", finished=time.time())
    time.sleep(0.01)  # 作った順（created）で前後を決めるので、同じ時刻にしない
    return job


def test_no_previous(tmp_path):
    store = Store(str(tmp_path / "db.sqlite3"))
    j = _job(store, [r"C:\共有"], {r"C:\共有\a.txt": (1, Status.SCANNED)})
    d = jobdiff.compute(store, j)
    assert d["previous"] is None and d["rows"] == []
    assert jobdiff.summary_text(d) == jobdiff.NO_PREVIOUS


def test_diff_kinds_and_notes(tmp_path):
    store = Store(str(tmp_path / "db.sqlite3"))
    root = r"C:\共有"
    S = Status.SCANNED
    prev = _job(
        store,
        [root],
        {
            root + r"\same.txt": (2, S),
            root + r"\more.txt": (1, S),
            root + r"\less.txt": (5, S),
            root + r"\fixed.txt": (3, S),
            root + r"\deleted.txt": (1, S),
            root + r"\locked.zip": (1, S),
            root + r"\was_clean.txt": (0, S),
        },
    )
    # 途中で止めた検査・別の場所の検査は「前回」にしない
    _job(store, [root], {root + r"\same.txt": (9, S)}, done=False)
    _job(store, [r"C:\別"], {r"C:\別\x.txt": (1, S)})
    cur = _job(
        store,
        [root.lower() + "\\"],  # 大小文字・末尾の区切りの違いは同じ場所とみなす
        {
            root + r"\same.txt": (2, S),
            root + r"\more.txt": (4, S),
            root + r"\less.txt": (2, S),
            root + r"\fixed.txt": (0, S),
            root + r"\locked.zip": (0, Status.NEEDS_PASSWORD),
            root + r"\was_clean.txt": (1, S),
            root + r"\brand_new.txt": (1, S),
        },
    )
    d = jobdiff.compute(store, cur)
    assert d["previous"]["id"] == prev
    got = {(r["kind"], r["top_path"].rsplit("\\", 1)[-1]): r for r in d["rows"]}
    assert set(got) == {
        ("new", "was_clean.txt"),
        ("new", "brand_new.txt"),
        ("gone", "fixed.txt"),
        ("gone", "deleted.txt"),
        ("gone", "locked.zip"),
        ("increased", "more.txt"),
        ("decreased", "less.txt"),
    }
    assert got[("new", "was_clean.txt")]["note"] == "前回は検出なし"
    assert got[("new", "brand_new.txt")]["note"] == "前回は無かったファイル"
    assert got[("gone", "fixed.txt")]["note"].startswith("検出が無くなった")
    assert got[("gone", "deleted.txt")]["note"].startswith("今回は見つからない")
    assert got[("gone", "locked.zip")]["note"] == "今回は調べられなかった（要パスワード）"
    assert (got[("increased", "more.txt")]["n_prev"], got[("increased", "more.txt")]["n_now"]) == (1, 4)
    assert d["counts"] == {"new": 2, "gone": 3, "increased": 1, "decreased": 1}
    text = jobdiff.summary_text(d, "「前回との違い」シート")
    assert "新しく見つかった 2 件" in text and "無くなった 3 件" in text and "シート" in text

    out = tmp_path / "r.xlsx"
    write_xlsx(store, cur, str(out))
    wb = load_workbook(out, read_only=True)
    rows = list(wb["前回との違い"].iter_rows(values_only=True))
    assert rows[1][0] == "要約" and "新しく見つかった 2 件" in rows[1][1]
    assert len(rows) == 2 + 7
    overview = {r[0]: r[1] for r in wb["判定"].iter_rows(values_only=True) if r and r[0] and len(r) > 1}
    assert "新しく見つかった 2 件" in overview["前回との違い"]


def test_connection_target_ignores_user(tmp_path):
    store = Store(str(tmp_path / "db.sqlite3"))
    root = r"\\nas01\総務"
    conn = {"kind": "smb", "host": "NAS01", "share": "総務", "path": "", "user": "taro", "port": None}
    prev = _job(store, [root], {root + r"\a.txt": (1, Status.SCANNED)}, connection=conn)
    other = dict(conn, user="hanako", host="nas01")
    cur = _job(store, [root], {root + r"\a.txt": (1, Status.SCANNED)}, connection=other)
    assert jobdiff.previous_job(store, cur)["id"] == prev
    # 共有が違えば別の検査先
    third = _job(store, [root], {}, connection=dict(conn, share="経理"))
    assert jobdiff.previous_job(store, third) is None
    # フォルダの指定（接続方式でない）とも別
    folder_job = _job(store, [root], {})
    assert jobdiff.previous_job(store, folder_job) is None
