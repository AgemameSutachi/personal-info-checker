"""許可リスト（拾わない値）と利用者の規則（自前の語・正規表現）の試験。

電話番号・氏名・社員番号はすべて架空。結果の DB は必ず一時フォルダを指す。
"""

from __future__ import annotations

import json
import socket
import sys
import threading
import time
import urllib.error
import urllib.request
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pii_checker import cli, engine, runner, userrules  # noqa: E402
from pii_checker.detectors import CRITICAL_TYPES, PERSON_CUE_TYPES, PII_TYPES, mask_value, normalize  # noqa: E402
from pii_checker.engine import Collector  # noqa: E402
from pii_checker.models import ItemResult, ScanOptions, Status  # noqa: E402
from pii_checker.report import write_xlsx  # noqa: E402
from pii_checker.signature import engine_signature  # noqa: E402
from pii_checker.store import Store  # noqa: E402
from pii_checker.userrules import RuleError, check_pattern, parse_allow_list, parse_rules  # noqa: E402
from pii_checker.verdict import decide  # noqa: E402

RULES_TEXT = "# 社内の番号\n社員番号 | EMP-\\d{6} | 中\n"


def _collector(allow=(), rules_text="") -> tuple[ItemResult, Collector]:
    item = ItemResult(path="t", kind="", status=Status.PENDING)
    rules = parse_rules(rules_text) if rules_text else []
    m = userrules.Matcher(list(allow), rules) if (allow or rules) else None
    return item, Collector(item, 20, matcher=m)


# ---------------------------------------------------------------- 許可リスト


def test_allow_list_normalizes_and_matches_whole_value():
    m = userrules.Matcher(["03-1234-5678", "山田太郎", "Info@Example.JP"], [])
    # 全角・空白・ハイフン・括弧・大文字小文字の違いは同じ値として扱う
    assert m.is_allowed("phone", "０３（１２３４）５６７８")
    assert m.is_allowed("phone", "03 1234 5678")
    assert m.is_allowed("email", "info@example.jp")
    # 氏名に付いた敬称は外して比べる
    assert m.is_allowed("name", "山田太郎様")
    # 部分一致はしない（外しすぎない）
    assert not m.is_allowed("phone", "03-1234-56789")
    assert not m.is_allowed("name", "山田太郎二")
    assert not m.is_allowed("name", "山田")
    assert not m.is_allowed("address", "東京都山田太郎")


def test_allow_list_parse_skips_comments_and_duplicates():
    assert parse_allow_list("# 自社\n03-1234-5678\n\n０３－１２３４－５６７８\n山田太郎\n") == [
        "03-1234-5678",
        "山田太郎",
    ]
    with pytest.raises(RuleError) as e:
        parse_allow_list("03-1234-5678\n---\n")
    assert "2行目" in str(e.value) and "許可リスト" in str(e.value)


def test_allowed_findings_are_not_counted_but_counted_as_allowed():
    text = "電話 090-1234-5678 / 携帯 090-9999-8888"
    item, col = _collector(allow=["090 1234 5678"])
    col.add_text(text, "1行目")
    col.finish()
    assert item.finding_counts == {"phone": 1}  # 許可リストに無い番号は残る
    assert item.allowed == 1
    assert all("5678" not in f.masked_value for f in item.findings)
    # 陽性対照: 許可リストが無ければ2件とも数える
    plain, col2 = _collector()
    col2.add_text(text, "1行目")
    assert plain.finding_counts == {"phone": 2} and plain.allowed == 0


def test_allowed_name_is_not_a_person_cue():
    """見本の「山田太郎」を外したら、それを手がかりに住所を 中 へ引き上げない。"""
    text = "氏名 山田太郎\n所在 東京都千代田区丸の内1-1-1"
    item, col = _collector(allow=["山田太郎"])
    col.add_text(text, "本文")
    col.finish()
    assert item.allowed == 1
    assert item.conf_counts.get("address:medium", 0) == 0 and item.conf_counts.get("address:low", 0) == 1
    plain, col2 = _collector()
    col2.add_text(text, "本文")
    col2.finish()
    assert plain.conf_counts.get("address:medium", 0) == 1  # 陽性対照


def test_allowed_count_is_not_doubled_by_reread():
    item, col = _collector(allow=["090-1234-5678"])
    col.add_text("電話 090-1234-5678", "本文")
    col.add_text("電話 090-1234-5678", "残留", raw=True)  # 生データ走査で同じ値を見直しても数え直さない
    assert item.allowed == 1


# ---------------------------------------------------------------- 利用者の規則


def test_custom_rule_detects_with_rule_name_and_masks():
    item, col = _collector(rules_text=RULES_TEXT)
    col.add_text("担当 EMP-123456 のほか ＥＭＰ－６５４３２１", "本文")
    assert item.finding_counts == {"custom": 2}
    assert item.conf_counts == {"custom:medium": 2}
    f = item.findings[0]
    assert f.pii_type == "custom" and f.rule == "利用者の規則「社員番号」"
    assert "123456" not in f.masked_value and f.masked_value.endswith("3456")
    assert "123456" not in f.context


def test_plain_word_rule_and_confidence():
    rules = parse_rules("機密\n顧客番号 | C\\d{8} | 高")
    assert [(r.name, r.pattern, r.confidence) for r in rules] == [
        ("機密", "機密", "medium"),
        ("顧客番号", "C\\d{8}", "high"),
    ]
    item, col = _collector(rules_text="機密")
    col.add_text("社外秘・機密資料", "本文")
    assert item.finding_counts == {"custom": 1}
    assert mask_value("custom", "機密") == "＊＊"  # 2文字以下の値は、先頭1文字を残すと半分見えるので全部伏せる


def test_custom_rule_alternation_in_parentheses():
    rules = parse_rules("番号 | (?:EMP|STF)-\\d{4} | 低")
    assert rules[0].pattern == "(?:EMP|STF)-\\d{4}" and rules[0].confidence == "low"


def test_custom_values_can_be_allowed():
    item, col = _collector(allow=["EMP-000001"], rules_text=RULES_TEXT)
    col.add_text("見本 EMP-000001 本物 EMP-123456", "本文")
    assert item.finding_counts == {"custom": 1} and item.allowed == 1


def test_custom_rule_found_once_across_window_boundary():
    m = userrules.Matcher([], parse_rules(RULES_TEXT))
    # 区切りの境目（STEP・WINDOW の前後）に当たりを置いても、1回ずつ数える
    for pos in (userrules.STEP - 5, userrules.STEP, userrules.WINDOW - 5, userrules.WINDOW + 3):
        text = "x" * pos + "EMP-123456" + "y" * 3000
        hits = m.find(text)
        assert [(h.start, h.value) for h in hits] == [(pos, "EMP-123456")], pos
    many = ("あ" * 900 + "EMP-123456") * 6
    assert len(m.find(many)) == 6
    # 区切りの終わりで切られた当たりは、当たりの頭から当て直す（短く切れた値や「$」の取り違えを残さない）
    long_rule = userrules.Matcher([], parse_rules("番号 | N\\d+"))
    text = "x" * (userrules.STEP - 5) + "N" + "1" * 300 + " y"
    assert [h.value for h in long_rule.find(text)] == ["N" + "1" * 300]
    tail = userrules.Matcher([], parse_rules("末尾 | A\\w+$"))
    assert tail.find("x" * 700 + "A" + "b" * 400 + " z") == []
    assert [h.value for h in tail.find("x" * 2000 + " Abc")] == ["Abc"]


@pytest.mark.parametrize(
    "text, where, what",
    [
        ("社員番号 | EMP-\\d{6} | 大", "1行目", "確度"),
        ("# 説明\n | EMP", "2行目", "名前が空"),
        ("社員番号 |  | 中", "1行目", "正規表現が空"),
        ("a | b | c | d", "1行目", "括弧で囲んで"),
        ("社員番号 | EMP-(\\d{6} | 中", "1行目", "「(」に対応する「)」"),
        ("x | [abc", "1行目", "「[」に対応する「]」"),
        ("x | *a", "1行目", "繰り返す文字がありません"),
    ],
)
def test_rule_format_errors_name_the_line(text, where, what):
    with pytest.raises(RuleError) as e:
        parse_rules(text)
    msg = str(e.value)
    assert "自分で足す規則" in msg and where in msg and what in msg, msg


@pytest.mark.parametrize(
    "pattern, what",
    [
        ("(\\d+)+", "入れ子"),
        ("(a*b?)*c", "入れ子"),
        ("(a)\\1", "後方参照"),
        ("\\d*", "長さ0"),
        ("a" * 201, "長すぎ"),
        (".*.*.*=", "続けた書き方"),
        ("\\s*\\s*\\s*x", "続けた書き方"),
        # 回数の上限が大きい繰り返しは、上限の無い繰り返しと同じに扱う
        ("[a-z]{1,200}[a-z]{1,200}!", "続けた書き方"),
        ("x{0,1000}x{0,1000}z", "続けた書き方"),
        # 繰り返しの中の選択肢が同じ文字で始まりうる・長さが違う
        ("(?:a|a)+b", "選択肢"),
        ("(a|aa)+b", "選択肢"),
        ("(?:a|b|ab)*c", "選択肢"),
        ("(\\w|\\d)+x", "選択肢"),
    ],
)
def test_slow_or_risky_patterns_are_refused(pattern, what):
    started = time.perf_counter()
    msg = check_pattern(pattern)
    assert msg and (what in msg or (what == "入れ子" and "繰り返しの中に" in msg)), msg
    assert time.perf_counter() - started < 2  # 断るまでに待たせない（画面の受け付けで待つ時間）


def test_slow_pattern_found_by_timing_is_refused_quickly():
    """形では見分けられないが当てるのが桁違いに遅い規則は、時間を測って断る。"""
    started = time.perf_counter()
    msg = check_pattern("\\w+\\w?\\d+x")
    assert msg and "時間がかかりすぎる" in msg, msg
    # 固まらずに断れることを確かめる（測って断るので、遅い PC・混んだ PC では数秒かかる。CI では 2 秒を超えた）
    assert time.perf_counter() - started < 15


@pytest.mark.parametrize(
    "pattern",
    [
        "EMP-\\d{6}",
        "C\\d{8}",
        "機密",
        "(?:EMP|STF)-\\d{6}",
        "(?:EMP|STF)-\\d{4}",
        "(\\d{3}-)+\\d",
        "\\w+@\\w+",
        ".*X",
        "^C\\d{3}",
        "\\w+\\d+",  # 最後の2つの繰り返しは、後ろで失敗しないので続いていてもよい
        "([0-9a-f])+x",  # 文字数の2乗で遅くなるが、区切って当てるので桁違いには遅くならない
        "株式会社（本社）",
    ],
)
def test_ordinary_patterns_are_accepted(pattern):
    started = time.perf_counter()
    assert check_pattern(pattern) is None
    # 普通の規則は受け付けで待たせない（手元では数ミリ秒。CI の共有の機械は遅く、0.5 秒では 0.52 秒で落ちた）
    assert time.perf_counter() - started < 2.0


def test_rules_use_multiline():
    """「^」「$」は行の頭・終わりに当たる（速さを確かめる段も同じ指定で当てる）。"""
    m = userrules.Matcher([], parse_rules("顧客 | ^C\\d{3}"))
    assert [h.value for h in m.find("x\nC123\nC456")] == ["C123", "C456"]
    end = userrules.Matcher([], parse_rules("末尾 | \\d{3}$"))
    assert [h.value for h in end.find("a123\nb456\nc7890x")] == ["123", "456"]


def test_fullwidth_letters_fold_but_fullwidth_symbols_stay_literal():
    """全角の英数字だけ半角に寄せ、全角の記号（＊・（）など）は正規表現の記号にしない。"""
    assert parse_rules("x | ＥＭＰ\\ｄ{3}")[0].pattern == "EMP\\d{3}"
    paren = userrules.Matcher([], parse_rules("本社 | 株式会社（本社）"))
    assert [h.value for h in paren.find(normalize("株式会社（本社）"))] == ["株式会社(本社)"]
    assert paren.find("株式会社本社") == []  # 「（）」をグループにしない
    star = userrules.Matcher([], parse_rules("注意 | 注意＊"))
    assert [h.value for h in star.find(normalize("注意＊"))] == ["注意*"]
    assert star.find("注意注意") == []  # 「＊」を繰り返しの記号にしない
    # 半角の記号は正規表現の記号のまま
    assert [h.value for h in userrules.Matcher([], parse_rules("x | A(?:B|C)")).find("AC")] == ["AC"]


@pytest.mark.parametrize(
    "line, pattern",
    [
        ("名前 | []|] | 高", "[]|]"),
        ("名前 | [^]|]+ | 高", "[^]|]+"),
        ("名前 | [a|b] | 高", "[a|b]"),
        ("名前 | (?:A|B)\\|C | 高", "(?:A|B)\\|C"),
    ],
)
def test_split_fields_keeps_bar_inside_class(line, pattern):
    rules = parse_rules(line)
    assert (rules[0].pattern, rules[0].confidence) == (pattern, "high")


def test_custom_type_is_not_critical_and_risk_follows_confidence():
    assert "custom" in PII_TYPES and "custom" not in CRITICAL_TYPES and "custom" not in PERSON_CUE_TYPES
    mid = ItemResult(path="a", kind="", status=Status.SCANNED, conf_counts={"custom:medium": 3})
    high = ItemResult(path="b", kind="", status=Status.SCANNED, conf_counts={"custom:high": 1})
    low = ItemResult(path="c", kind="", status=Status.SCANNED, conf_counts={"custom:low": 1})
    assert engine.risk_of([mid]) == "中"  # 重要な種類なら 中 でも「高」になる
    assert engine.risk_of([high]) == "高"
    assert engine.risk_of([low]) == "低"


def test_verdict_reports_allowed_count():
    base = {
        "item_status": {"scanned": 1},
        "confidence": {},
        "files_total": 1,
        "files_done": 1,
        "folder_errors": 0,
        "links": 0,
        "files_with_findings": 0,
    }
    v = decide({**base, "allowed": 3, "allow_list": 2}, "done")
    assert v.level == "clean" and any("拾わない値（許可リスト）で外した検出 3 件" in d for d in v.detail)
    # 外したために検出が無くなったときは、判定の文にその前提を添える
    assert v.detail[0].startswith("拾わない値（許可リスト）の値を除いて、検出は無かった")
    assert "拾わない値（許可リスト）の値を除いて、「個人情報は無い」" in v.next_step
    v0 = decide({**base, "allowed": 0, "allow_list": 2}, "done")
    assert any("拾わない値（許可リスト）で外した検出 0 件" in d for d in v0.detail)
    assert "値を除いて" not in v0.next_step  # 何も外していなければ添えない
    assert not any("許可リスト" in d for d in decide(base, "done").detail)  # 使っていなければ出さない


# ---------------------------------------------------------------- 使い回し・ジョブ・レポート


def test_signature_changes_with_allow_list_and_rules():
    rules = [r.to_dict() for r in parse_rules(RULES_TEXT)]
    base = engine_signature(ScanOptions())
    a = engine_signature(ScanOptions(allow_values=["03-1234-5678", "山田太郎"]))
    # 書き順・全角半角・ハイフンの違いだけでは変わらない
    a2 = engine_signature(ScanOptions(allow_values=["山田太郎", "０３１２３４５６７８"]))
    r = engine_signature(ScanOptions(custom_rules=rules))
    r2 = engine_signature(ScanOptions(custom_rules=[{**rules[0], "confidence": "high"}]))
    assert len({base, a, r, r2}) == 4
    assert a == a2


@pytest.fixture
def folder(tmp_path):
    d = tmp_path / "data"
    d.mkdir()
    (d / "f.txt").write_text("電話 090-1234-5678\n担当 EMP-123456\n", encoding="utf-8")
    return d


def _run(store, folder, tmp_path, **opts):
    job = store.create_job([str(folder)], ScanOptions(temp_dir=str(tmp_path), **opts))
    r = runner.JobRunner(store, job, workers=1)
    r.start()
    r.wait(180)
    assert store.job(job)["state"] == "done"
    return job


def _reused(store, job):
    return store.query("SELECT reused_from FROM files WHERE job_id=?", (job,))[0][0]


def test_job_counts_allowed_and_reuse_follows_settings(folder, tmp_path):
    store = Store(str(tmp_path / "jobs.sqlite3"))
    rules = [r.to_dict() for r in parse_rules(RULES_TEXT)]
    j1 = _run(store, folder, tmp_path, allow_values=["090-1234-5678"], custom_rules=rules)
    s = store.summary(j1)
    assert s["allowed"] == 1 and s["allow_list"] == 1 and s["custom_rules"] == 1
    assert s["types"] == {"custom": 1}
    v = decide(s, "done")
    assert v.level == "found" and any("拾わない値（許可リスト）で外した検出 1 件" in d for d in v.detail)

    # 同じ許可リスト（書き方だけ違う）・同じ規則なら使い回し、外した件数も写す
    j2 = _run(store, folder, tmp_path, allow_values=["０９０ １２３４ ５６７８"], custom_rules=rules)
    assert _reused(store, j2) == j1 and store.summary(j2)["allowed"] == 1
    # 許可リストを変えたら使い回さない
    j3 = _run(store, folder, tmp_path, custom_rules=rules)
    assert _reused(store, j3) is None
    assert store.summary(j3)["types"] == {"custom": 1, "phone": 1} and store.summary(j3)["allowed"] == 0
    # 規則を変えたら使い回さない
    j4 = _run(store, folder, tmp_path, custom_rules=[{**rules[0], "confidence": "low"}])
    assert _reused(store, j4) is None

    # レポート: 件数と規則の名前は出すが、許可リストの値そのものは出さない
    out = tmp_path / "out" / "r.xlsx"
    write_xlsx(store, j1, str(out))
    wb = load_workbook(out, read_only=True)
    cells = [str(c) for row in wb["判定"].iter_rows(values_only=True) for c in row if c is not None]
    wb.close()
    joined = "\n".join(cells)
    assert "拾わない値（許可リスト）で外した検出 1 件" in joined
    assert "設定: 拾わない値（許可リスト）" in cells and "1 件" in cells
    assert "設定: 自分で足す規則" in cells and "1 件（社員番号）" in cells
    assert "090-1234-5678" not in joined and "5678" not in joined


def test_resume_uses_saved_settings(folder, tmp_path):
    """再開しても、ジョブに保存した許可リスト・規則で検査する。"""
    store = Store(str(tmp_path / "jobs.sqlite3"))
    rules = [r.to_dict() for r in parse_rules(RULES_TEXT)]
    job = store.create_job(
        [str(folder)], ScanOptions(temp_dir=str(tmp_path), allow_values=["090-1234-5678"], custom_rules=rules)
    )
    r = runner.JobRunner(store, job, workers=1)
    assert r.options.allow_values == ["090-1234-5678"] and r.options.custom_rules == rules
    saved = json.loads(store.query("SELECT options FROM jobs WHERE id=?", (job,))[0][0])
    assert saved["allow_values"] == ["090-1234-5678"] and saved["custom_rules"] == rules


def test_old_db_gets_allowed_column(tmp_path):
    import sqlite3

    p = tmp_path / "old.sqlite3"
    con = sqlite3.connect(p)
    con.execute(
        "CREATE TABLE items (job_id TEXT, top_path TEXT, seq INTEGER, path TEXT, kind TEXT, status TEXT, reason TEXT,"
        " method TEXT, size INTEGER, mtime REAL, depth INTEGER, password_kind TEXT, counts TEXT, notes TEXT,"
        " conf_counts TEXT DEFAULT '{}', PRIMARY KEY (job_id, top_path, seq))"
    )
    con.commit()
    con.close()
    store = Store(str(p))
    assert "allowed" in {r[1] for r in store.query("PRAGMA table_info(items)")}
    store.set_setting("k", "v")
    assert store.get_setting("k") == "v" and store.get_setting("none", "d") == "d"


# ---------------------------------------------------------------- コマンド


def _cli_scan(tmp_path, folder, *extra):
    out = tmp_path / "out"
    out.mkdir(exist_ok=True)
    db = tmp_path / "cli.sqlite3"
    code = cli.main(["--db", str(db), "scan", str(folder), "--no-ask", "-o", str(out / "r.xlsx"), *extra])
    return code, db


def test_cli_allow_and_rules_files(tmp_path, folder, capsys):
    allow = tmp_path / "allow.txt"
    allow.write_text("# 自社の番号\n090-1234-5678\n", encoding="utf-8-sig")  # メモ帳の BOM 付きも読める
    rules = tmp_path / "rules.txt"
    rules.write_text(RULES_TEXT, encoding="utf-8")
    code, db = _cli_scan(tmp_path, folder, "--allow-file", str(allow), "--rules-file", str(rules))
    assert code == 0
    out = capsys.readouterr().out
    assert (
        "拾わない値（許可リスト）1 件・自分で足す規則 1 件（社員番号） を使います" in out
        and "拾わない値（許可リスト）で外した検出 1 件" in out
    )
    store = Store(str(db))
    job = store.jobs()[0]["id"]
    opts = store.job(job)["options"]
    assert opts["allow_values"] == ["090-1234-5678"] and opts["custom_rules"][0]["name"] == "社員番号"
    assert store.summary(job)["types"] == {"custom": 1}


def test_cli_bad_rules_file_is_refused_before_scan(tmp_path, folder, capsys):
    rules = tmp_path / "rules.txt"
    rules.write_text("社員番号 | EMP-\\d{6} | 中\n顧客 | (C\\d+ | 中\n", encoding="utf-8")
    code, db = _cli_scan(tmp_path, folder, "--rules-file", str(rules))
    assert code == 2
    err = capsys.readouterr().err
    assert "2行目" in err and "「(」に対応する「)」" in err
    assert not db.exists()  # ジョブも DB も作らない


def test_cli_rules_file_errors(tmp_path, folder, capsys):
    bad = tmp_path / "sjis.txt"
    bad.write_bytes("社員番号 | EMP".encode("cp932"))
    assert _cli_scan(tmp_path, folder, "--rules-file", str(bad))[0] == 2
    assert "UTF-8" in capsys.readouterr().err
    assert _cli_scan(tmp_path, folder, "--allow-file", str(tmp_path / "none.txt"))[0] == 2
    assert "見つかりません" in capsys.readouterr().err
    code = cli.main(["--db", str(tmp_path / "x.sqlite3"), "scan", "--resume", "job", "--allow-file", str(bad)])
    assert code == 2 and "一緒に使えません" in capsys.readouterr().err


# ---------------------------------------------------------------- Web の API

TOKEN = "test-token-rules"


@pytest.fixture
def web(tmp_path, folder):
    import uvicorn

    from pii_checker.web.server import create_app

    app = create_app(db=str(tmp_path / "web.sqlite3"), token=TOKEN)
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    srv = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
    th = threading.Thread(target=srv.run, daemon=True)
    th.start()
    for _ in range(100):
        if srv.started:
            break
        time.sleep(0.05)
    yield f"http://127.0.0.1:{port}", folder, tmp_path
    srv.should_exit = True
    th.join(timeout=10)


def _call(base, path, body=None):
    headers = {"X-Token": TOKEN}
    if body is not None:
        headers.update({"X-Requested-With": "pii-checker", "Content-Type": "application/json"})
    req = urllib.request.Request(
        base + path,
        data=json.dumps(body).encode() if body is not None else None,
        method="POST" if body is not None else "GET",
        headers=headers,
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read())


def test_web_api_takes_settings_and_remembers_them(web):
    base, folder, tmp_path = web
    assert _call(base, "/api/settings") == (200, {"allow_list": "", "custom_rules": ""})
    # 書き方の誤りは検査を始める前に断り、前回の値も書き換えない
    st, d = _call(base, "/api/jobs", {"roots": [str(folder)], "options": {"custom_rules": "a | b | c | d"}})
    assert st == 400 and "1行目" in d["detail"] and "自分で足す規則" in d["detail"]
    assert _call(base, "/api/jobs")[1] == []
    assert _call(base, "/api/settings")[1]["custom_rules"] == ""

    allow_text = "# 自社\r\n090-1234-5678\r\n"
    st, d = _call(
        base,
        "/api/jobs",
        {
            "roots": [str(folder)],
            "options": {"workers": 1, "temp_dir": str(tmp_path), "allow_list": allow_text, "custom_rules": RULES_TEXT},
        },
    )
    assert st == 200, d
    job_id = d["id"]
    for _ in range(360):
        st, j = _call(base, f"/api/jobs/{job_id}")
        if not j["running"] and j["state"] in ("done", "error", "cancelled"):
            break
        time.sleep(0.5)
    assert j["state"] == "done"
    assert j["summary"]["allowed"] == 1 and j["summary"]["types"] == {"custom": 1}
    assert any("拾わない値（許可リスト）で外した検出 1 件" in x for x in j["verdict"]["detail"])
    # 前回の値（入れたままの文字列）を返す
    assert _call(base, "/api/settings")[1] == {"allow_list": allow_text, "custom_rules": RULES_TEXT}
    # 画面の明細の種類は「利用者の規則（社員番号）」のように規則の名前つきで、生の値は出ない
    st, f = _call(base, f"/api/jobs/{job_id}/findings")
    assert f["rows"][0]["type"] == "利用者の規則（社員番号）" and f["rows"][0]["rule"] == "利用者の規則「社員番号」"
    # Excel・CSV の種類の列は今のまま（規則の名前は検出規則の列）
    from pii_checker.report import write_csv

    csv_path = write_csv(Store(str(tmp_path / "web.sqlite3")), job_id, str(tmp_path / "csv"))[0]
    csv_text = Path(csv_path).read_text(encoding="utf-8-sig")
    assert ",利用者の規則," in csv_text and "利用者の規則「社員番号」" in csv_text
    assert "123456" not in json.dumps(f, ensure_ascii=False)
    # レポートのページの見出しにも件数を出し、許可リストの値は出さない
    st, rep = _call(base, f"/api/jobs/{job_id}/report")
    flat = json.dumps(rep, ensure_ascii=False)
    assert "設定: 拾わない値（許可リスト）" in flat and "1 件（社員番号）" in flat and "090-1234-5678" not in flat
    # 一覧の内訳・詳細の異なる値・項目の件数は、利用者の規則を規則の名前ごとに出す
    st, files = _call(base, f"/api/jobs/{job_id}/files?only_findings=true")
    row = files["rows"][0]
    assert row["types"] == {"利用者の規則（社員番号）": 1}
    assert row["distinct"] == {"利用者の規則（社員番号）": 1}
    st, its = _call(base, f"/api/jobs/{job_id}/items?top_path={urllib.request.quote(row['top_path'])}")
    assert its[0]["counts"] == {"利用者の規則（社員番号）": 1} and its[0]["distinct"] == {"利用者の規則（社員番号）": 1}


def test_custom_counts_and_distinct_are_split_by_rule_name():
    """利用者の規則は、件数（rule_counts）と異なる値（custom/規則の名前）を規則ごとに数える。種類の件数は今のまま。"""
    item, col = _collector(rules_text="社員番号 | EMP-\\d{6}\n顧客番号 | C\\d{8} | 高\n")
    col.add_text("EMP-123456 EMP-123456 EMP-654321 C12345678", "本文")
    col.finish()
    assert item.finding_counts == {"custom": 4}
    assert item.rule_counts == {"社員番号": 3, "顧客番号": 1}
    assert item.distinct_counts == {"custom/社員番号:medium": 2, "custom/顧客番号:high": 1}
    from pii_checker import reportflags

    assert reportflags.distinct_text(json.dumps(item.distinct_counts, ensure_ascii=False)) == (
        "利用者の規則（社員番号） 2・利用者の規則（顧客番号） 1"
    )
    assert reportflags.type_label("custom/社:番号") == "利用者の規則（社:番号）"
    assert reportflags.distinct_medium({"custom/社:番号:medium": 3}) == {"custom/社:番号": 3}


def test_split_rule_counts_keeps_totals():
    from pii_checker.web.server import split_rule_counts

    assert split_rule_counts({"custom": 3, "phone": 1}, '{"社員番号": 2, "顧客番号": 1}') == {
        "phone": 1,
        "custom/社員番号": 2,
        "custom/顧客番号": 1,
    }
    # 規則ごとの件数が無い（古い版）・合計が合わないときは分けない
    assert split_rule_counts({"custom": 3}, "{}") == {"custom": 3}
    assert split_rule_counts({"custom": 3}, None) == {"custom": 3}
    assert split_rule_counts({"custom": 3}, '{"社員番号": 2}') == {"custom": 3}


@pytest.mark.parametrize("newline", ["\r\n", "\r", "\n"])
def test_rules_fold_crlf_before_matching(newline):
    """セル内改行・メールの本文の CR LF・CR でも、行末の規則が各行に当たる。位置は元の本文での位置を返す。"""
    m = userrules.Matcher([], parse_rules("番号 | ^C\\d{3}$"))
    text = f"C123{newline}C456{newline}end"
    hits = m.find(text)
    assert [h.value for h in hits] == ["C123", "C456"]
    assert [text[h.start : h.end] for h in hits] == ["C123", "C456"]


def test_rules_fold_crlf_positions_after_many_lines():
    m = userrules.Matcher([], parse_rules("番号 | C\\d{3}$"))
    text = "\r\n".join(f"x{i} C{i:03d}" for i in range(300))  # WINDOW をまたぐ長さ
    hits = m.find(text)
    assert len(hits) == 300
    assert all(text[h.start : h.end] == h.value for h in hits)


def test_busy_pc_does_not_let_a_slow_pattern_through(monkeypatch):
    """PC が混んで、どの試し文字列も時間の線を超えて見えるときも、遅い規則を通さず普通の規則は通す。

    線を超えて測り直し「遅くない」と出ても、そこで受け付けずに残りの試し文字列（本当に遅い並び）まで調べる。
    """
    monkeypatch.setattr(userrules, "_SPEED_ABORT", 0.0)
    assert check_pattern("\w+\w?\d+x"), "遅い規則を通した"
    for p in ("EMP-\d{6}", "\w+@\w+", "C\d{3}$"):
        assert check_pattern(p) is None, p


@pytest.mark.parametrize(
    "value, masked",
    [
        ("K-7", "*-*"),  # 短い値は全部伏せる（残す分で値が全部見えてしまうため）
        ("C123", "****"),
        ("EMP-123456", "***-**3456"),
        ("C12345678", "*****5678"),
        ("機密", "＊＊"),
        ("取扱注意", "取＊＊＊"),
    ],
)
def test_custom_values_are_masked_even_when_short(value, masked):
    assert mask_value("custom", value) == masked
