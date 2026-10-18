"""個人情報チェッカーの試験用検体（架空の個人情報を仕込んだファイル群）を生成する。

使い方:
    python tests/make_corpus.py <出力先フォルダ> [--clean] [--only 名前,名前...]

- 出力先には検体と manifest.json を作る。
- manifest.json の entries には、各ファイル（アーカイブ内なら「外側!内側」）ごとに
  expect（detect / clean / uncheckable）と仕込んだ値の種類 kinds を書く。
- 環境の都合で作れなかった検体は skipped に理由つきで書き、全体は止めない。
- 個人情報はすべて架空。マイナンバーは検査数字だけ正しい架空番号、クレジットカードは
  各ブランドが公開しているテスト用番号を使う。
- Office の COM 操作は窓を出さず（Visible=False / DisplayAlerts=False）、終わったら必ず Quit する。
  PowerPoint と Outlook は単一インスタンスのため、元から起動していた場合は Quit しない。
"""

from __future__ import annotations

import argparse
import csv
import email.encoders
import email.header
import email.mime.base
import email.mime.multipart
import email.mime.text
import email.utils
import gc
import io
import json
import os
import shutil
import subprocess
import sys
import tarfile
import tempfile
import traceback
import zipfile
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Callable

# ---------------------------------------------------------------------------
# 架空の個人情報
# ---------------------------------------------------------------------------

SEVEN_ZIP = r"C:\Program Files\7-Zip\7z.exe"
CSCRIPT32 = r"C:\Windows\SysWOW64\cscript.exe"
FONT_CANDIDATES = [
    r"C:\Windows\Fonts\meiryo.ttc",
    r"C:\Windows\Fonts\msgothic.ttc",
]
PASSWORD = "secret123"
JA_LOCALE = ";LANGID=0x0411;CP=932;COUNTRY=0"

ALL_KINDS = [
    "name",
    "address",
    "postal_code",
    "phone",
    "email",
    "birthdate",
    "my_number",
    "bank_account",
    "credit_card",
]


def my_number_check_digit(base11: str) -> int:
    """マイナンバー（個人番号）の上位11桁から検査数字を計算する。

    総務省令の式: 下から n 桁目の数字 Pn に重み Qn（n<=6 は n+1、7<=n<=11 は n-5）を掛けて合計し、
    11 で割った余りが 1 以下なら 0、それ以外は 11 - 余り。
    """
    if len(base11) != 11 or not base11.isdigit():
        raise ValueError(f"マイナンバーの上位は11桁の数字で渡してください: {base11!r}")
    total = 0
    for n in range(1, 12):
        digit = int(base11[11 - n])
        weight = n + 1 if n <= 6 else n - 5
        total += digit * weight
    remainder = total % 11
    return 0 if remainder <= 1 else 11 - remainder


def make_my_number(base11: str) -> str:
    """検査数字つきの12桁の架空マイナンバーを返す。"""
    return base11 + str(my_number_check_digit(base11))


def make_invalid_my_number(base11: str) -> str:
    """検査数字だけをわざと外した12桁の数字を返す（誤検出の試験用）。"""
    return base11 + str((my_number_check_digit(base11) + 1) % 10)


def luhn_ok(number: str) -> bool:
    """Luhn 検査が通るかを返す。"""
    digits = [int(c) for c in number if c.isdigit()]
    total = 0
    for i, d in enumerate(reversed(digits)):
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def make_luhn_invalid(prefix15: str) -> str:
    """Luhn 検査がわざと通らない16桁の数字を返す（伝票番号などの誤検出試験用）。"""
    for last in range(10):
        candidate = prefix15 + str(last)
        if not luhn_ok(candidate):
            return candidate
    raise RuntimeError("Luhn 不成立の番号を作れませんでした")


PEOPLE = [
    {
        "name": "山田 太郎",
        "kana": "ヤマダ タロウ",
        "postal": "〒100-0005",
        "address": "東京都千代田区丸の内1-2-3",
        "phone": "03-1234-5678",
        "mobile": "090-1234-5678",
        "email": "taro.yamada@example.jp",
        "birth": "1985年4月1日",
        "my_number": make_my_number("12345678901"),
        "bank": "サンプル銀行 丸の内支店 普通 1234567",
        "card": "4111 1111 1111 1111",
    },
    {
        "name": "佐藤 花子",
        "kana": "サトウ ハナコ",
        "postal": "〒530-0001",
        "address": "大阪府大阪市北区梅田9-8-7",
        "phone": "06-9876-5432",
        "mobile": "080-9876-5432",
        "email": "hanako.sato@example.com",
        "birth": "1990年12月24日",
        "my_number": make_my_number("98765432109"),
        "bank": "サンプル銀行 梅田支店 普通 7654321",
        "card": "5555 5555 5555 4444",
    },
    {
        "name": "鈴木 一郎",
        "kana": "スズキ イチロウ",
        "postal": "〒220-0012",
        "address": "神奈川県横浜市西区みなとみらい4-5-6",
        "phone": "045-123-4567",
        "mobile": "070-1111-2222",
        "email": "ichiro.suzuki@example.net",
        "birth": "1978年7月7日",
        "my_number": make_my_number("24681357902"),
        "bank": "サンプル銀行 横浜支店 当座 2468135",
        "card": "3530 1113 3330 0000",
    },
]

HEADERS = [
    "氏名",
    "フリガナ",
    "郵便番号",
    "住所",
    "電話番号",
    "携帯電話",
    "メールアドレス",
    "生年月日",
    "マイナンバー",
    "口座番号",
    "クレジットカード",
]
PERSON_KEYS = [
    "name",
    "kana",
    "postal",
    "address",
    "phone",
    "mobile",
    "email",
    "birth",
    "my_number",
    "bank",
    "card",
]


def table_rows() -> list[list[str]]:
    """見出し行＋データ行の表を返す。"""
    rows = [list(HEADERS)]
    for person in PEOPLE:
        rows.append([person[k] for k in PERSON_KEYS])
    return rows


def pii_text() -> str:
    """文章形式の個人情報（1人目を中心に全種類を含む）を返す。"""
    p = PEOPLE[0]
    return "\n".join(
        [
            "【顧客情報】",
            f"氏名：{p['name']}（{p['kana']}）",
            f"住所：{p['postal']} {p['address']}",
            f"電話番号：{p['phone']}　携帯：{p['mobile']}",
            f"メールアドレス：{p['email']}",
            f"生年月日：{p['birth']}",
            f"マイナンバー：{p['my_number']}",
            f"振込先口座：{p['bank']}",
            f"クレジットカード：{p['card']}",
        ]
    )


def short_pii_lines() -> list[str]:
    """画像に描く用の短い行（OCR しやすい大きさに収める）。"""
    p = PEOPLE[0]
    return [
        f"氏名 {p['name']}",
        f"住所 {p['postal']} {p['address']}",
        f"電話 {p['mobile']}",
        f"メール {p['email']}",
        f"生年月日 {p['birth']}",
        f"マイナンバー {p['my_number']}",
        f"口座 普通 1234567",
        f"カード {p['card']}",
    ]


def clean_rows() -> list[list[str]]:
    """個人情報を含まないが紛らわしい値の表（誤検出の試験用）。"""
    return [
        [
            "注文番号",
            "注文日",
            "取引先",
            "商品名",
            "JANコード",
            "数量",
            "金額",
            "管理番号",
            "伝票番号",
        ],
        [
            "ORD-20261017-0001",
            "2026年10月17日",
            "株式会社サンプル商事",
            "ステンレスボトル 500ml",
            "4901234567894",
            "12",
            "¥1,234,567",
            make_invalid_my_number("11122233344"),
            make_luhn_invalid("123456789012345"),
        ],
        [
            "ORD-20261017-0002",
            "2026/10/18",
            "サンプル物産有限会社",
            "A4コピー用紙 2500枚",
            "4987654321098",
            "3",
            "¥8,800",
            make_invalid_my_number("55566677788"),
            make_luhn_invalid("987654321098765"),
        ],
        [
            "ORD-20261017-0003",
            "令和8年11月1日",
            "テスト工業株式会社",
            "工具セット Ver.10.0.19045",
            "4512345678901",
            "1",
            "¥45,000",
            make_invalid_my_number("13579246801"),
            make_luhn_invalid("246813579024681"),
        ],
    ]


def clean_text() -> str:
    """個人情報を含まない文章（誤検出の試験用）。"""
    lines = ["【受注一覧】2026年10月17日 時点", "営業部 第二課 作成"]
    for row in clean_rows()[1:]:
        lines.append(
            f"注文番号 {row[0]} / {row[1]} / {row[2]} / {row[3]} / JAN {row[4]} / "
            f"数量 {row[5]} / 金額 {row[6]} / 管理番号 {row[7]} / 伝票番号 {row[8]}"
        )
    lines.append("合計金額 ¥1,288,367（税込）　納期 2026-11-30")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 生成の文脈（manifest・COM アプリ・一時フォルダ）
# ---------------------------------------------------------------------------


class Corpus:
    """出力先・manifest・共有する Office アプリを持つ。"""

    def __init__(self, out_dir: Path, work_dir: Path):
        self.out = out_dir
        self.work = work_dir
        self.entries: list[dict] = []
        self.skipped: list[dict] = []
        self.errors: list[str] = []
        self._apps: dict[str, object] = {}
        self._app_owned: dict[str, bool] = {}
        self._app_pid: dict[str, int | None] = {}
        self._cache: dict[str, bytes] = {}

    # --- manifest ---------------------------------------------------------
    def add(
        self,
        rel: str,
        expect: str,
        kinds: list[str] | None = None,
        fmt: str = "",
        note: str = "",
        password: str | None = None,
    ) -> None:
        """manifest に1件足す。rel は出力先からの相対パス（アーカイブ内は「!」区切り）。"""
        if expect not in ("detect", "clean", "uncheckable"):
            raise ValueError(f"expect が不正です: {expect}")
        top = rel.split("!")[0]
        if not (self.out / top).exists():
            raise FileNotFoundError(f"manifest に載せるファイルがありません: {top}")
        entry: dict = {
            "path": rel,
            "expect": expect,
            "kinds": list(kinds or []) if expect != "clean" else [],
            "format": fmt,
        }
        if note:
            entry["note"] = note
        if password:
            entry["password"] = password
        self.entries.append(entry)

    def skip(self, rel: str, reason: str) -> None:
        """作れなかった検体を記録する。"""
        if ("SaveAs" in reason or "保存できない" in reason) and current_session_id() == 0:
            reason += f"（{SESSION0_NOTE}）"
        self.skipped.append({"path": rel, "skipped": True, "reason": reason})
        print(f"  [skip] {rel}: {reason}")

    def path(self, rel: str) -> Path:
        """出力先の絶対パスを返し、親フォルダを作る。"""
        p = self.out / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        return p

    def winpath(self, rel: str) -> str:
        """COM に渡す Windows 形式の絶対パスを返す。"""
        return str(self.path(rel).resolve())

    def write_manifest(self) -> None:
        """manifest.json を書く（途中でも呼べる）。"""
        data = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "generator": "tests/make_corpus.py",
            "password_list_hint": [PASSWORD],
            "people": PEOPLE,
            "entries": self.entries,
            "skipped": self.skipped,
        }
        tmp = self.out / "manifest.json.tmp"
        tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(tmp, self.out / "manifest.json")

    # --- 共有する部品 -------------------------------------------------------
    def xlsx_bytes(self) -> bytes:
        """個人情報入りの xlsx をバイト列で返す（アーカイブ用に使い回す）。"""
        if "xlsx" not in self._cache:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "顧客名簿"
            for row in table_rows():
                ws.append(row)
            buf = io.BytesIO()
            wb.save(buf)
            self._cache["xlsx"] = buf.getvalue()
        return self._cache["xlsx"]

    def docx_bytes(self) -> bytes:
        """個人情報入りの docx（本文）をバイト列で返す。"""
        if "docx" not in self._cache:
            import docx

            d = docx.Document()
            d.add_heading("顧客情報", level=1)
            for line in pii_text().splitlines():
                d.add_paragraph(line)
            buf = io.BytesIO()
            d.save(buf)
            self._cache["docx"] = buf.getvalue()
        return self._cache["docx"]

    def csv_text(self) -> str:
        """個人情報入りの CSV 文字列を返す（改行は CRLF）。"""
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\r\n")
        writer.writerows(table_rows())
        return buf.getvalue()

    # --- COM アプリ ---------------------------------------------------------
    def app(self, progid: str, single_instance: bool = False):
        """Office アプリを遅延起動して返す。single_instance は既存起動を流用する種類。"""
        if progid in self._apps:
            return self._apps[progid]
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        owned = True
        app = None
        if single_instance:
            try:
                app = win32com.client.GetActiveObject(progid)
                owned = False
            except Exception:
                app = None
        if app is None:
            # DispatchEx で必ず新しいプロセスを起こす（利用者が開いている窓を乗っ取らない）
            app = win32com.client.DispatchEx(progid)
        name = progid.split(".")[0]
        if name in ("Excel", "Word"):
            app.Visible = False
            app.DisplayAlerts = False if name == "Excel" else 0
        if name == "Excel":
            app.ScreenUpdating = False
            app.EnableEvents = False
            app.AskToUpdateLinks = False
        if name == "PowerPoint":
            app.DisplayAlerts = 1  # ppAlertsNone
        self._apps[progid] = app
        self._app_owned[progid] = owned
        if owned:
            self._app_pid[progid] = _office_pid(app, name)
        return app

    def quit_apps(self) -> None:
        """起動した Office アプリをすべて終了する（失敗しても続ける）。

        Quit が失敗したら、起動時に控えた PID（自分で起こしたプロセス）だけを強制終了し、
        Office のプロセスを残さない。
        """
        for progid, app in list(self._apps.items()):
            owned = self._app_owned.get(progid, True)
            pid = self._app_pid.pop(progid, None)
            try:
                if owned:
                    app.Quit()
                    print(f"  [quit] {progid}")
                else:
                    print(f"  [keep] {progid}（元から起動していたため終了しない）")
            except Exception as exc:  # noqa: BLE001
                print(f"  [warn] {progid} の終了に失敗: {exc}")
                if owned and pid:
                    subprocess.run(
                        ["taskkill", "/PID", str(pid), "/F"],
                        capture_output=True,
                        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                    )
                    print(f"  [kill] {progid} PID={pid}（自分で起動したもの）")
        self._apps.clear()
        gc.collect()


# ---------------------------------------------------------------------------
# 便利関数
# ---------------------------------------------------------------------------


def _office_pid(app, name: str) -> int | None:
    """Office アプリのプロセス ID を窓ハンドルから求める（取れなければ None）。"""
    try:
        import win32process

        if name == "Excel":
            hwnd = app.Hwnd
        elif name == "Access":
            hwnd = app.hWndAccessApp()
        elif name == "PowerPoint":
            hwnd = app.HWND
        else:
            return None
        return int(win32process.GetWindowThreadProcessId(int(hwnd))[1]) or None
    except Exception:  # noqa: BLE001
        return None


def run(cmd: list[str], cwd: Path | None = None, timeout: int = 180) -> str:
    """外部コマンドを窓を出さずに実行し、失敗なら例外にする。"""
    proc = subprocess.run(
        cmd,
        cwd=str(cwd) if cwd else None,
        capture_output=True,
        timeout=timeout,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    out = proc.stdout.decode("cp932", errors="replace") + proc.stderr.decode("cp932", errors="replace")
    if proc.returncode != 0:
        raise RuntimeError(f"{Path(cmd[0]).name} が終了コード {proc.returncode}: {out.strip()[:500]}")
    return out


def find_font() -> str:
    """日本語を描ける TrueType フォントのパスを返す。"""
    for f in FONT_CANDIDATES:
        if Path(f).exists():
            return f
    raise FileNotFoundError("日本語フォント（meiryo.ttc / msgothic.ttc）が見つかりません")


def render_text_image(lines: list[str], font_size: int = 32):
    """文字を画像として描いた PIL.Image を返す（OCR の試験用）。"""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(find_font(), font_size)
    line_h = int(font_size * 1.6)
    width = 1400
    height = line_h * len(lines) + 80
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((40, 40 + i * line_h), line, fill="black", font=font)
    return img


def excel_write_table(ws, rows: list[list[str]]) -> None:
    """COM のワークシートへ表を書き込む（数字の列が数値化されないよう文字列書式にする）。"""
    n_rows, n_cols = len(rows), len(rows[0])
    rng = ws.Range(ws.Cells(1, 1), ws.Cells(n_rows, n_cols))
    rng.NumberFormat = "@"
    rng.Value = tuple(tuple(r) for r in rows)


def current_session_id() -> int:
    """このプロセスが動いている Windows のセッション番号（0 ならサービスのセッション）を返す。"""
    import ctypes

    sid = ctypes.c_ulong(0)
    pid = ctypes.windll.kernel32.GetCurrentProcessId()
    if not ctypes.windll.kernel32.ProcessIdToSessionId(pid, ctypes.byref(sid)):
        return -1
    return int(sid.value)


# セッション 0（サービスとして動くプログラム等）から起動した 64bit Office は、
# プロファイルが systemprofile になり、C:\Windows\System32\config\systemprofile\Desktop が
# 無いと Open / SaveAs が「SaveAs メソッドが失敗しました」「コマンドは正常終了できませんでした」で
# 失敗する（実測 2026-10-17）。そのときの回避策:
#   Excel: 新規ブックを作る前に DefaultSaveFormat を目的の形式にし、SaveCopyAs で書き出す
#   Word : Range.ExportFragment で docx / rtf / odt を書き出す（.doc は不可）
SESSION0_NOTE = (
    "セッション0（サービス）から起動した Office は systemprofile\\Desktop が無いと SaveAs/Open が失敗する。"
    "管理者で C:\\Windows\\System32\\config\\systemprofile\\Desktop を作るか、対話セッションで実行すれば作れる"
)


def excel_new_workbook(xl, file_format: int = 51):
    """指定の保存形式を既定にして新規ブックを作る（SaveCopyAs でその形式になるようにする）。"""
    changed = False
    if file_format != 51:
        try:
            xl.DefaultSaveFormat = file_format
            changed = True
        except Exception:  # noqa: BLE001
            changed = False  # 既定にできない形式（Excel 4.0 以前など）は SaveAs だけで試す
    try:
        wb = xl.Workbooks.Add()
    finally:
        if changed:
            xl.DefaultSaveFormat = 51
    return wb


def excel_save(wb, path: str, file_format: int) -> str:
    """SaveAs を試し、失敗したら（ブックの形式が合っていれば）SaveCopyAs で保存する。使った方法を返す。"""
    try:
        wb.SaveAs(path, file_format)
        return "SaveAs"
    except Exception as first:  # noqa: BLE001
        if int(wb.FileFormat) != file_format:
            raise RuntimeError(
                f"SaveAs(FileFormat={file_format}) が失敗し、SaveCopyAs でもこの形式にできない: "
                f"{com_error_text(first)}"
            ) from first
        wb.SaveCopyAs(path)
        return "SaveCopyAs"


WORD_EXPORTABLE = {
    12,
    13,
    16,
    6,
    23,
}  # ExportFragment で書き出せた形式（docx 系・rtf・odt）


def word_save(doc, path: str, file_format: int) -> str:
    """SaveAs2 を試し、失敗したら ExportFragment で保存する。使った方法を返す。"""
    try:
        doc.SaveAs2(path, file_format)
        return "SaveAs2"
    except Exception as first:  # noqa: BLE001
        if file_format not in WORD_EXPORTABLE:
            raise RuntimeError(
                f"SaveAs2(FileFormat={file_format}) が失敗し、ExportFragment も非対応の形式: "
                f"{com_error_text(first)}"
            ) from first
        doc.Content.ExportFragment(path, file_format)
        return "ExportFragment"


def encrypt_ooxml(plain: bytes, password: str) -> bytes:
    """msoffcrypto で OOXML（docx / xlsx）を読み取りパスワード付きに暗号化する。"""
    from msoffcrypto.format.ooxml import OOXMLFile

    src = io.BytesIO(plain)
    dst = io.BytesIO()
    OOXMLFile(src).encrypt(password, dst)
    return dst.getvalue()


# ---------------------------------------------------------------------------
# 1. Excel
# ---------------------------------------------------------------------------


def gen_excel_openpyxl(c: Corpus) -> None:
    """openpyxl で作れる Excel 検体（通常・非表示シート・コメント）。"""
    import openpyxl
    from openpyxl.comments import Comment

    c.path("excel/meibo.xlsx").write_bytes(c.xlsx_bytes())
    c.add("excel/meibo.xlsx", "detect", ALL_KINDS, "xlsx", "表形式（見出し＋3行）")

    # 非表示シートだけに値（表示シートは無害）
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表紙"
    ws["A1"] = "社内資料（表紙）"
    hidden = wb.create_sheet("hidden_list")
    for row in table_rows():
        hidden.append(row)
    hidden.sheet_state = "hidden"
    wb.save(c.path("excel/hidden_sheet.xlsx"))
    c.add("excel/hidden_sheet.xlsx", "detect", ALL_KINDS, "xlsx", "非表示シートにのみ値")

    # ごく非表示（veryHidden）シートだけに値
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表紙"
    ws["A1"] = "社内資料（表紙）"
    vh = wb.create_sheet("very_hidden_list")
    for row in table_rows():
        vh.append(row)
    vh.sheet_state = "veryHidden"
    wb.save(c.path("excel/very_hidden_sheet.xlsx"))
    c.add(
        "excel/very_hidden_sheet.xlsx",
        "detect",
        ALL_KINDS,
        "xlsx",
        "veryHidden シートにのみ値（Excel の画面からは再表示できない）",
    )

    # セルのコメントだけに値
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "集計"
    ws["A1"] = "担当者"
    ws["B1"] = "備考はコメント参照"
    ws["A1"].comment = Comment(pii_text(), "作成者")
    wb.save(c.path("excel/comment_only.xlsx"))
    c.add("excel/comment_only.xlsx", "detect", ALL_KINDS, "xlsx", "セルのコメントにのみ値")


def gen_excel_com(c: Corpus) -> None:
    """Excel COM で作る検体（テキストボックス・xlsm・xlsb・xls・古い形式・ods）。"""
    xl = c.app("Excel.Application")

    # テキストボックス（図形）だけに値
    wb = excel_new_workbook(xl, 51)
    try:
        ws = wb.Worksheets(1)
        ws.Range("A1").Value = "図形の中を参照"
        shape = ws.Shapes.AddTextbox(1, 50, 40, 420, 220)
        shape.TextFrame2.TextRange.Text = pii_text()
        how = excel_save(wb, c.winpath("excel/textbox_only.xlsx"), 51)
    finally:
        wb.Close(False)
    c.add(
        "excel/textbox_only.xlsx",
        "detect",
        ALL_KINDS,
        "xlsx",
        f"テキストボックス（図形）にのみ値（Excel COM {how}）",
    )

    # 表形式の各種保存形式
    formats = [
        ("excel/meibo.xlsm", 52, "xlsm", "マクロ有効ブック（マクロ自体は無し）"),
        ("excel/meibo.xlsb", 50, "xlsb", "バイナリブック"),
        ("excel/meibo.xls", 56, "xls (BIFF8)", "Excel 97-2003"),
        ("excel/meibo_excel95.xls", 39, "xls (BIFF5/7)", "Excel 5.0/95 形式"),
        ("excel/meibo_excel4.xls", 33, "xls (BIFF4)", "Excel 4.0 ワークシート形式"),
        ("excel/meibo_excel4w.xlw", 35, "xlw (BIFF4W)", "Excel 4.0 ブック形式"),
        ("excel/meibo_excel3.xls", 29, "xls (BIFF3)", "Excel 3.0 形式"),
        ("excel/meibo_excel2.xls", 16, "xls (BIFF2)", "Excel 2.1 形式"),
        ("excel/meibo.ods", 60, "ods", "OpenDocument スプレッドシート"),
    ]
    for rel, file_format, fmt, note in formats:
        wb = excel_new_workbook(xl, file_format)
        try:
            # 古い形式は1シートしか持てないので、余分なシートを消しておく
            while wb.Worksheets.Count > 1:
                wb.Worksheets(wb.Worksheets.Count).Delete()
            ws = wb.Worksheets(1)
            ws.Name = "顧客名簿"
            excel_write_table(ws, table_rows())
            try:
                how = excel_save(wb, c.winpath(rel), file_format)
            except Exception as exc:  # noqa: BLE001
                detail = com_error_text(exc) if type(exc).__name__ == "com_error" else str(exc)
                c.skip(rel, f"Excel で FileFormat={file_format} に保存できない: {detail}")
                continue
        finally:
            wb.Close(False)
        if c.path(rel).exists() and c.path(rel).stat().st_size > 0:
            c.add(
                rel,
                "detect",
                ALL_KINDS,
                fmt,
                f"{note}（Excel COM FileFormat={file_format}・{how}）",
            )
        else:
            c.skip(
                rel,
                f"FileFormat={file_format} の保存は例外なしだがファイルができなかった",
            )


def com_error_text(exc: Exception) -> str:
    """pywintypes.com_error から読める文を取り出す。"""
    try:
        args = exc.args
        if len(args) >= 3 and args[2] and len(args[2]) >= 3 and args[2][2]:
            return f"{args[1]} / {str(args[2][2]).strip()}"
        return str(args[1] if len(args) > 1 else exc)
    except Exception:  # noqa: BLE001
        return str(exc)


def gen_csv(c: Corpus) -> None:
    """文字コード違いの CSV 4種。"""
    text = c.csv_text()
    for rel, enc, label in [
        ("csv/meibo_utf8bom.csv", "utf-8-sig", "UTF-8 BOM付き"),
        ("csv/meibo_sjis.csv", "cp932", "Shift_JIS（cp932）"),
        ("csv/meibo_utf16.csv", "utf-16", "UTF-16 LE BOM付き"),
        ("csv/meibo_eucjp.csv", "euc_jp", "EUC-JP"),
    ]:
        c.path(rel).write_bytes(text.encode(enc))
        c.add(rel, "detect", ALL_KINDS, "csv", label)


# ---------------------------------------------------------------------------
# 2. Word
# ---------------------------------------------------------------------------


def gen_word_docx(c: Corpus) -> None:
    """python-docx で作れる Word 検体（本文・表・ヘッダー・フッター）。"""
    import docx

    c.path("word/body.docx").write_bytes(c.docx_bytes())
    c.add("word/body.docx", "detect", ALL_KINDS, "docx", "本文にのみ値")

    d = docx.Document()
    d.add_paragraph("顧客一覧（下表）")
    rows = table_rows()
    table = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            table.cell(i, j).text = val
    d.save(c.path("word/table.docx"))
    c.add("word/table.docx", "detect", ALL_KINDS, "docx", "表にのみ値（本文は無害）")

    d = docx.Document()
    d.add_paragraph("本文には個人情報はありません。")
    header = d.sections[0].header
    header.paragraphs[0].text = pii_text().replace("\n", "　")
    d.save(c.path("word/header_only.docx"))
    c.add("word/header_only.docx", "detect", ALL_KINDS, "docx", "ヘッダーにのみ値")

    d = docx.Document()
    d.add_paragraph("本文には個人情報はありません。")
    footer = d.sections[0].footer
    footer.paragraphs[0].text = pii_text().replace("\n", "　")
    d.save(c.path("word/footer_only.docx"))
    c.add("word/footer_only.docx", "detect", ALL_KINDS, "docx", "フッターにのみ値")


def gen_word_com(c: Corpus) -> None:
    """Word COM で作る検体（テキストボックス・コメント・doc・rtf・odt・Word 6/95）。"""
    wd = c.app("Word.Application")

    doc = wd.Documents.Add()
    try:
        doc.Content.Text = "本文には個人情報はありません（図形を参照）。"
        box = doc.Shapes.AddTextbox(1, 80, 120, 420, 220)
        box.TextFrame.TextRange.Text = pii_text()
        how = word_save(doc, c.winpath("word/textbox_only.docx"), 16)
    finally:
        doc.Close(False)
    c.add(
        "word/textbox_only.docx",
        "detect",
        ALL_KINDS,
        "docx",
        f"テキストボックスにのみ値（Word COM {how}）",
    )

    doc = wd.Documents.Add()
    try:
        doc.Content.Text = "本文には個人情報はありません（コメントを参照）。"
        doc.Comments.Add(doc.Range(0, 4), pii_text())
        how = word_save(doc, c.winpath("word/comment_only.docx"), 16)
    finally:
        doc.Close(False)
    c.add(
        "word/comment_only.docx",
        "detect",
        ALL_KINDS,
        "docx",
        f"コメントにのみ値（Word COM {how}）",
    )

    # 本文に個人情報を書いた新規文書を各形式で保存する
    # （既存ファイルを Open する方式はセッション0では Open が失敗するので使わない）
    targets = [
        ("word/body.doc", 0, "doc", "Word 97-2003"),
        ("word/body.rtf", 6, "rtf", "リッチテキスト"),
        ("word/body.odt", 23, "odt", "OpenDocument テキスト"),
    ]
    # Word 6.0/95 の保存用コンバーターがあれば足す
    word6_format = None
    try:
        for conv in wd.FileConverters:
            label = f"{conv.FormatName} {conv.ClassName}"
            if conv.CanSave and ("Word 6" in label or "Word 95" in label or "MSWord6" in label):
                word6_format = conv.SaveFormat
                break
    except Exception:  # noqa: BLE001
        word6_format = None
    if word6_format is not None:
        targets.append(("word/body_word6.doc", word6_format, "doc (Word 6.0/95)", "Word 6.0/95"))
    else:
        c.skip(
            "word/body_word6.doc",
            "Word に Word 6.0/95 形式で保存するコンバーターが無い",
        )

    for rel, file_format, fmt, note in targets:
        doc = wd.Documents.Add()
        try:
            doc.Content.Text = pii_text()
            try:
                how = word_save(doc, c.winpath(rel), file_format)
            except Exception as exc:  # noqa: BLE001
                detail = com_error_text(exc) if type(exc).__name__ == "com_error" else str(exc)
                c.skip(rel, f"Word で FileFormat={file_format} に保存できない: {detail}")
                continue
        finally:
            doc.Close(False)
        c.add(
            rel,
            "detect",
            ALL_KINDS,
            fmt,
            f"{note}（Word COM FileFormat={file_format}・{how}）",
        )


# ---------------------------------------------------------------------------
# 3. PowerPoint
# ---------------------------------------------------------------------------


def gen_powerpoint(c: Corpus) -> None:
    """PowerPoint の pptx・ppt。COM が使えなければ pptx だけ python-pptx で作る。"""
    note = "1枚目のテキストボックスに全種類、2枚目のノートに氏名と電話"
    failures: dict[str, str] = {}
    try:
        pp = c.app("PowerPoint.Application", single_instance=True)
    except Exception as exc:  # noqa: BLE001
        pp = None
        failures["powerpoint/meibo.pptx"] = failures["powerpoint/meibo.ppt"] = (
            f"PowerPoint COM が使えない: {com_error_text(exc)}"
        )
    if pp is not None:
        pres = pp.Presentations.Add(False)  # WithWindow=False で窓を出さない
        try:
            slide = pres.Slides.Add(1, 12)  # ppLayoutBlank
            box = slide.Shapes.AddTextbox(1, 30, 30, 650, 420)
            box.TextFrame.TextRange.Text = pii_text()
            slide2 = pres.Slides.Add(2, 12)
            slide2.Shapes.AddTextbox(1, 30, 30, 650, 100).TextFrame.TextRange.Text = "2枚目（無害）"
            slide2.NotesPage.Shapes.Placeholders(2).TextFrame.TextRange.Text = (
                f"発表者メモ：連絡先 {PEOPLE[1]['name']} {PEOPLE[1]['mobile']}"
            )
            for rel, file_format, fmt in [
                ("powerpoint/meibo.pptx", 24, "pptx"),
                ("powerpoint/meibo.ppt", 1, "ppt"),
            ]:
                try:
                    pres.SaveAs(c.winpath(rel), file_format)
                    c.add(
                        rel,
                        "detect",
                        ALL_KINDS,
                        fmt,
                        f"{note}（COM ppSaveAs={file_format}）",
                    )
                except Exception:  # noqa: BLE001
                    try:
                        pres.SaveCopyAs(c.winpath(rel), file_format)
                        c.add(
                            rel,
                            "detect",
                            ALL_KINDS,
                            fmt,
                            f"{note}（COM SaveCopyAs={file_format}）",
                        )
                    except Exception as exc:  # noqa: BLE001
                        failures[rel] = f"PowerPoint の SaveAs/SaveCopyAs({file_format}) が失敗: {com_error_text(exc)}"
        finally:
            pres.Close()

    # pptx は python-pptx で代わりに作れる
    if "powerpoint/meibo.pptx" in failures:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(5))
        tb.text_frame.text = pii_text()
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])
        slide2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1)).text_frame.text = "2枚目（無害）"
        slide2.notes_slide.notes_text_frame.text = f"発表者メモ：連絡先 {PEOPLE[1]['name']} {PEOPLE[1]['mobile']}"
        prs.save(c.path("powerpoint/meibo.pptx"))
        c.add(
            "powerpoint/meibo.pptx",
            "detect",
            ALL_KINDS,
            "pptx",
            f"{note}（python-pptx で作成。COM は失敗: {failures['powerpoint/meibo.pptx']}）",
        )
    if "powerpoint/meibo.ppt" in failures:
        c.skip("powerpoint/meibo.ppt", failures["powerpoint/meibo.ppt"])


# ---------------------------------------------------------------------------
# 4. PDF
# ---------------------------------------------------------------------------


def _pdf_text_page(doc, text: str):
    """日本語の本文を1ページに書く。"""
    page = doc.new_page()
    page.insert_text((50, 72), text, fontname="japan", fontsize=11)
    return page


def gen_pdf(c: Corpus) -> None:
    """PDF 検体（テキスト・注釈・フォーム・添付・暗号化・画像のみ）。"""
    import pymupdf

    # テキスト PDF
    doc = pymupdf.open()
    _pdf_text_page(doc, pii_text())
    doc.save(c.path("pdf/text.pdf"))
    doc.close()
    c.add(
        "pdf/text.pdf",
        "detect",
        ALL_KINDS,
        "pdf",
        "テキスト層あり（CJK フォント japan）",
    )

    # 注釈だけに値
    doc = pymupdf.open()
    page = _pdf_text_page(doc, "本文には個人情報はありません（注釈を参照）。")
    annot = page.add_text_annot((60, 120), pii_text())
    annot.set_info(title="確認者", content=pii_text())
    annot.update()
    doc.save(c.path("pdf/annotation_only.pdf"))
    doc.close()
    c.add("pdf/annotation_only.pdf", "detect", ALL_KINDS, "pdf", "付箋注釈の内容にのみ値")

    # フォーム欄
    doc = pymupdf.open()
    page = _pdf_text_page(doc, "申込書（入力欄を参照）")
    p = PEOPLE[0]
    fields = [
        ("name", p["name"]),
        ("address", f"{p['postal']} {p['address']}"),
        ("phone", p["mobile"]),
        ("email", p["email"]),
        ("birthdate", p["birth"]),
        ("my_number", p["my_number"]),
        ("bank_account", p["bank"]),
        ("credit_card", p["card"]),
    ]
    for i, (fname, value) in enumerate(fields):
        w = pymupdf.Widget()
        w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
        w.field_name = fname
        w.field_value = value
        w.rect = pymupdf.Rect(150, 110 + i * 30, 500, 132 + i * 30)
        w.text_fontsize = 10
        page.add_widget(w)
    doc.save(c.path("pdf/form_fields.pdf"))
    doc.close()
    c.add(
        "pdf/form_fields.pdf",
        "detect",
        ALL_KINDS,
        "pdf",
        "フォーム欄の値にのみ（本文は無害）",
    )

    # 添付ファイル（中に xlsx）
    doc = pymupdf.open()
    _pdf_text_page(doc, "添付ファイルを参照してください。")
    doc.embfile_add("meibo.xlsx", c.xlsx_bytes(), filename="meibo.xlsx", desc="顧客名簿")
    doc.save(c.path("pdf/attachment.pdf"))
    doc.close()
    c.add("pdf/attachment.pdf", "detect", ALL_KINDS, "pdf", "埋め込み添付の xlsx にのみ値")
    c.add(
        "pdf/attachment.pdf!meibo.xlsx",
        "detect",
        ALL_KINDS,
        "xlsx",
        "PDF の添付ファイル",
    )

    # 空パスワードで暗号化（開ける）
    doc = pymupdf.open()
    _pdf_text_page(doc, pii_text())
    doc.save(
        c.path("pdf/encrypted_empty_password.pdf"),
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner-only-pass",
        user_pw="",
        permissions=pymupdf.PDF_PERM_PRINT,
    )
    doc.close()
    c.add(
        "pdf/encrypted_empty_password.pdf",
        "detect",
        ALL_KINDS,
        "pdf",
        "AES-256 暗号化だが閲覧パスワードは空（パスワードなしで開ける）",
    )

    # パスワード付き
    doc = pymupdf.open()
    _pdf_text_page(doc, pii_text())
    doc.save(
        c.path("pdf/encrypted_password.pdf"),
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner-secret-999",
        user_pw=PASSWORD,
    )
    doc.close()
    c.add(
        "pdf/encrypted_password.pdf",
        "uncheckable",
        ALL_KINDS,
        "pdf",
        "閲覧パスワード付き。パスワード一覧に secret123 を渡せば detect になるべき",
        password=PASSWORD,
    )

    # 画像だけの PDF（OCR が必要）
    img = render_text_image(short_pii_lines(), font_size=34)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(
        pymupdf.Rect(20, 40, 575, 40 + 555 * img.height / img.width),
        stream=buf.getvalue(),
    )
    doc.save(c.path("pdf/image_only.pdf"))
    doc.close()
    c.add(
        "pdf/image_only.pdf",
        "detect",
        [
            "name",
            "address",
            "postal_code",
            "phone",
            "email",
            "birthdate",
            "my_number",
            "bank_account",
            "credit_card",
        ],
        "pdf (image only)",
        "テキスト層なし。OCR しないと読めない（OCR なしなら検出できないのが正しい）",
    )


# ---------------------------------------------------------------------------
# 5. Access
# ---------------------------------------------------------------------------

DB_VERSION_40 = 64  # dbVersion40（Jet 4 / Access 2000-2003 .mdb）
DB_VERSION_120 = 128  # dbVersion120（ACE / .accdb）
DB_ENCRYPT = 2  # dbEncrypt


def _sql_quote(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _create_meibo_table(db, table: str = "顧客名簿") -> None:
    cols = ", ".join(f"[{h}] TEXT(255)" for h in HEADERS)
    db.Execute(f"CREATE TABLE [{table}] (ID COUNTER PRIMARY KEY, {cols})")


def _insert_people(db, table: str = "顧客名簿", repeat: int = 1) -> None:
    names = ", ".join(f"[{h}]" for h in HEADERS)
    for _ in range(repeat):
        for person in PEOPLE:
            values = ", ".join(_sql_quote(person[k]) for k in PERSON_KEYS)
            db.Execute(f"INSERT INTO [{table}] ({names}) VALUES ({values})")


def gen_access(c: Corpus) -> None:
    """Access COM（64bit）の DBEngine で accdb / mdb(Jet4) と派生検体を作る。"""
    acc = c.app("Access.Application")
    dbe = acc.DBEngine

    def create(rel: str, option: int, password: str | None = None):
        target = c.winpath(rel)
        if Path(target).exists():
            Path(target).unlink()
        locale = JA_LOCALE + (f";pwd={password}" if password else "")
        return dbe.CreateDatabase(target, locale, option)

    # 通常の accdb
    db = create("access/meibo.accdb", DB_VERSION_120)
    try:
        _create_meibo_table(db)
        _insert_people(db)
    finally:
        db.Close()
    c.add(
        "access/meibo.accdb",
        "detect",
        ALL_KINDS,
        "accdb",
        "テーブル「顧客名簿」（Access COM DBEngine）",
    )

    # Jet4（Access 2000-2003）の mdb
    db = create("access/meibo_jet4.mdb", DB_VERSION_40)
    try:
        _create_meibo_table(db)
        _insert_people(db)
    finally:
        db.Close()
    c.add("access/meibo_jet4.mdb", "detect", ALL_KINDS, "mdb (Jet4)", "Access 2000 形式")

    # 挿入してから削除し、最適化していない accdb
    rel = "access/deleted_rows_not_compacted.accdb"
    db = create(rel, DB_VERSION_120)
    try:
        _create_meibo_table(db)
        _insert_people(db, repeat=20)
        db.Execute("DELETE FROM [顧客名簿]")
        db.Execute("CREATE TABLE [備品] (ID COUNTER PRIMARY KEY, [品名] TEXT(100), [数量] LONG)")
        db.Execute("INSERT INTO [備品] ([品名], [数量]) VALUES ('ボールペン', 120)")
        db.Execute("INSERT INTO [備品] ([品名], [数量]) VALUES ('クリアファイル', 300)")
    finally:
        db.Close()
    raw = c.path(rel).read_bytes()
    found = {
        "utf-16le": PEOPLE[0]["email"].encode("utf-16-le") in raw,
        "ascii(圧縮)": PEOPLE[0]["email"].encode("ascii") in raw,
        "氏名utf-16le": PEOPLE[0]["name"].encode("utf-16-le") in raw,
    }
    residual = ", ".join(k for k, v in found.items() if v) or "見つからず"
    c.add(
        rel,
        "detect",
        ALL_KINDS,
        "accdb",
        "削除済み領域にのみ残る値（表「顧客名簿」は空、表「備品」に無害な2行）。"
        f"生成直後のファイル内に残っていた形: {residual}。"
        "通常の SQL 読み取りでは見えないので、生バイト走査をしない検査では clean になる",
    )

    # データベースパスワード付き
    rel = "access/password_protected.accdb"
    db = create(rel, DB_VERSION_120 + DB_ENCRYPT, password=PASSWORD)
    try:
        _create_meibo_table(db)
        _insert_people(db)
    finally:
        db.Close()
    c.add(
        rel,
        "uncheckable",
        ALL_KINDS,
        "accdb (encrypted)",
        "データベースパスワード付き（暗号化）。パスワード一覧に secret123 を渡せば detect になるべき",
        password=PASSWORD,
    )

    rel = "access/password_protected_jet4.mdb"
    db = create(rel, DB_VERSION_40 + DB_ENCRYPT, password=PASSWORD)
    try:
        _create_meibo_table(db)
        _insert_people(db)
    finally:
        db.Close()
    c.add(
        rel,
        "uncheckable",
        ALL_KINDS,
        "mdb (Jet4, encrypted)",
        "データベースパスワード付き。パスワード一覧に secret123 を渡せば detect になるべき",
        password=PASSWORD,
    )

    # OLE オブジェクト型フィールドにファイルを格納
    rel = "access/ole_field.accdb"
    db = create(rel, DB_VERSION_120)
    try:
        db.Execute("CREATE TABLE [書類] (ID COUNTER PRIMARY KEY, [書類名] TEXT(100), [本体] LONGBINARY)")
        rs = db.OpenRecordset("書類")
        try:
            rs.AddNew()
            rs.Fields("書類名").Value = "名簿（Excel）"
            rs.Fields("本体").AppendChunk(c.xlsx_bytes())
            rs.Update()
            rs.AddNew()
            rs.Fields("書類名").Value = "名簿（CSV）"
            rs.Fields("本体").AppendChunk(c.csv_text().encode("cp932"))
            rs.Update()
        finally:
            rs.Close()
    finally:
        db.Close()
    c.add(
        rel,
        "detect",
        ALL_KINDS,
        "accdb",
        "OLE オブジェクト型（LONGBINARY）列に xlsx と Shift_JIS の CSV の生バイトを格納。"
        "Access の OLE Package ヘッダーは付けていない（生のファイルそのもの）。テキスト列は無害",
    )


def gen_access97(c: Corpus) -> None:
    """32bit の DAO.DBEngine.36 を SysWOW64 の cscript から呼び、Access 97（Jet3）形式を作る。"""
    rel = "access/meibo_access97.mdb"
    target = c.winpath(rel)
    if Path(target).exists():
        Path(target).unlink()
    if not Path(CSCRIPT32).exists():
        c.skip(rel, "32bit の cscript.exe（SysWOW64）が無い")
        return
    cols = ", ".join(f"[{h}] TEXT(255)" for h in HEADERS)
    names = ", ".join(f"[{h}]" for h in HEADERS)
    lines = [
        "Option Explicit",
        "Dim e, db",
        'Set e = CreateObject("DAO.DBEngine.36")',
        f'Set db = e.CreateDatabase(CStr(WScript.Arguments(0)), "{JA_LOCALE}", 32)',
        f'db.Execute "CREATE TABLE [顧客名簿] (ID COUNTER PRIMARY KEY, {cols})"',
    ]
    for person in PEOPLE:
        values = ", ".join(_sql_quote(person[k]) for k in PERSON_KEYS)
        sql = f"INSERT INTO [顧客名簿] ({names}) VALUES ({values})".replace('"', '""')
        lines.append(f'db.Execute "{sql}"')
    lines += ['WScript.Echo "VERSION=" & db.Version', "db.Close", 'WScript.Echo "OK"']
    vbs = c.work / "make_access97.vbs"
    # WSH は UTF-16 LE（BOM 付き）のスクリプトなら日本語をそのまま扱える
    vbs.write_text("\r\n".join(lines) + "\r\n", encoding="utf-16")
    out = run([CSCRIPT32, "//nologo", str(vbs), target])
    if "OK" not in out:
        c.skip(rel, f"cscript の出力に OK が無い: {out.strip()[:300]}")
        return
    version = out.split("VERSION=")[1].split()[0] if "VERSION=" in out else "?"
    c.add(
        rel,
        "detect",
        ALL_KINDS,
        "mdb (Jet3)",
        f"Access 97 形式（DAO 3.6・dbVersion30・CP932 ロケール・Version={version}）。"
        "ACE（Access 2013 以降）では開けない",
    )


# ---------------------------------------------------------------------------
# 6. メール
# ---------------------------------------------------------------------------


def gen_eml(c: Corpus) -> None:
    """ISO-2022-JP 本文＋CSV 添付の eml。"""
    msg = email.mime.multipart.MIMEMultipart()
    msg["Subject"] = email.header.Header("顧客名簿の送付", "iso-2022-jp")
    msg["From"] = email.utils.formataddr((str(email.header.Header("山田 太郎", "iso-2022-jp")), PEOPLE[0]["email"]))
    msg["To"] = "sales@example.co.jp"
    msg["Date"] = email.utils.formatdate(localtime=True)
    body = "営業部各位\n\n顧客名簿を添付します。\n\n" + pii_text() + "\n"
    msg.attach(email.mime.text.MIMEText(body, "plain", "iso-2022-jp"))
    part = email.mime.base.MIMEBase("text", "csv")
    part.set_payload(c.csv_text().encode("cp932"))
    email.encoders.encode_base64(part)
    part.add_header(
        "Content-Disposition",
        "attachment",
        filename=("iso-2022-jp", "", "顧客名簿.csv"),
    )
    msg.attach(part)
    c.path("mail/meibo.eml").write_bytes(msg.as_bytes())
    c.add(
        "mail/meibo.eml",
        "detect",
        ALL_KINDS,
        "eml",
        "本文 ISO-2022-JP、Shift_JIS の CSV を添付",
    )
    c.add(
        "mail/meibo.eml!顧客名簿.csv",
        "detect",
        ALL_KINDS,
        "csv",
        "eml の添付（Shift_JIS）",
    )


def gen_msg(c: Corpus) -> None:
    """Outlook COM で .msg を作る（送信はしない）。"""
    rel = "mail/meibo.msg"
    if current_session_id() == 0:
        # セッション0から Outlook を COM で起こすと応答が返らず固まる（実測: 90秒で打ち切り、
        # OUTLOOK.EXE がセッション0に残った）。固まると後片付けができないので試さない
        c.skip(
            rel,
            "セッション0（サービス）では Outlook の COM 自動操作が応答せず固まるため作らない。対話セッションで実行すれば作れる",
        )
        return
    try:
        ol = c.app("Outlook.Application", single_instance=True)
    except Exception as exc:  # noqa: BLE001
        c.skip(rel, f"Outlook COM が使えない: {com_error_text(exc)}")
        return
    attach = c.work / "顧客名簿.csv"
    attach.write_bytes(c.csv_text().encode("cp932"))
    mail = ol.CreateItem(0)  # olMailItem
    try:
        mail.Subject = "顧客名簿の送付"
        mail.To = "sales@example.co.jp"
        mail.Body = "営業部各位\n\n顧客名簿を添付します。\n\n" + pii_text()
        mail.Attachments.Add(str(attach))
        mail.SaveAs(c.winpath(rel), 9)  # olMSGUnicode
    finally:
        try:
            mail.Close(1)  # olDiscard（下書きフォルダへは保存しない）
        except Exception:  # noqa: BLE001
            pass
    c.add(
        rel,
        "detect",
        ALL_KINDS,
        "msg",
        "Outlook COM で作成（olMSGUnicode）。CSV 添付つき・未送信",
    )
    c.add(f"{rel}!顧客名簿.csv", "detect", ALL_KINDS, "csv", "msg の添付（Shift_JIS）")


# ---------------------------------------------------------------------------
# 7. 画像
# ---------------------------------------------------------------------------

IMAGE_KINDS = ALL_KINDS


def gen_images(c: Corpus) -> None:
    """日本語の個人情報を文字として描いた画像（OCR 用）。"""
    img = render_text_image(short_pii_lines(), font_size=34)
    for rel, fmt, kwargs in [
        ("image/meibo.png", "PNG", {}),
        ("image/meibo.jpg", "JPEG", {"quality": 92}),
        ("image/meibo.tif", "TIFF", {"compression": "tiff_lzw"}),
    ]:
        img.save(c.path(rel), format=fmt, **kwargs)
        c.add(
            rel,
            "detect",
            IMAGE_KINDS,
            fmt.lower(),
            "文字は画像のみ。OCR しないと読めない",
        )


# ---------------------------------------------------------------------------
# 8. アーカイブ
# ---------------------------------------------------------------------------


class Cp932ZipInfo(zipfile.ZipInfo):
    """ファイル名を UTF-8 フラグなしの cp932 で格納する ZipInfo（日本語版 Windows の旧来の zip を再現）。"""

    def _encodeFilenameFlags(self):  # noqa: N802（標準ライブラリの名前に合わせる）
        return self.filename.encode("cp932"), self.flag_bits & ~0x800


def gen_archives(c: Corpus) -> None:
    """zip / 7z / tar.gz / cab など。"""
    # zip（中に xlsx）
    rel = "archive/meibo.zip"
    with zipfile.ZipFile(c.path(rel), "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("meibo.xlsx", c.xlsx_bytes())
        zf.writestr("readme.txt", "資料一式です。")
    c.add(rel, "detect", ALL_KINDS, "zip")
    c.add(f"{rel}!meibo.xlsx", "detect", ALL_KINDS, "xlsx")
    c.add(f"{rel}!readme.txt", "clean", [], "txt")

    # 日本語ファイル名を cp932 で格納した zip
    rel = "archive/cp932_names.zip"
    with zipfile.ZipFile(c.path(rel), "w", zipfile.ZIP_DEFLATED) as zf:
        for inner, data in [
            ("顧客資料/顧客名簿.xlsx", c.xlsx_bytes()),
            ("顧客資料/名簿.csv", c.csv_text().encode("cp932")),
        ]:
            zi = Cp932ZipInfo(inner, date_time=(2026, 10, 17, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zf.writestr(zi, data)
    c.add(rel, "detect", ALL_KINDS, "zip", "ファイル名が cp932（UTF-8 フラグなし）")
    c.add(
        f"{rel}!顧客資料/顧客名簿.xlsx",
        "detect",
        ALL_KINDS,
        "xlsx",
        "名前は cp932 で格納",
    )
    c.add(
        f"{rel}!顧客資料/名簿.csv",
        "detect",
        ALL_KINDS,
        "csv",
        "名前は cp932・中身 Shift_JIS",
    )

    # zip の中に zip の中に docx（3重）
    rel = "archive/nested3.zip"
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("meibo.docx", c.docx_bytes())
    middle = io.BytesIO()
    with zipfile.ZipFile(middle, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("inner.zip", inner.getvalue())
    with zipfile.ZipFile(c.path(rel), "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("middle.zip", middle.getvalue())
    c.add(rel, "detect", ALL_KINDS, "zip", "zip>zip>zip>docx の3重")
    c.add(f"{rel}!middle.zip!inner.zip!meibo.docx", "detect", ALL_KINDS, "docx")

    # tar.gz
    rel = "archive/meibo.tar.gz"
    with tarfile.open(c.path(rel), "w:gz") as tf:
        for name, data in [
            ("meibo.docx", c.docx_bytes()),
            ("meibo.csv", c.csv_text().encode("utf-8")),
        ]:
            ti = tarfile.TarInfo(name)
            ti.size = len(data)
            ti.mtime = int(datetime(2026, 10, 17).timestamp())
            tf.addfile(ti, io.BytesIO(data))
    c.add(rel, "detect", ALL_KINDS, "tar.gz")
    c.add(f"{rel}!meibo.docx", "detect", ALL_KINDS, "docx")
    c.add(f"{rel}!meibo.csv", "detect", ALL_KINDS, "csv")

    # 7z（py7zr）
    rel = "archive/meibo.7z"
    try:
        import py7zr

        with py7zr.SevenZipFile(c.path(rel), "w") as sz:
            sz.writestr(c.xlsx_bytes(), "meibo.xlsx")
        c.add(rel, "detect", ALL_KINDS, "7z", "py7zr で作成")
        c.add(f"{rel}!meibo.xlsx", "detect", ALL_KINDS, "xlsx")
    except Exception as exc:  # noqa: BLE001
        c.skip(rel, f"py7zr で作れない: {exc}")

    # 7-Zip で作るもの（作業フォルダに中身を置いて相対名で格納する）
    stage = c.work / "stage"
    stage.mkdir(exist_ok=True)
    (stage / "meibo.xlsx").write_bytes(c.xlsx_bytes())
    if not Path(SEVEN_ZIP).exists():
        for r in (
            "archive/zipcrypto_password.zip",
            "archive/password_header_encrypted.7z",
        ):
            c.skip(r, "7z.exe が無い")
    else:
        rel = "archive/zipcrypto_password.zip"
        run(
            [
                SEVEN_ZIP,
                "a",
                "-tzip",
                f"-p{PASSWORD}",
                "-mem=ZipCrypto",
                c.winpath(rel),
                "meibo.xlsx",
            ],
            cwd=stage,
        )
        c.add(
            rel,
            "uncheckable",
            ALL_KINDS,
            "zip (ZipCrypto)",
            "中身の名前は見えるが内容は暗号化。パスワード一覧に secret123 を渡せば detect になるべき",
            password=PASSWORD,
        )
        rel = "archive/aes256_password.zip"
        run(
            [
                SEVEN_ZIP,
                "a",
                "-tzip",
                f"-p{PASSWORD}",
                "-mem=AES256",
                c.winpath(rel),
                "meibo.xlsx",
            ],
            cwd=stage,
        )
        c.add(
            rel,
            "uncheckable",
            ALL_KINDS,
            "zip (AES-256)",
            "WinZip AES 暗号化。標準の zipfile では解けない。パスワード secret123",
            password=PASSWORD,
        )
        rel = "archive/password_header_encrypted.7z"
        run(
            [
                SEVEN_ZIP,
                "a",
                "-t7z",
                f"-p{PASSWORD}",
                "-mhe=on",
                c.winpath(rel),
                "meibo.xlsx",
            ],
            cwd=stage,
        )
        c.add(
            rel,
            "uncheckable",
            ALL_KINDS,
            "7z (AES, header encrypted)",
            "ヘッダーも暗号化（中身の名前も見えない）。パスワード secret123",
            password=PASSWORD,
        )

    # cab（Windows 標準の makecab）
    rel = "archive/meibo.cab"
    makecab = shutil.which("makecab") or r"C:\Windows\System32\makecab.exe"
    if Path(makecab).exists():
        out_dir = str(c.path(rel).parent.resolve())
        run(
            [
                makecab,
                "/D",
                "CompressionType=MSZIP",
                "/L",
                out_dir,
                "meibo.xlsx",
                "meibo.cab",
            ],
            cwd=stage,
        )
        c.add(rel, "detect", ALL_KINDS, "cab", "makecab（MSZIP）で作成")
        c.add(f"{rel}!meibo.xlsx", "detect", ALL_KINDS, "xlsx")
    else:
        c.skip(rel, "makecab.exe が無い")

    # lzh / rar は作る道具が無ければ skipped
    lha = shutil.which("lha") or shutil.which("lhaz")
    if lha:
        rel = "archive/meibo.lzh"
        run([lha, "a", c.winpath(rel), "meibo.xlsx"], cwd=stage)
        c.add(rel, "detect", ALL_KINDS, "lzh")
    else:
        c.skip(
            "archive/meibo.lzh",
            "LZH を作る道具（lha 等）が無い。7-Zip は LZH の展開のみ対応で作成不可",
        )
    rar = shutil.which("rar") or (
        r"C:\Program Files\WinRAR\Rar.exe" if Path(r"C:\Program Files\WinRAR\Rar.exe").exists() else None
    )
    if rar:
        rel = "archive/meibo.rar"
        run([rar, "a", c.winpath(rel), "meibo.xlsx"], cwd=stage)
        c.add(rel, "detect", ALL_KINDS, "rar")
    else:
        c.skip(
            "archive/meibo.rar",
            "RAR を作る道具（WinRAR の Rar.exe）が無い。RAR 形式は作成にライセンスが要る",
        )


# ---------------------------------------------------------------------------
# 9. 偽装
# ---------------------------------------------------------------------------


def gen_disguised(c: Corpus) -> None:
    """拡張子と中身が一致しない検体。"""
    c.path("disguised/meibo_backup.dat").write_bytes(c.xlsx_bytes())
    c.add(
        "disguised/meibo_backup.dat",
        "detect",
        ALL_KINDS,
        "xlsx (ext .dat)",
        "中身は xlsx",
    )
    c.path("disguised/document_noext").write_bytes(c.docx_bytes())
    c.add(
        "disguised/document_noext",
        "detect",
        ALL_KINDS,
        "docx (no ext)",
        "拡張子なし・中身は docx",
    )
    text = "顧客名簿（社外秘）\r\n" + "\r\n".join("\t".join(r) for r in table_rows()) + "\r\n"
    c.path("disguised/meibo_sjis.txt").write_bytes(text.encode("cp932"))
    c.add(
        "disguised/meibo_sjis.txt",
        "detect",
        ALL_KINDS,
        "txt (Shift_JIS)",
        "タブ区切り・Shift_JIS・BOM なし",
    )


# ---------------------------------------------------------------------------
# 10. 暗号化 Office
# ---------------------------------------------------------------------------


def _is_office_encrypted(path: Path) -> bool:
    import msoffcrypto

    with open(path, "rb") as f:
        try:
            return bool(msoffcrypto.OfficeFile(f).is_encrypted())
        except Exception:  # noqa: BLE001
            return False


def _can_decrypt(path: Path, password: str) -> bool:
    import msoffcrypto

    with open(path, "rb") as f:
        try:
            of = msoffcrypto.OfficeFile(f)
            of.load_key(password=password)
            of.decrypt(io.BytesIO())
            return True
        except Exception:  # noqa: BLE001
            return False


def gen_encrypted_office(c: Corpus) -> None:
    """読み取りパスワード付き xlsx と、既定パスワード VelvetSweatshop で暗号化された xls。"""
    # 読み取りパスワード付き docx / xlsx（msoffcrypto で OOXML 暗号化。Office COM に依らない）
    for rel, plain, fmt in [
        ("encrypted/password.docx", c.docx_bytes(), "docx (encrypted)"),
        ("encrypted/password_msoffcrypto.xlsx", c.xlsx_bytes(), "xlsx (encrypted)"),
    ]:
        c.path(rel).write_bytes(encrypt_ooxml(plain, PASSWORD))
        ok = _is_office_encrypted(c.path(rel)) and _can_decrypt(c.path(rel), PASSWORD)
        c.add(
            rel,
            "uncheckable",
            ALL_KINDS,
            fmt,
            f"msoffcrypto で暗号化（Agile, 確認={ok}）。パスワード一覧に secret123 を渡せば detect になるべき",
            password=PASSWORD,
        )

    xl = c.app("Excel.Application")

    # Excel 自身が暗号化した xlsx / xls（Workbook.Password を設定して保存）
    for rel, file_format, fmt in [
        ("encrypted/password.xlsx", 51, "xlsx (encrypted)"),
        ("encrypted/password.xls", 56, "xls (RC4 CryptoAPI)"),
    ]:
        wb = excel_new_workbook(xl, file_format)
        try:
            excel_write_table(wb.Worksheets(1), table_rows())
            wb.Password = PASSWORD
            how = excel_save(wb, c.winpath(rel), file_format)
        finally:
            wb.Close(False)
        ok = _is_office_encrypted(c.path(rel)) and _can_decrypt(c.path(rel), PASSWORD)
        c.add(
            rel,
            "uncheckable",
            ALL_KINDS,
            fmt,
            f"Excel COM で読み取りパスワード付き保存（{how}・暗号化と復号の確認={ok}）。"
            "パスワード一覧に secret123 を渡せば detect になるべき",
            password=PASSWORD,
        )

    # VelvetSweatshop: 読み取りパスワードなしで開けるのに FILEPASS（暗号化）を持つ xls。
    # 実測では「ブック構成の保護（パスワード付き）」で既定パスワード暗号化になった。
    # 念のため複数の作り方を試し、msoffcrypto で「暗号化されていて VelvetSweatshop で解ける」ものを採る。
    rel = "encrypted/velvetsweatshop.xls"
    attempts: list[str] = []
    trials: list[tuple[str, Callable]] = [
        (
            "ブック構成の保護（パスワード付き）",
            lambda wb: wb.Protect("protect123", True, False),
        ),
        (
            "書き込みパスワード（WritePassword）",
            lambda wb: setattr(wb, "WritePassword", "modify123"),
        ),
        (
            "シート保護（パスワード付き）",
            lambda wb: wb.Worksheets(1).Protect("protect123"),
        ),
    ]
    done = False
    for label, action in trials:
        if c.path(rel).exists():
            c.path(rel).unlink()
        wb = excel_new_workbook(xl, 56)
        try:
            excel_write_table(wb.Worksheets(1), table_rows())
            action(wb)
            excel_save(wb, c.winpath(rel), 56)
        except Exception as exc:  # noqa: BLE001
            attempts.append(f"{label}: 例外 {com_error_text(exc)}")
            continue
        finally:
            wb.Close(False)
        enc = _is_office_encrypted(c.path(rel))
        dec = enc and _can_decrypt(c.path(rel), "VelvetSweatshop")
        attempts.append(f"{label}: 暗号化={enc} 既定パスワードで復号={dec}")
        if enc and dec:
            c.add(
                rel,
                "detect",
                ALL_KINDS,
                "xls (VelvetSweatshop)",
                f"既定パスワード VelvetSweatshop で暗号化（作り方: {label}）。パスワードなしで検査できるべき",
            )
            done = True
            break
    if not done:
        if c.path(rel).exists():
            c.path(rel).unlink()
        c.skip(rel, "VelvetSweatshop 暗号化の xls を作れなかった: " + " / ".join(attempts))


# ---------------------------------------------------------------------------
# 11. 対照（clean）
# ---------------------------------------------------------------------------


def gen_clean(c: Corpus) -> None:
    """個人情報を含まない対照（紛らわしい値は入れる）。"""
    import docx
    import openpyxl
    import pymupdf

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "受注一覧"
    for row in clean_rows():
        ws.append(row)
    wb.save(c.path("clean/orders.xlsx"))
    note = "会社名・商品名・日付・金額・JAN・注文番号・検査数字が合わない12桁・Luhn 不成立の16桁"
    c.add("clean/orders.xlsx", "clean", [], "xlsx", note)

    d = docx.Document()
    for line in clean_text().splitlines():
        d.add_paragraph(line)
    d.save(c.path("clean/orders.docx"))
    c.add("clean/orders.docx", "clean", [], "docx", note)

    doc = pymupdf.open()
    _pdf_text_page(doc, clean_text())
    doc.save(c.path("clean/orders.pdf"))
    doc.close()
    c.add("clean/orders.pdf", "clean", [], "pdf", note)

    c.path("clean/orders.txt").write_text(clean_text() + "\n", encoding="utf-8")
    c.add("clean/orders.txt", "clean", [], "txt (UTF-8)", note)


# ---------------------------------------------------------------------------
# 実行
# ---------------------------------------------------------------------------

STEPS: list[tuple[str, Callable[[Corpus], None]]] = [
    ("excel_openpyxl", gen_excel_openpyxl),
    ("excel_com", gen_excel_com),
    ("csv", gen_csv),
    ("word_docx", gen_word_docx),
    ("word_com", gen_word_com),
    ("powerpoint", gen_powerpoint),
    ("pdf", gen_pdf),
    ("access", gen_access),
    ("access97", gen_access97),
    ("eml", gen_eml),
    ("msg", gen_msg),
    ("images", gen_images),
    ("archives", gen_archives),
    ("disguised", gen_disguised),
    ("encrypted_office", gen_encrypted_office),
    ("clean", gen_clean),
]


def prepare_out_dir(out: Path, clean: bool) -> None:
    """出力先を用意する。中身があれば、以前この道具で作ったもの（manifest.json あり）に限り --clean で空にする。"""
    if out.exists() and any(out.iterdir()):
        if not clean:
            raise SystemExit(f"出力先が空ではありません: {out}\n作り直すときは --clean を付けてください。")
        if not (out / "manifest.json").exists():
            raise SystemExit(f"出力先に manifest.json が無いため、この道具の出力とは判断できず消しません: {out}")
        shutil.rmtree(out)
    out.mkdir(parents=True, exist_ok=True)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="個人情報チェッカーの試験用検体を生成する")
    parser.add_argument("out_dir", help="出力先フォルダ（一時フォルダを推奨）")
    parser.add_argument(
        "--clean",
        action="store_true",
        help="以前の出力（manifest.json あり）を消して作り直す",
    )
    parser.add_argument("--only", default="", help="作る段の名前をカンマ区切りで限定（例: pdf,csv）")
    args = parser.parse_args(argv)

    out = Path(args.out_dir).resolve()
    prepare_out_dir(out, args.clean)
    only = {s.strip() for s in args.only.split(",") if s.strip()}
    unknown = only - {name for name, _ in STEPS}
    if unknown:
        raise SystemExit(f"--only に知らない段があります: {', '.join(sorted(unknown))}")

    with tempfile.TemporaryDirectory(prefix="corpus_work_", dir=str(out.parent)) as work:
        c = Corpus(out, Path(work))
        try:
            for name, func in STEPS:
                if only and name not in only:
                    continue
                print(f"[{name}]")
                try:
                    func(c)
                except Exception as exc:  # noqa: BLE001
                    detail = com_error_text(exc) if type(exc).__name__ == "com_error" else str(exc)
                    c.skip(
                        f"<step:{name}>",
                        f"段の途中で失敗（以降のこの段の検体は未作成）: {type(exc).__name__}: {detail}",
                    )
                    c.errors.append(traceback.format_exc())
                finally:
                    # 保存に失敗した Office は後で「呼び出しを拒否」して応答しなくなることがある
                    # （実測）。段ごとに終了し、次の段は新しいプロセスで始める
                    c.quit_apps()
                c.write_manifest()  # 途中で止まっても残るよう段ごとに保存
        finally:
            c.quit_apps()
        c.write_manifest()

    files = [p for p in out.rglob("*") if p.is_file() and p.name != "manifest.json"]
    top_entries = {e["path"] for e in c.entries if "!" not in e["path"]}
    print(
        f"\n完了: 検体ファイル {len(files)} 個 / manifest の entries {len(c.entries)} 件"
        f"（うちファイル単位 {len(top_entries)} 件・アーカイブ内 {len(c.entries) - len(top_entries)} 件）"
        f" / skipped {len(c.skipped)} 件"
    )
    for err in c.errors:
        print(err, file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
