"""PDF・Word（docx）・PowerPoint（pptx）への対処の試験。

守ること: ほかの形式と同じ（検査は何も書き込まない・書き換えは一時の名前の新ファイルから確定・失敗で元を壊さない・記録とジョブの DB に生の値を入れない）に加えて、
- PDF: 黒塗りのあと、検査し直して消え・ページの文字の抽出にも生のバイト列（全オブジェクトを展開して）にも値が無く・関係ない文字は残り・メタデータの値も消える
- docx: 値が複数の run に割れていても置き換わり・段落・表の行・列を削除でき・ほかの部品がバイト単位で同じ・変更履歴・ヘッダー・脚注・コメントの値も消える
- pptx: ダミー化・スライドの削除・ノート
検体はすべて pytest の tmp_path に作る。本物のごみ箱へは送らない（送る関数を差し替える）。画面・Word・Acrobat は開かない。
"""

from __future__ import annotations

import glob
import io
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from pii_checker import remediate  # noqa: E402
from pii_checker.models import ScanOptions  # noqa: E402
from pii_checker.ooxml_text import W, docx_para_text, pptx_para_text, pptx_paragraphs  # noqa: E402
from pii_checker.remedy_write import RemedyError  # noqa: E402
from test_remediate import (  # noqa: E402
    ALL_RANKS,
    FakeTrash,
    SimulatedCrash,
    _leak,
    analyze,
    crash_after,
    make_service,
    new_tmp_files,
    run,
    run_until_crash,
    scan_counts,
)


@pytest.fixture(autouse=True)
def trash(tmp_path, monkeypatch):
    """どの試験でも、ごみ箱へ送る関数を差し替える（本物のごみ箱へ送らない）。ごみ箱のある場所として扱う。"""
    t = FakeTrash(tmp_path / "_trash")
    monkeypatch.setattr(remediate, "send_to_recycle_bin", t)
    monkeypatch.setattr(remediate, "has_recycle_bin", lambda path: True)
    return t


# ---------------------------------------------------------------- docx の検体

NS_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
PHONE_SPLIT = "090-1234-5678"  # 本文で2つの run に割れる値
PHONE_DEL = "080-9999-8888"  # 変更履歴の削除文字列にだけある値
MAIL_BODY = "taro.yamada@example.com"
MAIL_BOX = "textbox.mail@example.com"
MAIL_HEADER = "header.mail@example.com"
PHONE_FOOTER = "070-1111-2222"
MAIL_NOTE = "footnote.mail@example.com"
PHONE_COMMENT = "090-3333-4444"
MAIL_TITLE = "title.mail@example.com"
TABLE_ROWS = [
    ("table1.a@example.com", "090-5000-0001"),
    ("table2.b@example.com", "090-5000-0002"),
    ("table3.c@example.com", "090-5000-0003"),
]
DOCX_NEEDLES = [
    PHONE_SPLIT,
    PHONE_DEL,
    MAIL_BODY,
    MAIL_BOX,
    MAIL_HEADER,
    PHONE_FOOTER,
    MAIL_NOTE,
    PHONE_COMMENT,
    MAIL_TITLE,
    *[v for r in TABLE_ROWS for v in r],
]
STYLES_BYTES = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:styles ' + NS_W.encode() + b"/>\n"
IMAGE_BYTES = b"\x89PNG\r\n\x1a\n" + bytes(range(256)) * 8


def _run(text: str, bold: bool = False) -> str:
    pr = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f'<w:r>{pr}<w:t xml:space="preserve">{text}</w:t></w:r>'


def _para(*runs: str) -> str:
    return "<w:p>" + "".join(runs) + "</w:p>"


def docx_parts() -> dict[str, bytes]:
    """値を、割れた run・表・テキストボックス・変更履歴・ヘッダー・フッター・脚注・コメント・プロパティに散らした docx の部品。"""
    table = "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>"
    table += "<w:tr><w:tc>" + _para(_run("メール")) + "</w:tc><w:tc>" + _para(_run("電話")) + "</w:tc></w:tr>"
    for mail, phone in TABLE_ROWS:
        table += "<w:tr><w:tc>" + _para(_run(mail)) + "</w:tc><w:tc>" + _para(_run(phone)) + "</w:tc></w:tr>"
    table += "</w:tbl>"
    textbox = (
        '<w:p><w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox><w:txbxContent>'
        + _para(_run("ボックス " + MAIL_BOX))
        + "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    tracked = (
        '<w:p><w:del w:id="1" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>'
        + PHONE_DEL
        + '</w:delText></w:r></w:del><w:ins w:id="2" w:author="A" w:date="2026-01-01T00:00:00Z">'
        + _run("新しい番号")
        + "</w:ins></w:p>"
    )
    body = (
        _para(_run("連絡先は "), _run(PHONE_SPLIT[:6], bold=True), _run(PHONE_SPLIT[6:]), _run(" です。"))
        + _para(_run("メール: " + MAIL_BODY))
        + _para(_run("これは関係のない文章です。"))
        + table
        + textbox
        + tracked
        + _para(_run("おわり"))
    )
    sect = '<w:sectPr><w:headerReference w:type="default" r:id="rIdH"/><w:footerReference w:type="default" r:id="rIdF"/></w:sectPr>'
    rid = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    decl = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    doc = f"{decl}<w:document {NS_W} {rid}><w:body>{body}{sect}</w:body></w:document>"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    doc_rels = (
        decl
        + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + f'<Relationship Id="rIdH" Type="{rel}/header" Target="header1.xml"/>'
        + f'<Relationship Id="rIdF" Type="{rel}/footer" Target="footer1.xml"/>'
        + f'<Relationship Id="rIdN" Type="{rel}/footnotes" Target="footnotes.xml"/>'
        + f'<Relationship Id="rIdC" Type="{rel}/comments" Target="comments.xml"/>'
        + f'<Relationship Id="rIdS" Type="{rel}/styles" Target="styles.xml"/>'
        + f'<Relationship Id="rIdI" Type="{rel}/image" Target="media/image1.png"/>'
        + "</Relationships>"
    )
    root_rels = (
        decl
        + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="word/document.xml"/>'
        + '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        + '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail" Target="docProps/thumbnail.jpeg"/>'
        + "</Relationships>"
    )
    ctypes = (
        decl
        + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        + '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        + '<Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>'
        + '<Default Extension="jpeg" ContentType="image/jpeg"/>'
        + '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        + "</Types>"
    )
    core = (
        decl
        + '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        + 'xmlns:dc="http://purl.org/dc/elements/1.1/">'
        + f"<dc:title>連絡 {MAIL_TITLE}</dc:title><dc:creator>Hanako</dc:creator></cp:coreProperties>"
    )

    def part(tag: str, *paras: str) -> bytes:
        return f"{decl}<w:{tag} {NS_W}>{''.join(paras)}</w:{tag}>".encode("utf-8")

    return {
        "[Content_Types].xml": ctypes.encode("utf-8"),
        "_rels/.rels": root_rels.encode("utf-8"),
        "word/document.xml": doc.encode("utf-8"),
        "word/_rels/document.xml.rels": doc_rels.encode("utf-8"),
        "word/styles.xml": STYLES_BYTES,
        "word/header1.xml": part("hdr", _para(_run("ヘッダー " + MAIL_HEADER))),
        "word/footer1.xml": part("ftr", _para(_run("フッター " + PHONE_FOOTER))),
        "word/footnotes.xml": part("footnotes", f'<w:footnote w:id="1">{_para(_run("注 " + MAIL_NOTE))}</w:footnote>'),
        "word/comments.xml": part(
            "comments", f'<w:comment w:id="0" w:author="A">{_para(_run("コメント " + PHONE_COMMENT))}</w:comment>'
        ),
        "word/media/image1.png": IMAGE_BYTES,
        "docProps/core.xml": core.encode("utf-8"),
        "docProps/thumbnail.jpeg": b"\xff\xd8\xff\xe0JFIF-thumbnail",
    }


def write_zip(path: Path, parts: dict[str, bytes], stored: set[str] = frozenset()) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as z:
        for name, data in parts.items():
            zi = zipfile.ZipInfo(name, date_time=(2024, 5, 6, 7, 8, 10))
            zi.compress_type = zipfile.ZIP_STORED if name in stored else zipfile.ZIP_DEFLATED
            z.writestr(zi, data)


def build_docx(tmp_path: Path, name: str = "memo.docx") -> Path:
    p = tmp_path / "src" / name
    write_zip(p, docx_parts())
    return p


def zip_raw(path: Path) -> dict[str, tuple[bytes, int, tuple]]:
    """部品ごとの（圧縮済みのバイト列, 圧縮方式, 日時）。バイト単位で同じかの比較に使う。"""
    out = {}
    with zipfile.ZipFile(path) as z, open(path, "rb") as f:
        for i in z.infolist():
            f.seek(i.header_offset)
            hdr = f.read(30)
            n, m = int.from_bytes(hdr[26:28], "little"), int.from_bytes(hdr[28:30], "little")
            f.seek(i.header_offset + 30 + n + m)
            out[i.filename] = (f.read(i.compress_size), i.compress_type, i.date_time)
    return out


def docx_paragraph_texts(path: Path, part: str = "word/document.xml") -> list[str]:
    from lxml import etree

    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read(part))
    return [docx_para_text(p) for p in root.iter(W + "p")]


def all_text(path: Path) -> str:
    """zip の全部品（XML）を展開してつないだ文字列（値が残っていないかを探す）。"""
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels")))


# ---------------------------------------------------------------- docx: 位置の集め方


def test_docx_positions_are_paragraph_level_and_cover_every_part(tmp_path):
    p = build_docx(tmp_path)
    got = remediate.collect_task({"path": str(p), "options": ScanOptions(temp_dir=str(tmp_path))})
    assert got["layout"] == "docx" and got["status"] == "scanned" and not got["overflow"]
    parts = {k[0] for k in got["positions"]}
    assert parts == {
        "word/document.xml",
        "word/header1.xml",
        "word/footer1.xml",
        "word/footnotes.xml",
        "word/comments.xml",
    }
    # 位置は数値だけ（生の値を持たない）
    assert all(isinstance(v, int) for arr in got["positions"].values() for v in arr)
    assert "taro" not in repr(got["positions"])


def test_docx_paragraph_numbering_matches_the_scanner():
    """検査（formats.office._para_text）と対処（ooxml_text.docx_para_text）の段落の文字の並べ方が同じ。"""
    import xml.etree.ElementTree as ET

    from pii_checker.formats.office import _para_text

    root = ET.fromstring(docx_parts()["word/document.xml"])
    assert [docx_para_text(p) for p in root.iter(W + "p")] == [_para_text(p) for p in root.iter(W + "p")]
    assert any(PHONE_DEL in docx_para_text(p) for p in root.iter(W + "p"))  # 削除文字列も段落の文字


# ---------------------------------------------------------------- docx: ダミー化


UNCHANGED_PARTS = ("word/styles.xml", "word/media/image1.png", "word/_rels/document.xml.rels", "[Content_Types].xml")


@pytest.mark.parametrize("mode", ["same", "shape", "fixed"])
def test_docx_dummy_replaces_split_runs_and_every_part(tmp_path, mode, trash):
    p = build_docx(tmp_path)
    before = zip_raw(p)
    # 陽性対照: 元のファイルには値がある（割れた値は、段落の文字として読んだときだけ見つかる）
    raw = all_text(p)
    assert all(n in raw for n in DOCX_NEEDLES if n != PHONE_SPLIT) and PHONE_SPLIT not in raw
    assert any(PHONE_SPLIT in t for t in docx_paragraph_texts(p))
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="dummy", dummy_mode=mode, confidences=ALL_RANKS)
    assert r["state"] == "done", r
    text = all_text(p)
    for needle in DOCX_NEEDLES:
        assert (
            needle not in text
        ), needle  # 割れた run・削除文字列・ヘッダー・脚注・コメント・テキストボックス・プロパティのどこにも残らない
    assert r["remaining"]["total"] == 0
    # 関係ない文字は残る・書式は値の最初の run
    body = docx_paragraph_texts(p)
    assert "これは関係のない文章です。" in body and "おわり" in body and any(t.endswith("新しい番号") for t in body)
    first = next(t for t in body if t.startswith("連絡先は "))
    assert first.endswith(" です。") and PHONE_SPLIT[:6] not in first
    from lxml import etree

    with zipfile.ZipFile(p) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    bold = [
        r_
        for r_ in root.iter(W + "r")
        if r_.find(W + "rPr") is not None and r_.find(W + "rPr").find(W + "b") is not None
    ]
    assert (
        bold and bold[0].find(W + "t").text and bold[0].find(W + "t").text not in ("090-12",)
    )  # 置き換えは最初の run（太字）に入る
    # ほかの部品は圧縮済みのバイト列・圧縮方式・日時までそのまま
    after = zip_raw(p)
    for name in UNCHANGED_PARTS:
        assert after[name] == before[name], name
    assert "docProps/thumbnail.jpeg" not in after and "docProps/thumbnail.jpeg" in before  # 縮小画像は取り除く
    assert "thumbnail" not in zipfile.ZipFile(p).read("_rels/.rels").decode()
    assert "Hanako" in zipfile.ZipFile(p).read("docProps/core.xml").decode()  # 値の無いプロパティは残す
    assert MAIL_TITLE not in zipfile.ZipFile(p).read("docProps/core.xml").decode()
    assert r["changed"]["cells"] >= 9 and r["changed"]["thumbnails"] == 1
    assert [f.read_bytes() for f in trash.files()] and not new_tmp_files(p.parent)
    _total, values = scan_counts(
        p, tmp_path
    )  # 列見出しが「メール」の列のダミーは、検査し直しに拾われうる。元の値は拾われない
    assert not [v for _t, v in values if v in DOCX_NEEDLES]


def test_docx_value_found_only_through_a_column_header_is_replaced_in_its_own_paragraph(tmp_path):
    """見出し「電話番号」の列にある数字だけの値（本文として読んでも検出されない）を、セルの中の段落ごとの位置で置き換える（セルの値は、複数の段落にまたがる）。"""
    parts = docx_parts()
    keep = ("word/header1.xml", "word/footer1.xml", "word/footnotes.xml", "word/comments.xml")
    parts = {k: v for k, v in parts.items() if k not in keep}
    parts["docProps/core.xml"] = parts["docProps/core.xml"].replace(f"連絡 {MAIL_TITLE}".encode(), b"memo")
    rows = [
        ("備考", "電話番号"),
        ("a", ["内線なし", "09012345678"]),
        ("b", ["０９０１２３４５６７９"]),  # 全角
        ("c", ["関係なし"]),
    ]
    tbl = "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>"
    for name, phone in rows:
        paras = phone if isinstance(phone, list) else [phone]
        tbl += (
            "<w:tr><w:tc>"
            + _para(_run(name))
            + "</w:tc><w:tc>"
            + "".join(_para(_run(t)) for t in paras)
            + "</w:tc></w:tr>"
        )
    tbl += "</w:tbl>"
    doc = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f"<w:document {NS_W}><w:body>{tbl}{_para(_run('おわり'))}</w:body></w:document>"
    )
    parts["word/document.xml"] = doc.encode("utf-8")
    parts["word/_rels/document.xml.rels"] = (
        b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
    )
    p = tmp_path / "src" / "header.docx"
    write_zip(p, parts)
    got = remediate.collect_task({"path": str(p), "options": ScanOptions(temp_dir=str(tmp_path))})
    assert {k[0] for k in got["positions"]} == {
        "word/document.xml"
    }  # 表の列見出しを使った検出も、段落の位置になっている
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="dummy", dummy_mode="fixed", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    from lxml import etree

    with zipfile.ZipFile(p) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    cells = [
        [docx_para_text(x) for x in tc.findall(W + "p")] for tr in root.iter(W + "tr") for tc in tr.findall(W + "tc")
    ]
    # 見出しの列の値はセルの全体が1件の検出（検出の仕様）なので、セルの中の2つの段落とも置き換わる。ほかのセルは元のまま
    assert cells[3] == ["＊＊＊", "＊＊＊"] and cells[5] == ["＊＊＊"] and cells[7] == ["関係なし"]
    assert cells[0] == ["備考"] and cells[2] == ["a"] and cells[4] == ["b"]
    assert "09012345678" not in all_text(p) and "０９０１２３４５６７９" not in all_text(p)


def test_docx_value_in_unreadable_xml_part_and_link_target_is_replaced(tmp_path):
    """段落の外（XML として読めない古い VML・リンク先）に残った、消した値と同じ文字も置き換える。"""
    parts = docx_parts()
    parts["word/drawings/legacy.vml"] = f"<html><br>contact {PHONE_SPLIT} <p>".encode()
    parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(
        b"</Relationships>",
        f'<Relationship Id="rIdM" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        f'Target="mailto:{MAIL_BODY}" TargetMode="External"/></Relationships>'.encode(),
    )
    p = tmp_path / "src" / "links.docx"
    write_zip(p, parts)
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="dummy", dummy_mode="same", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    with zipfile.ZipFile(p) as z:
        vml = z.read("word/drawings/legacy.vml").decode()
        rels = z.read("word/_rels/document.xml.rels").decode()
    assert PHONE_SPLIT not in vml and "contact " in vml and "電話番号_" in vml
    assert MAIL_BODY not in rels and 'Target="mailto:' in rels and 'TargetMode="External"' in rels
    assert r["changed"]["cleaned"] >= 2


# ---------------------------------------------------------------- docx: 段落・表の行・列の削除


def test_docx_delete_paragraphs_keeps_required_paragraphs_empty(tmp_path, trash):
    p = build_docx(tmp_path)
    before = zip_raw(p)
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_paragraphs", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    text = all_text(p)
    for needle in DOCX_NEEDLES:
        assert needle not in text, needle
    body = docx_paragraph_texts(p)
    assert "これは関係のない文章です。" in body and "おわり" in body
    assert not any("連絡先は" in t for t in body)  # 値のあった段落は消えた
    # ヘッダー・フッター・脚注・コメント・セル・テキストボックスは、段落が1つは要るので、中身だけ空になる
    for part in ("word/header1.xml", "word/footer1.xml", "word/footnotes.xml", "word/comments.xml"):
        assert docx_paragraph_texts(p, part) == [""], part
    from lxml import etree

    with zipfile.ZipFile(p) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    for tc in root.iter(W + "tc"):
        assert tc.find(W + "p") is not None  # セルは段落で終わる
    assert all(tc.getparent().tag == W + "tr" for tc in root.iter(W + "tc"))
    after = zip_raw(p)
    for name in UNCHANGED_PARTS:
        assert after[name] == before[name], name
    assert r["changed"]["paragraphs"] >= 9 and r["remaining"]["total"] == 0


def test_docx_delete_rows_only_touches_table_rows(tmp_path):
    p = build_docx(tmp_path)
    svc, store, job = make_service(tmp_path, [p])
    pv_a = analyze(svc, job, p)
    assert (
        "delete_rows" in pv_a.info["methods"] and pv_a.info["methods"]["delete_rows"]["label"] == "表の該当する行を削除"
    )
    pv = svc.preview(job, str(p), {"method": "delete_rows", "confidences": ALL_RANKS})
    assert pv["changes"]["rows"] == 3 and any("表の外にある検出" in w for w in pv["warnings"] + pv["notes"])
    r = run(svc, job, p, method="delete_rows", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    assert r["changed"]["rows"] == 3 and r["changed"]["outside_table"] >= 1
    text = all_text(p)
    for mail, phone in TABLE_ROWS:
        assert mail not in text and phone not in text
    assert MAIL_BODY in text and MAIL_HEADER in text  # 表の外の値は、行の削除では変わらない
    from lxml import etree

    with zipfile.ZipFile(p) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    rows = list(root.iter(W + "tr"))
    assert len(rows) == 1 and docx_para_text(rows[0].find(W + "tc").find(W + "p")) == "メール"  # 見出し行は残る


def build_wide_table_docx(tmp_path: Path, merged: bool = False) -> Path:
    """3列の表（備考・電話・メモ）。電話の列だけに個人情報がある。"""
    parts = docx_parts()
    rows = [("備考", "電話", "メモ")] + [(f"memo{i}", f"090-6000-000{i}", f"note{i}") for i in range(1, 4)]
    tbl = "<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/><w:gridCol/></w:tblGrid>"
    for n, (a, b, c) in enumerate(rows):
        cells = [a, b, c]
        tbl += "<w:tr>"
        for k, v in enumerate(cells):
            pr = ""
            if merged and n == 1 and k == 1:
                pr = '<w:tcPr><w:gridSpan w:val="2"/></w:tcPr>'
            if merged and n == 1 and k == 2:
                continue
            tbl += f"<w:tc>{pr}{_para(_run(v))}</w:tc>"
        tbl += "</w:tr>"
    tbl += "</w:tbl>"
    doc = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<w:document {NS_W}><w:body>{_para(_run("表です"))}{tbl}{_para(_run("おわり"))}</w:body></w:document>'
    )
    keep = ("word/header1.xml", "word/footer1.xml", "word/footnotes.xml", "word/comments.xml")
    parts = {k: v for k, v in parts.items() if k not in keep}
    parts["docProps/core.xml"] = parts["docProps/core.xml"].replace(f"連絡 {MAIL_TITLE}".encode(), b"memo")
    parts["word/document.xml"] = doc.encode("utf-8")
    parts["word/_rels/document.xml.rels"] = (
        b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
    )
    p = tmp_path / "src" / "wide.docx"
    write_zip(p, parts)
    return p


def test_docx_delete_columns_removes_the_grid_column_everywhere(tmp_path):
    p = build_wide_table_docx(tmp_path)
    svc, store, job = make_service(tmp_path, [p])
    analyze(svc, job, p)
    pv = svc.preview(job, str(p), {"method": "delete_columns", "confidences": ALL_RANKS})
    assert pv["changes"]["cols"] == 1 and "1 列" in pv["summary"] and pv["changes"]["detections"] == 3
    r = run(svc, job, p, method="delete_columns", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    assert r["changed"]["cols"] == 1
    from lxml import etree

    with zipfile.ZipFile(p) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    rows = [[docx_para_text(tc.find(W + "p")) for tc in tr.findall(W + "tc")] for tr in root.iter(W + "tr")]
    assert rows == [["備考", "メモ"]] + [
        [f"memo{i}", f"note{i}"] for i in range(1, 4)
    ]  # 電話の列だけが消える（見出しも）
    assert len(list(root.iter(W + "gridCol"))) == 2
    assert "090-6000" not in all_text(p) and r["remaining"]["total"] == 0


def test_docx_delete_columns_refuses_merged_cell_with_detection(tmp_path):
    p = build_wide_table_docx(tmp_path, merged=True)
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_columns", confidences=ALL_RANKS)
    assert r["state"] == "error" and "結合したセル" in r["message"]
    assert p.read_bytes() == original and not new_tmp_files(p.parent) and store.remediations(job) == []


# ---------------------------------------------------------------- pptx


PPT_PHONE = "090-2468-1357"
PPT_MAIL = "slide2.mail@example.com"
PPT_NOTE = "notes.mail@example.com"
PPT_TITLE = "deck.title@example.com"
PPTX_NEEDLES = [PPT_PHONE, PPT_MAIL, PPT_NOTE, PPT_TITLE]


def build_pptx(tmp_path: Path, name: str = "deck.pptx") -> Path:
    """3 枚のスライド。1枚目に割れた run の電話・2枚目にメールと画像とノート・3枚目は関係ない文章と、2枚目へのジャンプのリンク。"""
    from PIL import Image
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.util import Inches
    from lxml import etree

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1, s2, s3 = (prs.slides.add_slide(blank) for _ in range(3))
    tf = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame
    para = tf.paragraphs[0]
    para.add_run().text = "連絡先 " + PPT_PHONE[:6]
    second = para.add_run()
    second.text = PPT_PHONE[6:]
    second.font.bold = True
    s2.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "メール " + PPT_MAIL
    img = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 10, 10)).save(img, "PNG")
    img.seek(0)
    s2.shapes.add_picture(img, Inches(1), Inches(3))
    s2.notes_slide.notes_text_frame.text = "ノート " + PPT_NOTE
    box = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "関係のない文章です"
    rid = s3.part.relate_to(s2.part, RT.SLIDE)
    rpr = box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    link = etree.SubElement(rpr, "{http://schemas.openxmlformats.org/drawingml/2006/main}hlinkClick")
    link.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", rid)
    link.set("action", "ppaction://hlinksldjump")
    prs.core_properties.title = "資料 " + PPT_TITLE
    prs.core_properties.author = "Hanako"
    p = tmp_path / "src" / name
    p.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(p))
    return p


def pptx_slide_texts(path: Path) -> list[str]:
    """スライドの順に、各スライドの段落の文字をつないだもの。"""
    from lxml import etree

    from pii_checker.ooxml_text import pptx_slide_parts

    with zipfile.ZipFile(path) as z:
        out = []
        for _no, part, _notes in pptx_slide_parts(z.read):
            root = etree.fromstring(z.read(part))
            out.append("\n".join(pptx_para_text(p) for _n, _s, p in pptx_paragraphs(root)))
        return out


def test_pptx_positions_cover_slides_and_notes(tmp_path):
    p = build_pptx(tmp_path)
    got = remediate.collect_task({"path": str(p), "options": ScanOptions(temp_dir=str(tmp_path))})
    assert got["layout"] == "pptx" and got["status"] == "scanned"
    sheets = {k[0] for k in got["positions"]}
    assert {s.split(":", 1)[0] for s in sheets} == {"スライド1", "スライド2", "スライド2（ノート）"}
    assert all(isinstance(v, int) for arr in got["positions"].values() for v in arr)


def test_pptx_dummy_replaces_split_runs_and_notes(tmp_path, trash):
    p = build_pptx(tmp_path)
    before = zip_raw(p)
    raw = all_text(p)
    assert (
        PPT_MAIL in raw and PPT_NOTE in raw and PPT_TITLE in raw and PPT_PHONE not in raw
    )  # 陽性対照（電話は割れている）
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="dummy", dummy_mode="same", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    text = all_text(p)
    for needle in PPTX_NEEDLES:
        assert needle not in text, needle
    texts = pptx_slide_texts(p)
    assert texts[0].startswith("連絡先 ") and PPT_PHONE[:6] not in texts[0] and texts[2] == "関係のない文章です"
    assert r["remaining"]["total"] == 0 and r["changed"]["cells"] == 3
    after = zip_raw(p)
    untouched = [
        n for n in before if n.startswith(("ppt/slideLayouts/", "ppt/slideMasters/", "ppt/theme/", "ppt/media/"))
    ]
    assert untouched and all(after[n] == before[n] for n in untouched)  # ほかの部品は圧縮済みのバイト列までそのまま
    assert "Hanako" in zipfile.ZipFile(p).read("docProps/core.xml").decode() and PPT_TITLE not in all_text(p)
    from pptx import Presentation

    assert len(Presentation(str(p)).slides) == 3  # 開ける
    assert [f.read_bytes() for f in trash.files()]


def test_pptx_delete_slides_removes_exclusive_parts_and_dangling_links(tmp_path):
    p = build_pptx(tmp_path)
    before_names = set(zipfile.ZipFile(p).namelist())
    assert any(n.startswith("ppt/media/") for n in before_names) and any(
        n.startswith("ppt/notesSlides/") for n in before_names
    )
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_slides", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    assert r["changed"]["slides"] == 2
    names = set(zipfile.ZipFile(p).namelist())
    assert not any(n.startswith(("ppt/media/", "ppt/notesSlides/")) for n in names)  # 2枚目の画像・ノートも消える
    assert sum(1 for n in names if n.startswith("ppt/slides/slide")) == 1
    text = all_text(p)
    for needle in PPTX_NEEDLES:
        assert needle not in text, needle
    assert pptx_slide_texts(p) == ["関係のない文章です"]
    # 3枚目から2枚目へのジャンプのリンクは、関係ごと外れている（壊れた参照を残さない）
    slide_xml = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")][0]
    assert "hlinkClick" not in zipfile.ZipFile(p).read(slide_xml).decode()
    rels = zipfile.ZipFile(p).read(slide_xml.replace("slides/", "slides/_rels/") + ".rels").decode()
    assert 'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"' not in rels
    ct = zipfile.ZipFile(p).read("[Content_Types].xml").decode()
    assert all(f"/{n}" not in ct for n in before_names - names)
    from pptx import Presentation

    prs = Presentation(str(p))
    assert len(prs.slides) == 1 and r["remaining"]["total"] == 0
    assert PPT_TITLE not in zipfile.ZipFile(p).read("docProps/core.xml").decode()


def test_pptx_delete_every_slide_is_refused_and_original_kept(tmp_path):
    p = build_pptx(tmp_path)
    parts = {}
    with zipfile.ZipFile(p) as z:
        for n in z.namelist():
            parts[n] = z.read(n)
    # 3 枚目にも個人情報を入れて、全部のスライドが対象になるようにする
    s3 = [n for n in parts if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    for n in s3:
        if b"slide3" in n.encode() or n.endswith("slide3.xml"):
            parts[n] = parts[n].replace("関係のない文章です".encode(), "call 090-1111-9999".encode())
    write_zip(p, parts)
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])
    r = run(svc, job, p, method="delete_slides", confidences=ALL_RANKS)
    assert r["state"] == "error" and "すべてのスライド" in r["message"]
    assert p.read_bytes() == original and not new_tmp_files(p.parent) and store.remediations(job) == []


# ---------------------------------------------------------------- 確定の手順（形式を問わず同じ）・記録に生の値を入れない


def _sample(fmt: str, tmp_path: Path):
    """（検体のパス, 実行の指定, 新しい中身に値が無いことを確かめる関数）。"""
    import test_remediate_pdf as pdf_tests

    if fmt == "docx":
        p = build_docx(tmp_path)
        return (
            p,
            {"method": "dummy", "dummy_mode": "same", "confidences": ALL_RANKS},
            lambda q: not any(n in all_text(q) for n in DOCX_NEEDLES),
        )
    if fmt == "pptx":
        p = build_pptx(tmp_path)
        return (
            p,
            {"method": "dummy", "dummy_mode": "same", "confidences": ALL_RANKS},
            lambda q: not any(n in all_text(q) for n in PPTX_NEEDLES),
        )
    p = pdf_tests.build_pdf(tmp_path)
    return (
        p,
        {"method": "redact", "confidences": ALL_RANKS},
        lambda q: not any(pdf_tests.has_bytes(pdf_tests.pdf_dump(q), n) for n in pdf_tests.PDF_NEEDLES),
    )


def _crash_stages():
    return [
        ("a1-元を改名した直後", "rename", 1),
        ("a2-元の改名を記録した直後", "state", 2),
        ("b1-新を元の名前へ改名した直後", "rename", 2),
        ("b2-新の改名を記録した直後", "state", 3),
        ("c1-ごみ箱へ送った直後", "trash", 1),
        ("c2-ごみ箱へ送って記録を済みにした直後", "state", 4),
    ]


@pytest.mark.parametrize("fmt", ["docx", "pptx", "pdf"])
@pytest.mark.parametrize("stage", _crash_stages(), ids=lambda s: s[0])
def test_crash_in_each_commit_stage_is_recovered(tmp_path, monkeypatch, trash, fmt, stage):
    """確定の各段の直後に落ちても、立て直しのあとに、元か新の中身が元の名前で残り、元の中身は必ずどこかに残る。"""
    from pii_checker.remediate import RemedyService, inline_runner

    _, kind, nth = stage
    p, params, clean = _sample(fmt, tmp_path)
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])
    if kind == "rename":
        crash_after(monkeypatch, remediate, "_fs_rename", nth)
    elif kind == "state":
        crash_after(monkeypatch, store, "update_remediation", nth)
    else:

        def trash_then_crash(path):
            trash(path)
            raise SimulatedCrash("trash")

        monkeypatch.setattr(remediate, "send_to_recycle_bin", trash_then_crash)
    with pytest.raises(SimulatedCrash):
        run_until_crash(svc, job, p, params)
    assert len(store.unfinished_remediations()) <= 1  # 最後の記録まで済んでいれば、途中のものは無い
    monkeypatch.undo()
    monkeypatch.setattr(remediate, "send_to_recycle_bin", trash)
    monkeypatch.setattr(remediate, "has_recycle_bin", lambda path: True)
    svc2 = RemedyService(store, runner_fn=inline_runner)
    svc2.recover()
    assert store.unfinished_remediations() == []
    final = p.read_bytes()
    assert final == original or clean(p)  # 元か、値の消えた新が、元の名前で残る
    assert not new_tmp_files(p.parent) and not list(p.parent.glob("*対処前*"))
    assert original in [final] + [f.read_bytes() for f in trash.files()]  # 元の中身はどこかに残っている
    assert svc2.recover() == []


@pytest.mark.parametrize("fmt", ["docx", "pptx", "pdf"])
def test_without_recycle_bin_content_is_replaced_in_one_step(tmp_path, monkeypatch, trash, fmt):
    p, params, clean = _sample(fmt, tmp_path)
    monkeypatch.setattr(remediate, "has_recycle_bin", lambda path: False)
    svc, store, job = make_service(tmp_path, [p])
    pv = svc.preview(job, str(p), params) if analyze(svc, job, p) else None
    assert pv["disposal"] == "replaced" and "元のファイルは残りません" in pv["warnings"][0]
    r = run(svc, job, p, **params)
    assert r["state"] == "done" and r["disposal"] == "replaced", r
    assert clean(p) and trash.sent == [] and not new_tmp_files(p.parent) and not list(p.parent.glob("*対処前*"))


@pytest.mark.parametrize("fmt", ["docx", "pptx", "pdf"])
def test_failure_while_writing_keeps_original_and_leaves_no_record(tmp_path, monkeypatch, trash, fmt):
    from pii_checker import remedy_office, remedy_pdf

    p, params, _clean = _sample(fmt, tmp_path)
    original = p.read_bytes()
    svc, store, job = make_service(tmp_path, [p])

    def boom(*a, **k):
        raise RemedyError("書き込みに失敗した（試験）")

    if fmt == "pdf":
        monkeypatch.setattr(remedy_pdf, "rewrite_pdf_file", boom)
    else:
        monkeypatch.setattr(remedy_office, "rewrite_zip", boom)
    r = run(svc, job, p, **params)
    assert r["state"] == "error" and "試験" in r["message"]
    assert p.read_bytes() == original and not new_tmp_files(p.parent) and trash.sent == []
    assert store.remediations(job) == []


@pytest.mark.parametrize("fmt", ["docx", "pptx", "pdf"])
def test_changed_since_preview_is_refused_for_every_format(tmp_path, fmt):
    p, params, _clean = _sample(fmt, tmp_path)
    svc, store, job = make_service(tmp_path, [p])
    analyze(svc, job, p)
    pv = svc.preview(job, str(p), params)
    if fmt == "pdf":
        with open(p, "ab") as f:
            f.write(b"\n")  # 下見のあとにファイルが変わった
    else:
        with zipfile.ZipFile(p, "a") as z:
            z.writestr("customXml/added.xml", b"<x/>")
    started = svc.start_run(job, str(p), params, True, pv["expected_sha256"])
    svc._runs[started["run_id"]].thread.join(120)
    r = svc.run_status(job, started["run_id"])
    assert r["state"] == "error" and "ファイルが変わりました" in r["message"]
    assert not new_tmp_files(p.parent) and store.remediations(job) == []


def test_records_and_job_db_have_no_raw_values_for_pdf_docx_pptx(tmp_path, trash):
    import test_remediate_pdf as pdf_tests
    from pii_checker.runner import JobRunner

    docx, pptx, pdf = build_docx(tmp_path), build_pptx(tmp_path), pdf_tests.build_pdf(tmp_path)
    needles = sorted({*DOCX_NEEDLES, *PPTX_NEEDLES, *pdf_tests.PDF_NEEDLES})
    svc, store, job = make_service(tmp_path, [docx, pptx, pdf])
    # 実際のジョブの検査（DB に伏せた値の明細が入る）も同じ DB に走らせておく
    full = store.create_job([str(docx.parent)], ScanOptions(temp_dir=str(tmp_path / "scantmp")))
    runner = JobRunner(store, full, workers=1)
    runner.start()
    runner.wait(300)
    for path, params in (
        (docx, {"method": "dummy", "dummy_mode": "same", "confidences": ALL_RANKS}),
        (pptx, {"method": "dummy", "dummy_mode": "shape", "confidences": ALL_RANKS}),
        (pdf, {"method": "redact", "confidences": ALL_RANKS}),
    ):
        assert run(svc, job, path, **params)["state"] == "done"
    db_files = glob.glob(str(tmp_path / "jobs.sqlite3*"))
    assert db_files and _leak(db_files, needles) == []
    for table in ("jobs", "files", "items", "findings", "remediations"):
        for row in store.query(f"SELECT * FROM {table}"):
            for v in tuple(row):
                if isinstance(v, str):
                    assert not any(n in v for n in needles), (table, v[:80])
    recs = store.remediations(job)
    assert sorted(x["method"] for x in recs) == ["dummy", "dummy", "redact"] and all(
        x["backup_path"] is None for x in recs
    )
    # 陽性対照: 生の値を入れた DB なら、この探し方で見つかる
    import sqlite3

    q = tmp_path / "control.sqlite3"
    con = sqlite3.connect(q)
    con.execute("CREATE TABLE t(a TEXT)")
    con.execute("INSERT INTO t VALUES (?)", (f"x {PPT_MAIL} y",))
    con.commit()
    con.close()
    assert set(_leak([str(q)], needles)) == {PPT_MAIL}


# ---------------------------------------------------------------- 書庫の中の docx・pptx・PDF／旧形式／暗号化された Office


def test_docx_pptx_pdf_inside_a_zip_are_handled_like_real_files(tmp_path, trash):
    import test_remediate_pdf as pdf_tests
    from test_remediate_archive import DEFLATED, build_zip, find_item, make_archive_service, run_inner

    docx_p, pptx_p = build_docx(tmp_path), build_pptx(tmp_path)
    pdf_p = pdf_tests.build_pdf(tmp_path)
    archive = tmp_path / "pack" / "bundle.zip"
    archive.parent.mkdir()
    date = (2022, 7, 8, 9, 10, 12)
    build_zip(
        archive,
        [
            ("memo.docx", docx_p.read_bytes(), DEFLATED, date),
            ("deck.pptx", pptx_p.read_bytes(), DEFLATED, date),
            ("form.pdf", pdf_p.read_bytes(), DEFLATED, date),
            ("note.txt", b"nothing here\n", DEFLATED, date),
        ],
    )
    svc, store, job = make_archive_service(tmp_path, archive)

    def inner_file(name: str, suffix: str) -> Path:
        out = tmp_path / f"out-{name}"
        with zipfile.ZipFile(archive) as z:
            out.write_bytes(z.read(name))
        return out

    r, pv = run_inner(
        svc, job, archive, find_item(store, job, archive, "memo.docx"), method="dummy", confidences=ALL_RANKS
    )
    assert r["state"] == "done", r
    assert not any(n in all_text(inner_file("memo.docx", ".docx")) for n in DOCX_NEEDLES)
    r, pv = run_inner(
        svc, job, archive, find_item(store, job, archive, "deck.pptx"), method="delete_slides", confidences=ALL_RANKS
    )
    assert r["state"] == "done", r
    assert pptx_slide_texts(inner_file("deck.pptx", ".pptx")) == ["関係のない文章です"]
    r, pv = run_inner(
        svc, job, archive, find_item(store, job, archive, "form.pdf"), method="redact", confidences=ALL_RANKS
    )
    assert r["state"] == "done", r
    assert PDF_MAIL_INNER not in pdf_text_of(inner_file("form.pdf", ".pdf"))
    with zipfile.ZipFile(archive) as z:
        assert z.read("note.txt") == b"nothing here\n"  # ほかの項目はそのまま


PDF_MAIL_INNER = "taro.yamada@example.com"


def pdf_text_of(path: Path) -> str:
    import pymupdf as fitz

    with fitz.open(str(path)) as doc:
        return "\n".join(pg.get_text("text") for pg in doc)


def test_old_office_formats_and_encrypted_office_only_allow_file_delete():
    for kind in ("Word 97-2003（doc）", "PowerPoint 97-2003（ppt）"):
        got = {"status": "scanned", "kind": kind, "layout": "", "overflow": False}
        reason = remediate.unsupported_reason("dummy", "old.doc", got)
        assert "旧形式" in reason and "ファイルの削除だけ" in reason and "docx" in reason
        assert remediate.unsupported_reason("delete_file", "old.doc", got) == ""
    got = {"status": "scanned", "kind": "暗号化された Office 文書（docx/xlsx/pptx）", "layout": "", "overflow": False}
    assert "暗号化された Office" in remediate.unsupported_reason("dummy", "secret.docx", got)
    # 暗号化された docx が、検査のとき入れたパスワード無しの状態: 理由つきで、方法の一覧には docx の方法が出る
    got = {"status": "needs_password", "kind": "Word（docx）", "layout": "", "overflow": False}
    assert "パスワード" in remediate.unsupported_reason("dummy", "x.docx", got)
    assert remediate.layout_from_kind("Word（docx）") == "docx" and remediate.layout_from_kind("PDF") == "pdf"


def test_methods_listed_per_format_and_old_methods_unchanged(tmp_path):
    p = build_docx(tmp_path)
    svc, store, job = make_service(tmp_path, [p])
    methods = analyze(svc, job, p).info["methods"]
    assert set(methods) == {"delete_file", "dummy", "delete_paragraphs", "delete_rows", "delete_columns"}
    assert all(m["available"] for m in methods.values())
    q = build_pptx(tmp_path)
    svc, store, job = make_service(tmp_path, [q])
    assert set(analyze(svc, job, q).info["methods"]) == {"delete_file", "dummy", "delete_slides"}


def test_real_worker_process_round_trip_for_docx_and_pptx(tmp_path, trash):
    """本番と同じ別プロセス（pebble・メモリ上限）で、docx・pptx の分析・書き換えが通る。"""
    docx, pptx = build_docx(tmp_path), build_pptx(tmp_path)
    svc, store, job = make_service(tmp_path, [docx, pptx], runner=remediate._run_in_process)
    r = run(svc, job, docx, method="dummy", dummy_mode="same", confidences=ALL_RANKS)
    assert r["state"] == "done" and r["remaining"]["total"] == 0, r
    assert not any(n in all_text(docx) for n in DOCX_NEEDLES)
    r = run(svc, job, pptx, method="delete_slides", confidences=ALL_RANKS)
    assert r["state"] == "done", r
    assert pptx_slide_texts(pptx) == ["関係のない文章です"]
