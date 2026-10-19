"""レポートを少しずつ読んで書く作り（全行をメモリに載せない）の試験。"""

from __future__ import annotations

import json

from openpyxl import load_workbook

from pii_checker import cli
from pii_checker.detectors import PII_TYPES
from pii_checker.report import _counts_by_top
from pii_checker.store import Store


def _scan_folder(tmp_path):
    root = tmp_path / "root"
    # 並びの前後が入れ替わりやすい名前（ASCII・日本語・入れ子・検出なし）を混ぜる
    files = {
        "a.txt": "顧客 携帯 090-1234-5678\n",
        "b.txt": "会議の議題\n",
        "あ名簿.txt": "氏名 山田太郎 様 携帯 090-2222-3333 メール taro@example.com\n",
        "sub/z.txt": "連絡先 090-4444-5555\n",
        "sub/い.txt": "議事録\n",
        "Z.txt": "携帯 090-6666-7777\n",
    }
    for rel, text in files.items():
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")
    db = tmp_path / "jobs.sqlite3"
    out = tmp_path / "out" / "r.xlsx"
    out.parent.mkdir()
    cli.main(["--db", str(db), "scan", str(root), "--no-ask", "-o", str(out)])
    store = Store(str(db))
    return store, store.jobs()[0]["id"], out


def test_counts_by_top_matches_whole_dict(tmp_path):
    """ファイルごとの件数を、全ファイル分の辞書で数えたときと同じに返す（呼ばれないファイルが挟まっても崩れない）。"""
    store, job_id, _ = _scan_folder(tmp_path)
    want: dict[str, dict[str, int]] = {}
    for r in store.query("SELECT top_path, counts FROM items WHERE job_id=? AND counts != '{}'", (job_id,)):
        d = want.setdefault(r["top_path"], {})
        for k, n in json.loads(r["counts"]).items():
            d[k] = d.get(k, 0) + n
    tops = [
        r["top_path"] for r in store.query("SELECT top_path FROM files WHERE job_id=? ORDER BY top_path", (job_id,))
    ]
    assert len(tops) == 6 and len(want) >= 4
    counts_of = _counts_by_top(store, job_id)
    assert {t: counts_of(t) for t in tops} == {t: want.get(t, {}) for t in tops}
    # 途中のファイルを飛ばして呼んでも、後ろのファイルの件数は正しい
    counts_of = _counts_by_top(store, job_id)
    assert counts_of(tops[-1]) == want.get(tops[-1], {})


def test_excel_file_sheet_counts_match_db(tmp_path):
    """Excel の「ファイル別」の種類別の件数が、保存した検出の件数と一致する。"""
    store, job_id, out = _scan_folder(tmp_path)
    ws = load_workbook(out, read_only=True)["ファイル別"]
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[0])
    phone = head.index(PII_TYPES["phone"])
    got = {r[0]: r for r in rows[1:] if r[0]}
    assert len(got) == 6
    want: dict[str, int] = {}
    for r in store.query("SELECT top_path, counts FROM items WHERE job_id=? AND counts != '{}'", (job_id,)):
        n = json.loads(r["counts"]).get("phone", 0)
        if n:
            want[r["top_path"]] = want.get(r["top_path"], 0) + n
    assert len(want) == 4  # a・あ名簿・sub/z・Z
    assert {t: got[t][phone] for t in want} == want
    assert all(got[t][phone] in ("", None) for t in got if t not in want)


def test_iter_query_reads_while_writer_is_open(tmp_path):
    """検査中（書き込みの接続が開いたまま）でも、別の接続で少しずつ読める。"""
    store, job_id, _ = _scan_folder(tmp_path)
    rows = list(store.iter_query("SELECT top_path FROM files WHERE job_id=? ORDER BY top_path", (job_id,), batch=2))
    assert len(rows) == 6
